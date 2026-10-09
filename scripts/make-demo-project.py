# Copyright (C) 2026 Kirill Borovoy
# SPDX-License-Identifier: GPL-3.0-or-later
"""Build a manual-test project: Demo/project.stanok + Данные + Шаблоны.

Usage: .venv/Scripts/python scripts/make-demo-project.py [dest_dir]
Output dir is gitignored (see demo/ in .gitignore).
"""

from __future__ import annotations

import json
import sys
from pathlib import Path


def main(dest: Path) -> None:
    from docx import Document as DocxDocument
    from openpyxl import Workbook

    data_dir = dest / "Данные"
    tpl_dir = dest / "Шаблоны"
    data_dir.mkdir(parents=True, exist_ok=True)
    tpl_dir.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.append(["ФИО", "Сумма"])
    ws.append(["Иван", 100])
    ws.append(["Петр", 200])
    wb.save(data_dir / "data.xlsx")

    doc = DocxDocument()
    doc.add_paragraph("Договор {{ФИО}} на сумму {{Сумма}} №{{Номер}}")
    doc.save(tpl_dir / "tpl.docx")

    pj = {
        "version": "0.0.0",
        "templates": {
            "Договор": {
                "file": "Шаблоны/tpl.docx",
                "fields": {
                    "ФИО": {"source": "table", "value": "ФИО"},
                    "Сумма": {"source": "table", "value": "Сумма"},
                    "Номер": {"source": "counter", "value": "num"},
                },
            }
        },
        "data_sources": [
            {"file": "Данные/data.xlsx", "mode": "sequential", "start_row": 0}
        ],
        "counters": {"num": {"last": 0, "format": "plain"}},
        "filename_template": "Договор_{ФИО}_{i}.docx",
    }
    (dest / "project.stanok").write_text(
        json.dumps(pj, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"demo project ready: {dest}")


if __name__ == "__main__":
    main(Path(sys.argv[1]) if len(sys.argv) > 1 else Path("demo"))
