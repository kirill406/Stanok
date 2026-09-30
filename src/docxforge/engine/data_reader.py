# -*- coding: utf-8 -*-
"""Read Excel files into list-of-dicts format used by the renderer."""

import logging
import os
from datetime import date, datetime, time
from typing import Any, Dict, List
import openpyxl

from docxforge.engine.schema import BatchSourceConfig, RowIterationMode


logger = logging.getLogger(__name__)


def _coerce_cell(val: Any) -> Any:
    """Normalize one Excel cell without destroying its type (M6).

    Old code did ``str(val).strip()`` for everything, mangling dates
    (``datetime`` → ``'2024-05-06 00:00:00'``), integral floats (``10000.0``
    → ``'10000.0'``) and bools. New contract: ``None`` → ``''``; ``bool``,
    ``int``, ``datetime``/``date``/``time`` pass through natively; integral
    floats become ``int``; only ``str`` values are stripped. Anything else
    is kept as-is (render boundaries apply ``str()`` when needed).
    """
    if val is None:
        return ''
    if isinstance(val, bool):
        return val
    if isinstance(val, int):
        return val
    if isinstance(val, float):
        return int(val) if val.is_integer() else val
    if isinstance(val, (datetime, date, time)):
        return val
    if isinstance(val, str):
        return val.strip()
    return val


class DataReader:
    """Reads Excel data files."""

    def read_excel(self, path: str, sheet_name: str = None) -> List[Dict[str, Any]]:
        """Read an Excel file and return list of rows as dicts.

        First row is treated as headers.
        Rows where all cells are empty are skipped.
        Cell types are preserved (see :func:`_coerce_cell`).

        Raises:
            FileNotFoundError: the file does not exist (M6: absence is
                an explicit error, distinct from a corrupt file which
                logs an error and yields ``[]``).
        """
        if not os.path.exists(path):
            raise FileNotFoundError(f'Excel file not found: {path}')
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
        except PermissionError as e:
            logger.error(f"Permission denied reading {path}: {e}")
            return []
        except Exception as e:
            logger.exception(f"Error reading {path}: {e}")
            return []

        if sheet_name:
            try:
                ws = wb[sheet_name]
            except KeyError as e:
                logger.error(f"Sheet {sheet_name!r} not found in {path}: {e}")
                wb.close()
                return []
        else:
            ws = wb.active

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            wb.close()
            return []

        # Find header row (first non-empty row)
        header = None
        data_start = 0
        for i, row in enumerate(rows):
            # Check if this row looks like a header
            non_empty = sum(1 for c in row if c is not None)
            if non_empty >= 1:
                header = [str(c).strip() if c is not None else f'col_{j}'
                          for j, c in enumerate(row)]
                data_start = i + 1
                break

        if header is None:
            wb.close()
            return []

        result = []
        for row in rows[data_start:]:
            values = {}
            empty_count = 0
            for j, val in enumerate(row):
                if j < len(header):
                    if val is None:
                        values[header[j]] = ''
                        empty_count += 1
                    else:
                        values[header[j]] = _coerce_cell(val)
            # Skip completely empty rows
            if empty_count < len(header):
                result.append(values)

        wb.close()
        return result

    def get_columns(self, path: str, sheet_name: str = None) -> List[str]:
        """Get list of column names from an Excel file.

        Raises:
            FileNotFoundError: the file does not exist (see :meth:`read_excel`).
        """
        if not os.path.exists(path):
            raise FileNotFoundError(f'Excel file not found: {path}')
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
        except PermissionError as e:
            logger.error(f"Permission denied reading {path}: {e}")
            return []
        except Exception as e:
            logger.exception(f"Error reading {path}: {e}")
            return []

        if sheet_name:
            try:
                ws = wb[sheet_name]
            except KeyError as e:
                logger.error(f"Sheet {sheet_name!r} not found in {path}: {e}")
                wb.close()
                return []
        else:
            ws = wb.active

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            wb.close()
            return []

        for row in rows:
            non_empty = sum(1 for c in row if c is not None)
            if non_empty >= 1:
                cols = [str(c).strip() if c is not None else '' for c in row]
                wb.close()
                return cols

        wb.close()
        return []

    def get_distinct_values(self, path: str, column: str,
                             sheet_name: str = None) -> List[str]:
        """Get distinct values from a column for dropdown lists."""
        data = self.read_excel(path, sheet_name)
        seen = set()
        result = []
        for row in data:
            val = row.get(column, '')
            if val and val not in seen:
                seen.add(val)
                result.append(val)
        return result

    def read_all_batch_sources(self, project_dir: str,
                                 batch_sources: Dict[str, BatchSourceConfig]) -> Dict[str, List[Dict[str, Any]]]:
        """Read all batch source files and return dict of file -> rows."""
        result = {}
        for source_name, bsc in batch_sources.items():
            if bsc.mode == RowIterationMode.SEQUENTIAL:
                file_path = os.path.join(project_dir, 'Данные', bsc.file)
                if os.path.exists(file_path):
                    result[bsc.file] = self.read_excel(file_path)
                else:
                    logger.warning(f"Batch source file not found: {file_path}")
                    result[bsc.file] = []
        return result
