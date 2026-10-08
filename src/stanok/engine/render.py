# Copyright (C) 2026 Kirill Borovoy
# SPDX-License-Identifier: GPL-3.0-or-later
"""Render: FillingJSON + .docx template → Document. No disk writes."""

import logging
from pathlib import Path

from docx import Document

from .schema import FillingJSON, RenderError
from .xmlops import _stringify, substitute_paragraph

logger = logging.getLogger(__name__)


def _iter_paragraphs(doc: Document):
    yield from doc.paragraphs
    for table in doc.tables:
        yield from _iter_table_paragraphs(table)


def _iter_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from _iter_table_paragraphs(nested)


def render(fj: FillingJSON, template_path: Path) -> Document:
    """Substitute fj.fields into template, return in-memory document."""
    try:
        doc = Document(str(template_path))
    except Exception as e:
        logger.error(f"render open {template_path}: {e}", exc_info=True)
        raise RenderError("template", [f"cannot open template: {e}"]) from e

    mapping = {k: _stringify(v) for k, v in fj.fields.items()}
    try:
        for paragraph in _iter_paragraphs(doc):
            substitute_paragraph(paragraph, mapping)
    except RenderError:
        raise
    except Exception as e:
        logger.error(f"render {template_path}: {e}", exc_info=True)
        raise RenderError("template", [str(e)]) from e
    return doc
