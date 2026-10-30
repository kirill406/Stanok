# Copyright (C) 2026 Kirill Borovoy
# SPDX-License-Identifier: GPL-3.0-or-later
"""Project storage: atomic save/load, Home migration, recent projects."""

from __future__ import annotations

import codecs
import json
import logging
import os
import shutil
import threading
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from ..engine.schema import (
    ApplicationJSON,
    CounterDef,
    DataSourceDef,
    FieldDef,
    FieldSource,
    FormatTooNewError,
    MigrationError,
    ProjectJSON,
    RecentItem,
    SchemaError,
    StorageError,
    TemplateDef,
    current_version,
    migrate,
    validate_aj,
    validate_pj,
)
from ..engine.resolve import COUNTER_WIDTHS
from ..engine.xmlops import PLACEHOLDER_RE, iter_paragraphs
from ..gui.strings import STRINGS

logger = logging.getLogger(__name__)


def _remove_tmp(path: Path) -> None:
    """Best-effort removal of a temp file after a failed write."""
    if not path.exists():
        return
    try:
        path.unlink()
    except OSError as e:
        logger.warning(f"remove temp {path}: {e}", exc_info=True)


def project_folder(
    ref: str | Path, pj: ProjectJSON, config_path: Path, store: ProjectStore
) -> Path:
    """Find project folder: explicit dir ref, else PJ.folder, else recent."""
    ref_path = Path(ref) if isinstance(ref, str) else ref
    if ref_path.exists() and ref_path.is_dir():
        return ref_path.resolve()
    if pj.folder:
        folder = Path(pj.folder)
        if folder.is_dir():
            return folder.resolve()
        raise StorageError(
            "resolve", [f"project folder not found: {pj.folder}"]
        )
    # Legacy config without folder: folder remembered in recent projects.
    for item in store.get_recent():
        if item.config == config_path.stem:
            folder = Path(item.folder)
            if folder.is_dir():
                return folder.resolve()
    raise StorageError(
        "resolve",
        [f"project folder not found for '{ref}': pass a project folder path"],
    )


def _scan_placeholders(docx_path: Path) -> list[str]:
    """Ordered unique {{name}} placeholders from docx paragraphs + tables."""
    from docx import Document

    try:
        doc = Document(str(docx_path))
    except (OSError, ValueError, KeyError) as e:
        logger.warning(f"unreadable template {docx_path}: {e}", exc_info=True)
        raise StorageError(
            "init", [f"cannot read template {docx_path.name}: {e}"]
        ) from e
    except Exception as e:
        logger.critical(
            f"unexpected unreadable template {docx_path}: {e}", exc_info=True
        )
        raise StorageError(
            "init", [f"cannot read template {docx_path.name}: {e}"]
        ) from e
    names: list[str] = []
    for p in iter_paragraphs(doc):
        for m in PLACEHOLDER_RE.finditer(p.text):
            name = m.group(1).strip()
            if name and name not in names:
                names.append(name)
    return names


def _unique_sibling(path: Path) -> Path:
    """Return path or first free `stem (1).suffix` sibling."""
    if not path.exists():
        return path
    i = 1
    while True:
        candidate = path.with_name(f"{path.stem} ({i}){path.suffix}")
        if not candidate.exists():
            return candidate
        i += 1


def _counter_name(existing: set[str], file: str) -> str:
    """Auto-counter name from a source file stem; `_2`, `_3` on collision (021)."""
    stem = Path(file).stem or "counter"
    name = stem
    i = 1
    while name in existing:
        i += 1
        name = f"{stem}_{i}"
    return name


def _first_bound_counter(pj: ProjectJSON) -> str | None:
    """First counter bound to a data source (default target for new fields)."""
    for name, cdef in pj.counters.items():
        if cdef.table is not None:
            return name
    return None


@dataclass
class ImportReport:
    """Result of importing .stanok files from one folder."""

    imported: list[str] = field(default_factory=list)
    failed: list[str] = field(default_factory=list)
    missing_folder: list[str] = field(default_factory=list)
    found: int = 0


class ProjectStore:
    """Manages project configs in ~/.stanok/ with atomic operations."""

    def __init__(self, home_dir: Path | None = None):
        if home_dir is None:
            home_dir = Path.home() / ".stanok"
        self.home_dir = Path(home_dir).resolve()
        # RLock: init_project re-enters via save()/add_recent().
        self._lock = threading.RLock()
        self._ensure_home()

    def _ensure_home(self) -> None:
        self.home_dir.mkdir(parents=True, exist_ok=True)
        settings_file = self.home_dir / "settings.json"
        if not settings_file.exists():
            default_aj = ApplicationJSON(
                version=current_version(), recent=[], settings={}
            )
            self._atomic_write_json(
                settings_file, default_aj.model_dump(mode="json")
            )

    def _validate_path(self, path: Path, field: str = "path") -> Path:
        """Validate path: no traversal, must be under home_dir."""
        if ".." in path.parts:
            raise StorageError(field, [f"path traversal not allowed: {path}"])
        resolved = path.resolve()
        try:
            resolved.relative_to(self.home_dir)
        except ValueError:
            raise StorageError(
                field, [f"path must be under home directory: {path}"]
            )
        return resolved

    @staticmethod
    def _backup_path(path: Path) -> Path:
        """`<name>.stanok` → `<name>.stanok.bak`."""
        return path.with_suffix(path.suffix + ".bak")

    @staticmethod
    def _tmp_path(path: Path) -> Path:
        """`<name>.stanok` → `<name>.stanok.tmp`."""
        return path.with_suffix(path.suffix + ".tmp")

    def _has_config(self, path: Path) -> bool:
        """Whether the main config or a recoverable copy exists."""
        return any(
            candidate.exists()
            for candidate in (path, self._tmp_path(path), self._backup_path(path))
        )

    def _atomic_write_json(self, path: Path, data: dict, backup: bool = True) -> None:
        """Write JSON atomically: tmp + fsync -> backup copy + replace.

        `backup=False` skips the copy, so the previous `.bak` survives writes
        that are not user saves (for example, a version relabel on open).
        """
        path = self._validate_path(path)
        path.parent.mkdir(parents=True, exist_ok=True)

        tmp_path = self._tmp_path(path)
        bak_path = self._backup_path(path)

        try:
            with open(tmp_path, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
                f.flush()
                os.fsync(f.fileno())
        except (OSError, TypeError, ValueError) as e:
            logger.error(f"write temp config {tmp_path}: {e}", exc_info=True)
            _remove_tmp(tmp_path)
            raise
        except Exception as e:
            logger.critical(
                f"unexpected write temp config {tmp_path}: {e}", exc_info=True
            )
            _remove_tmp(tmp_path)
            raise

        if backup and path.exists():
            # Copy, not rename: the main config must exist at every moment.
            try:
                shutil.copy2(path, bak_path)
            except OSError as e:
                logger.warning(f"backup config {path}: {e}", exc_info=True)
            except Exception as e:
                logger.critical(
                    f"unexpected backup config {path}: {e}", exc_info=True
                )

        try:
            os.replace(tmp_path, path)
        except OSError as e:
            logger.warning(
                f"atomic replace failed for {path}: {e}", exc_info=True
            )
            try:
                shutil.copy2(tmp_path, path)
                tmp_path.unlink()
            except OSError as copy_error:
                logger.warning(
                    f"copy config fallback {path}: {copy_error}; "
                    f"keeping temp {tmp_path}",
                    exc_info=True,
                )
                raise copy_error from e
            except Exception as copy_error:
                logger.critical(
                    f"unexpected copy config fallback {path}: {copy_error}",
                    exc_info=True,
                )
                raise
        logger.debug(f"wrote config atomically: {path}")

    def _read_pj(self, path: Path, persist: bool = True) -> ProjectJSON:
        """Read a PJ file: migrate an older format, then validate.

        With `persist`, a migrated file is rewritten with the current version.
        A failed rewrite (for example, a read-only file) only logs a warning.
        """
        raw = self._read_json(path)
        migrated = migrate(raw)
        pj = validate_pj(migrated)
        if persist and migrated != raw:
            try:
                self._atomic_write_json(path, migrated, backup=False)
            except OSError as e:
                logger.warning(f"cannot save migrated config {path}: {e}", exc_info=True)
            else:
                logger.info(
                    f"migrated config {path.name}: "
                    f"{raw.get('version')} -> {migrated.get('version')}"
                )
        return pj

    def _read_json(self, path: Path) -> dict:
        path = self._validate_path(path)
        raw = path.read_bytes()
        logger.debug(f"read json {path}: {len(raw)} bytes")
        if raw.startswith(codecs.BOM_UTF8):
            payload = raw.removeprefix(codecs.BOM_UTF8)
            data = json.loads(payload)
            self._rewrite_without_bom(path, payload)
            logger.info(f"removed UTF-8 BOM from {path}")
            return data
        return json.loads(raw)

    def _rewrite_without_bom(self, path: Path, payload: bytes) -> None:
        """Atomically rewrite a config as UTF-8 without BOM; raise on failure."""
        path = self._validate_path(path)
        tmp_path = self._tmp_path(path)
        try:
            with open(tmp_path, "wb") as f:
                f.write(payload)
                f.flush()
                os.fsync(f.fileno())
            os.replace(tmp_path, path)
        except (OSError, TypeError, ValueError) as e:
            logger.warning(f"rewrite without BOM {path}: {e}", exc_info=True)
            _remove_tmp(tmp_path)
            raise
        except Exception as e:
            logger.critical(
                f"unexpected rewrite without BOM {path}: {e}", exc_info=True
            )
            _remove_tmp(tmp_path)
            raise

    def config_exists(self, name: str) -> bool:
        """Whether a project config (or its recoverable copy) exists in Home."""
        return self._has_config(self._get_config_path(name))

    def _get_config_path(self, name: str) -> Path:
        """Get path for config file in Home."""
        safe_name = name.replace("/", "_").replace("\\", "_")
        if not safe_name:
            safe_name = "project"
        return self.home_dir / f"{safe_name}.stanok"

    def _extract_project_name(self, pj: ProjectJSON) -> str:
        """Extract project name from PJ or use fallback."""
        template = pj.filename_template or ""
        name = template.split("{")[0].strip()
        if name:
            return name.replace("/", "_").replace("\\", "_")
        if pj.templates:
            first_template = next(iter(pj.templates.values()))
            file = getattr(first_template, "file", "")
            if file:
                stem = Path(file).stem
                if stem:
                    return stem
        return "project"

    def save(self, pj: ProjectJSON, name: str | None = None) -> Path:
        """Save project config atomically, return path."""
        with self._lock:
            path = self._get_config_path(name or self._extract_project_name(pj))
            logger.debug(f"saving project config: name={name} path={path}")
            self._atomic_write_json(path, pj.model_dump(mode="json"))
            logger.info(f"saved project config to {path}")
            return path

    def load(self, name: str) -> ProjectJSON:
        """Load a project config; recover it from `.tmp`/`.bak` if unreadable."""
        with self._lock:
            path = self._get_config_path(name)
            logger.debug(f"loading project config: name={name} path={path}")
            if not path.exists():
                recovered = self._recover_config(path, name)
                if recovered is not None:
                    return recovered
                raise StorageError("load", [f"project not found: {name}"])
            try:
                pj = self._read_pj(path)
            except StorageError:
                raise
            except FormatTooNewError as e:
                # Not recovered from .tmp/.bak: the file is valid, just not for this program.
                logger.error(f"load config {path}: {e}")
                raise StorageError("load", [f"config {name}: {e}"]) from e
            except (OSError, ValueError, SchemaError) as e:
                logger.error(f"load config {path}: {e}", exc_info=True)
                recovered = self._recover_config(path, name)
                if recovered is not None:
                    return recovered
                raise StorageError("load", [f"invalid config {name}: {e}"]) from e
            except Exception as e:
                logger.critical(
                    f"unexpected load config {path}: {e}", exc_info=True
                )
                raise StorageError("load", [f"invalid config {name}: {e}"]) from e
            logger.debug(
                f"loaded project {name}: version={pj.version} "
                f"templates={len(pj.templates)} sources={len(pj.data_sources)}"
            )
            return pj

    def _recover_config(self, path: Path, name: str) -> ProjectJSON | None:
        """Recover an unreadable config: newest valid `.tmp`, then `.bak`."""
        for source in (self._tmp_path(path), self._backup_path(path)):
            if not source.exists():
                continue
            try:
                pj = self._read_pj(source, persist=False)
            except (OSError, ValueError, SchemaError) as e:
                logger.warning(f"recovery source {source}: {e}", exc_info=True)
                continue
            except Exception as e:
                logger.critical(
                    f"unexpected recovery source {source}: {e}", exc_info=True
                )
                continue
            logger.warning(f"restored project config {name} from {source}")
            self._promote_recovered(source, path, pj)
            return pj
        return None

    def _promote_recovered(self, source: Path, path: Path, pj: ProjectJSON) -> None:
        """Make the recovered copy the main config; keep `.bak` as a backup."""
        if source == self._tmp_path(path):
            # The interrupted save already wrote a complete tmp file.
            try:
                os.replace(source, path)
                logger.debug(f"promoted recovered config {source} to {path}")
                return
            except OSError as e:
                logger.warning(f"promote {source} to {path}: {e}", exc_info=True)
                return
            except Exception as e:
                logger.critical(
                    f"unexpected promote {source} to {path}: {e}", exc_info=True
                )
                return
        self._write_recovered(path, pj)

    def _write_recovered(self, path: Path, pj: ProjectJSON) -> None:
        """Rewrite a corrupt/missing main config from the loaded backup."""
        tmp_path = self._tmp_path(path)
        try:
            path = self._validate_path(path)
            with open(tmp_path, "w", encoding="utf-8") as f:
                json.dump(pj.model_dump(mode="json"), f, ensure_ascii=False, indent=2)
                f.flush()
                os.fsync(f.fileno())
            os.replace(tmp_path, path)
        except (OSError, TypeError, ValueError) as e:
            logger.warning(f"restore config {path}: {e}", exc_info=True)
            _remove_tmp(tmp_path)
        except Exception as e:
            logger.critical(
                f"unexpected restore config {path}: {e}", exc_info=True
            )
            _remove_tmp(tmp_path)

    def delete(self, name: str) -> None:
        """Delete project config and its recoverable copies."""
        with self._lock:
            path = self._get_config_path(name)
            for candidate in (
                path,
                self._tmp_path(path),
                self._backup_path(path),
            ):
                if candidate.exists():
                    candidate.unlink()
                    logger.debug(f"deleted config file {candidate}")
            logger.info(f"deleted project config {name}")

    def list(self) -> list[ProjectJSON]:
        """List all project configs sorted by mtime (newest first)."""
        with self._lock:
            configs = []
            for path in self.home_dir.glob("*.stanok"):
                try:
                    pj = self._read_pj(path)
                    configs.append((path.stat().st_mtime, pj))
                except (OSError, ValueError, SchemaError) as e:
                    logger.warning(f"skipping invalid config {path}: {e}")
                except Exception as e:
                    logger.critical(
                        f"unexpected list config {path}: {e}", exc_info=True
                    )
            configs.sort(key=lambda x: x[0], reverse=True)
            logger.debug(f"listed {len(configs)} project configs")
            return [pj for _, pj in configs]

    def list_named(self) -> list[tuple[str, ProjectJSON]]:
        """List (config name, PJ) pairs sorted by name (020 open dialog)."""
        with self._lock:
            out: list[tuple[str, ProjectJSON]] = []
            for path in self.home_dir.glob("*.stanok"):
                try:
                    out.append((path.stem, self._read_pj(path)))
                except (OSError, ValueError, SchemaError) as e:
                    logger.warning(f"skipping invalid config {path}: {e}")
                except Exception as e:
                    logger.critical(
                        f"unexpected list named config {path}: {e}", exc_info=True
                    )
            out.sort(key=lambda x: x[0].lower())
            logger.debug(f"listed {len(out)} named project configs")
            return out

    def _unique_config_name(self, base: str) -> str:
        """Free config stem: base, base (1), base (2), ... (§5.3 spec.md)."""
        if not self._has_config(self._get_config_path(base)):
            return base
        i = 1
        while True:
            candidate = f"{base} ({i})"
            if not self._has_config(self._get_config_path(candidate)):
                return candidate
            i += 1

    def import_configs(self, folder: str | Path) -> ImportReport:
        """Move every *.stanok under folder into Home; update recent projects."""
        report = ImportReport()
        root = Path(folder).expanduser()
        if not root.is_dir():
            return report
        with self._lock:
            for src in self._config_files(root.resolve()):
                report.found += 1
                try:
                    with open(src, encoding="utf-8") as f:
                        pj = validate_pj(migrate(json.load(f)))
                except (OSError, ValueError, SchemaError) as e:
                    logger.warning(f"import config {src}: {e}", exc_info=True)
                    report.failed.append(src.name)
                    continue
                except Exception as e:
                    logger.critical(
                        f"unexpected import config {src}: {e}", exc_info=True
                    )
                    report.failed.append(src.name)
                    continue
                project_dir = self._import_project_folder(pj, src)
                if project_dir is not None:
                    pj.folder = str(project_dir)
                name = self._unique_config_name(src.stem or "project")
                try:
                    self.save(pj, name=name)
                except (OSError, SchemaError) as e:
                    logger.warning(
                        f"save imported config {src}: {e}", exc_info=True
                    )
                    report.failed.append(src.name)
                    continue
                except Exception as e:
                    logger.critical(
                        f"unexpected save imported config {src}: {e}",
                        exc_info=True,
                    )
                    report.failed.append(src.name)
                    continue
                try:
                    src.unlink()
                except OSError as e:
                    logger.warning(
                        f"remove imported config {src}: {e}", exc_info=True
                    )
                report.imported.append(name)
                if project_dir is None:
                    report.missing_folder.append(name)
                    continue
                try:
                    self.add_recent(str(project_dir), name)
                except (OSError, SchemaError) as e:
                    logger.warning(
                        f"add imported project {name} to recent: {e}",
                        exc_info=True,
                    )
                except Exception as e:
                    logger.critical(
                        f"unexpected add imported project {name} to recent: {e}",
                        exc_info=True,
                    )
        logger.info(
            f"imported configs from {root}: found={report.found} "
            f"imported={len(report.imported)} failed={len(report.failed)} "
            f"missing_folder={len(report.missing_folder)}"
        )
        return report

    def _config_files(self, root: Path) -> list[Path]:
        """All *.stanok under root, except configs already inside Home."""
        files: list[Path] = []
        for path in sorted(root.rglob("*.stanok")):
            if not path.is_file():
                continue
            try:
                path.resolve().relative_to(self.home_dir)
                continue
            except ValueError:
                pass
            files.append(path)
        return files

    def _import_project_folder(
        self, pj: ProjectJSON, src: Path
    ) -> Path | None:
        """Valid PJ.folder, else a nearby folder with Данные and Шаблоны."""
        if pj.folder:
            folder = Path(pj.folder).expanduser()
            if folder.is_dir():
                return folder.resolve()
        for parent in (src.parent, *src.parents):
            if (parent / STRINGS.DATA_DIR).is_dir() and (
                parent / STRINGS.TPL_DIR
            ).is_dir():
                return parent.resolve()
        return None

    def migrate_from_project(self, project_dir: Path) -> ProjectJSON:
        """Migrate project config from project folder to Home."""
        if ".." in project_dir.parts:
            raise StorageError(
                "migrate", [f"path traversal not allowed: {project_dir}"]
            )
        legacy_config = project_dir.resolve() / "project.stanok"
        if not legacy_config.exists():
            raise StorageError(
                "migrate", [f"legacy config not found: {legacy_config}"]
            )
        with open(legacy_config, "r", encoding="utf-8") as f:
            pj = validate_pj(migrate(json.load(f)))
        pj.folder = str(project_dir.resolve())
        name = self._unique_config_name(project_dir.resolve().name)
        self.save(pj, name=name)
        legacy_config.unlink()
        logger.info(f"migrated project config from {project_dir} as {name}")
        return pj

    def resolve_project(self, ref: str | Path) -> tuple[ProjectJSON, Path]:
        """Resolve project reference to (PJ, config_path)."""
        logger.debug(f"resolving project ref: {ref}")
        ref_path = Path(ref) if isinstance(ref, str) else ref

        if ref_path.exists() and ref_path.is_dir():
            legacy = ref_path / "project.stanok"
            if legacy.exists():
                pj = self.migrate_from_project(ref_path)
                return pj, self._get_config_path(ref_path.resolve().name)
            config_path = self._get_config_path(ref_path.name)
            if self._has_config(config_path):
                pj = self.load(ref_path.name)
                self._remember_folder(pj, ref_path, config_path)
                logger.debug(f"resolved project folder {ref_path} -> {config_path}")
                return pj, config_path
            raise StorageError("resolve", [f"project not found: {ref}"])

        config_path = self._get_config_path(ref_path.name)
        if self._has_config(config_path):
            logger.debug(f"resolved project config name {ref_path.name}")
            return self.load(ref_path.name), config_path

        # Config name known only from recent: resolve by remembered folder.
        for item in self.get_recent():
            if item.config == ref_path.name:
                folder = Path(item.folder)
                if folder.is_dir():
                    return self.resolve_project(folder)
        raise StorageError("resolve", [f"project not found: {ref}"])

    def _remember_folder(
        self, pj: ProjectJSON, folder: str | Path, config_path: Path | None = None
    ) -> None:
        """Persist resolved folder into PJ (self-heal / legacy migration)."""
        resolved = str(Path(folder).resolve())
        if pj.folder == resolved:
            return
        pj.folder = resolved
        self.save(pj, name=config_path.stem if config_path else None)
        logger.info(f"remembered project folder {resolved}")

    def set_project_folder(self, ref: str | Path, folder: str | Path) -> None:
        """Point an existing project at a new folder (relocate, 020)."""
        with self._lock:
            pj, config_path = self.resolve_project(ref)
            self._remember_folder(pj, folder, config_path)
            self.add_recent(str(Path(folder).resolve()), config_path.stem)

    def get_recent(self) -> list[RecentItem]:
        """Get recent projects from ApplicationJSON."""
        settings_file = self.home_dir / "settings.json"
        if not settings_file.exists():
            logger.debug("recent projects: settings.json is absent")
            return []
        recent = validate_aj(self._read_json(settings_file)).recent
        logger.debug(f"recent projects: {len(recent)}")
        return recent

    def add_recent(self, folder: str, config: str) -> None:
        """Add or update recent project entry (dedupe, cap 10)."""
        with self._lock:
            settings_file = self.home_dir / "settings.json"
            logger.debug(f"add recent project: folder={folder} config={config}")
            data = self._read_json(settings_file) if settings_file.exists() else {}
            aj = validate_aj(data)
            key = (folder, config)
            aj.recent = [r for r in aj.recent if (r.folder, r.config) != key]
            # In memory newest-first (validator order): new entry goes on top.
            aj.recent.insert(
                0,
                RecentItem(
                    folder=folder, config=config, opened_at=datetime.now()
                ),
            )
            aj.recent = aj.recent[:10]
            self._write_settings(aj)
            logger.info(f"updated recent projects, count={len(aj.recent)}")

    def _write_settings(self, aj: ApplicationJSON) -> None:
        """Write settings.json; `recent` is stored oldest-first on disk."""
        data = aj.model_dump(mode="json")
        data["recent"].reverse()
        self._atomic_write_json(self.home_dir / "settings.json", data)


    def init_project(
        self,
        folder: str | Path,
        name: str,
        xlsx: list[str | Path],
        docx: list[str | Path],
        copy_files: bool = True,
    ) -> ProjectJSON:
        """Create project folder with Данные/Шаблоны + PJ from files (FR-12).

        copy_files=True copies sources into the folder (unique names on
        collision); False only links files already inside Данные/Шаблоны.
        Template fields default to table source named as the placeholder.
        """
        with self._lock:
            name = (name or "").strip()
            if not name or "/" in name or "\\" in name or name in (".", ".."):
                raise StorageError("init", [f"invalid project name: {name!r}"])
            if self.config_exists(name):
                raise StorageError("init", [f"project already exists: {name}"])
            if not xlsx:
                raise StorageError("init", ["no Excel files selected"])
            if not docx:
                raise StorageError("init", ["no Word templates selected"])

            folder = Path(folder).resolve()
            logger.debug(
                f"init project: name={name} folder={folder} "
                f"xlsx={len(xlsx)} docx={len(docx)} copy_files={copy_files}"
            )
            if folder.name in (STRINGS.DATA_DIR, STRINGS.TPL_DIR):
                raise StorageError(
                    "init",
                    [f"cannot create project inside {folder.name}/"],
                )
            data_dir = folder / STRINGS.DATA_DIR
            tpl_dir = folder / STRINGS.TPL_DIR
            result_dir = folder / STRINGS.GEN_RESULT_DIR
            try:
                data_dir.mkdir(parents=True, exist_ok=True)
                tpl_dir.mkdir(parents=True, exist_ok=True)
                result_dir.mkdir(parents=True, exist_ok=True)
            except OSError as e:
                logger.warning(
                    f"cannot create project folder {folder}: {e}", exc_info=True
                )
                raise StorageError(
                    "init", [f"cannot create folder {folder}: {e}"]
                ) from e

            rel_xlsx = [
                self._place_source(Path(f), data_dir, folder, ".xlsx", copy_files)
                for f in xlsx
            ]
            counters: dict[str, dict] = {}
            for r in rel_xlsx:
                cname = _counter_name(set(counters) | self._shared_names(), r)
                counters[cname] = {"pos": 0, "format": "1", "table": r}
            default_counter = next(iter(counters), None)
            templates: dict[str, dict] = {}
            for f in docx:
                rel = self._place_source(
                    Path(f), tpl_dir, folder, ".docx", copy_files
                )
                stem = Path(rel).stem
                base, i = stem, 1
                while stem in templates:
                    i += 1
                    stem = f"{base} ({i})"
                placeholders = _scan_placeholders(folder / rel)
                fields: dict[str, dict] = {}
                for ph in placeholders:
                    fdef: dict = {"source": "table", "value": ph}
                    if default_counter is not None:
                        fdef["counter_name"] = default_counter
                    fields[ph] = fdef
                templates[stem] = {"file": rel, "fields": fields}

            pj = validate_pj(
                {
                    "version": current_version(),
                    "templates": templates,
                    "default_template": next(iter(templates)),
                    "data_sources": [
                        {"file": r, "mode": "sequential"} for r in rel_xlsx
                    ],
                    "counters": counters,
                    "filename_template": name,
                    "folder": str(folder),
                }
            )
            self.save(pj, name=name)
            self.add_recent(str(folder), name)
            logger.info(f"initialized project {name} in {folder}")
            return pj

    def _place_source(
        self,
        src: Path,
        dest_dir: Path,
        folder: Path,
        suffix: str,
        copy_files: bool,
    ) -> str:
        """Copy (or link in place) one source file; return PJ-relative posix path."""
        if src.suffix.lower() != suffix:
            raise StorageError(
                "init", [f"wrong file type {src.name}: expected *{suffix}"]
            )
        if copy_files:
            if not src.is_file():
                raise StorageError("init", [f"file not found: {src}"])
            try:
                rel = src.resolve().relative_to(folder)
            except ValueError:
                rel = None
            if rel is not None and rel.parts and rel.parts[0] == dest_dir.name:
                return rel.as_posix()
            try:
                dst = _unique_sibling(dest_dir / src.name)
                shutil.copy2(src, dst)
            except OSError as e:
                logger.warning(
                    f"copy {src} -> {dest_dir} failed: {e}", exc_info=True
                )
                raise StorageError(
                    "init", [f"cannot copy {src.name}: {e}"]
                ) from e
            return dst.relative_to(folder).as_posix()
        resolved = src.resolve()
        try:
            rel = resolved.relative_to(folder)
        except ValueError:
            raise StorageError(
                "init", [f"file outside project folder: {src}"]
            ) from None
        if rel.parts[0] != dest_dir.name:
            raise StorageError(
                "init", [f"file must be in {dest_dir.name}/: {src}"]
            )
        if not resolved.is_file():
            raise StorageError("init", [f"file not found: {src}"])
        return rel.as_posix()


    def sync_template_fields(self, ref: str | Path) -> dict[str, list[str]]:
        """Add docx placeholders missing from PJ (020); return template → added.

        Only adds; never removes (user mappings are kept). Unreadable or
        missing template files are skipped with a warning.
        """
        with self._lock:
            pj, config_path = self.resolve_project(ref)
            try:
                folder = project_folder(ref, pj, config_path, self)
            except StorageError as e:
                logger.warning(f"sync skipped, no folder: {e}")
                return {}
            added: dict[str, list[str]] = {}
            default_counter = _first_bound_counter(pj)
            for name, tpl in pj.templates.items():
                tpl_path = folder / tpl.file
                if not tpl_path.is_file():
                    logger.warning(f"sync skipped, missing template file: {tpl.file}")
                    continue
                try:
                    placeholders = _scan_placeholders(tpl_path)
                except StorageError as e:
                    logger.warning(f"sync skipped, unreadable template: {e}")
                    continue
                new = [ph for ph in placeholders if ph not in tpl.fields]
                if new and default_counter is None:
                    logger.warning(
                        f"sync skipped, no bound counter for new fields: {tpl.file}"
                    )
                    continue
                for ph in new:
                    tpl.fields[ph] = FieldDef(
                        source=FieldSource.TABLE,
                        value=ph,
                        counter_name=default_counter,
                    )
                if new:
                    added[name] = new
            if added:
                self.save(pj, name=config_path.stem)
                logger.info(f"synced template fields: {added}")
            return added

    def touch_recent(self, ref: str | Path) -> None:
        """Move a project to the top of recent (folder resolved from PJ)."""
        with self._lock:
            pj, config_path = self.resolve_project(ref)
            try:
                folder = project_folder(ref, pj, config_path, self)
            except StorageError as e:
                logger.warning(f"touch_recent {ref}: {e}")
                return
            self.add_recent(str(folder), config_path.stem)

    def remove_recent(self, folder: str, config: str) -> None:
        """Remove one recent project entry (no-op when absent)."""
        with self._lock:
            settings_file = self.home_dir / "settings.json"
            data = self._read_json(settings_file) if settings_file.exists() else {}
            aj = validate_aj(data)
            aj.recent = [
                r
                for r in aj.recent
                if (r.folder, r.config) != (folder, config)
            ]
            self._write_settings(aj)
            logger.info(f"removed recent project {config}")

    def get_setting(self, key: str, default=None):
        """Read one AJ setting (018)."""
        settings_file = self.home_dir / "settings.json"
        if not settings_file.exists():
            return default
        logger.debug(f"reading setting {key} from {settings_file}")
        try:
            data = self._read_json(settings_file)
            value = validate_aj(data).settings.get(key, default)
            logger.debug(f"read setting {key}: {value!r}")
            return value
        except (OSError, ValueError, SchemaError) as e:
            logger.warning(f"read setting {key} failed: {e}", exc_info=True)
            return default
        except Exception as e:
            logger.critical(
                f"unexpected read setting {key}: {e}", exc_info=True
            )
            return default

    def set_setting(self, key: str, value) -> None:
        """Write one AJ setting atomically (018)."""
        with self._lock:
            settings_file = self.home_dir / "settings.json"
            logger.debug(f"setting {key}={value!r} in {settings_file}")
            data = self._read_json(settings_file) if settings_file.exists() else {}
            aj = validate_aj(data)
            aj.settings[key] = value
            self._write_settings(aj)
            logger.info(f"updated setting {key}")

    # -- shared (сквозные) counters, 026 ------------------------------------
    def global_counters(self) -> dict[str, CounterDef]:
        """Shared numeric counters from settings.json (026)."""
        with self._lock:
            settings_file = self.home_dir / "settings.json"
            if not settings_file.exists():
                return {}
            return dict(validate_aj(self._read_json(settings_file)).counters)

    def _project_counter_names(self) -> set[str]:
        names: set[str] = set()
        for _cfg, pj in self.list_named():
            names.update(pj.counters)
        return names

    def global_counter_users(self, name: str) -> list[str]:
        """Config names of projects whose counter fields read the shared counter."""
        with self._lock:
            return [
                cfg
                for cfg, pj in self.list_named()
                if any(
                    fdef.source == FieldSource.COUNTER and fdef.value == name
                    for tdef in pj.templates.values()
                    for fdef in tdef.fields.values()
                )
            ]

    def _shared_names(self) -> set[str]:
        """Names of shared counters: project counters must not reuse them (026)."""
        return set(self.global_counters())

    def _settings_aj(self) -> ApplicationJSON:
        settings_file = self.home_dir / "settings.json"
        data = self._read_json(settings_file) if settings_file.exists() else {}
        return validate_aj(data)

    def add_global_counter(self, name: str, fmt: str = "1", pos: int = 0) -> None:
        """Create a shared counter (026); its name must be free in settings and projects."""
        with self._lock:
            name = (name or "").strip()
            if not name:
                raise StorageError("counters", [STRINGS.GCNT_NAME_EMPTY])
            if "/" in name or "\\" in name:
                raise StorageError("counters", [STRINGS.GCNT_NAME_BAD])
            if fmt not in COUNTER_WIDTHS:
                raise StorageError("counters", [STRINGS.GCNT_FORMAT_BAD.format(fmt=fmt)])
            aj = self._settings_aj()
            if name in aj.counters or name in self._project_counter_names():
                raise StorageError("counters", [STRINGS.GCNT_NAME_TAKEN.format(name=name)])
            aj.counters[name] = CounterDef(pos=max(0, int(pos)), format=fmt)
            self._write_settings(aj)
            logger.info(f"added shared counter {name} (format {fmt})")

    def set_global_counter_pos(self, name: str, pos: int) -> None:
        """Set the current value of a shared counter (also used by generation)."""
        with self._lock:
            aj = self._settings_aj()
            if name not in aj.counters:
                raise StorageError("counters", [STRINGS.GCNT_UNKNOWN.format(name=name)])
            aj.counters[name].pos = max(0, int(pos))
            self._write_settings(aj)
            logger.info(f"shared counter {name} -> {pos}")

    def rename_global_counter(self, old: str, new: str) -> None:
        """Rename a shared counter and point every project field that reads it (026)."""
        with self._lock:
            aj = self._settings_aj()
            if old not in aj.counters:
                raise StorageError("counters", [STRINGS.GCNT_UNKNOWN.format(name=old)])
            new = (new or "").strip()
            if not new:
                raise StorageError("counters", [STRINGS.GCNT_NAME_EMPTY])
            if "/" in new or "\\" in new:
                raise StorageError("counters", [STRINGS.GCNT_NAME_BAD])
            if new == old:
                return
            if new in aj.counters or new in self._project_counter_names():
                raise StorageError("counters", [STRINGS.GCNT_NAME_TAKEN.format(name=new)])
            for cfg, pj in self.list_named():
                changed = False
                for tdef in pj.templates.values():
                    for fdef in tdef.fields.values():
                        if fdef.source == FieldSource.COUNTER and fdef.value == old:
                            fdef.value = new
                            changed = True
                if changed:
                    self.save(pj, name=cfg)
            aj.counters[new] = aj.counters.pop(old)
            self._write_settings(aj)
            logger.info(f"renamed shared counter {old} -> {new}")

    def remove_global_counter(self, name: str) -> None:
        """Delete a shared counter; refused while any project's field reads it."""
        with self._lock:
            users = self.global_counter_users(name)
            if users:
                raise StorageError(
                    "counters",
                    [STRINGS.GCNT_IN_USE.format(name=name, projects=", ".join(users))],
                )
            aj = self._settings_aj()
            if name not in aj.counters:
                raise StorageError("counters", [STRINGS.GCNT_UNKNOWN.format(name=name)])
            del aj.counters[name]
            self._write_settings(aj)
            logger.info(f"removed shared counter {name}")

    def add_template(
        self, ref: str | Path, docx: str | Path, copy_files: bool = True
    ) -> str:
        """Attach a docx template to the project (FR-13); return its name."""
        with self._lock:
            pj, config_path = self.resolve_project(ref)
            folder = project_folder(ref, pj, config_path, self)
            tpl_dir = folder / STRINGS.TPL_DIR
            try:
                tpl_dir.mkdir(parents=True, exist_ok=True)
            except OSError as e:
                logger.warning(f"cannot create templates dir {tpl_dir}: {e}")
                raise StorageError(
                    "templates", [f"cannot create folder {tpl_dir}: {e}"]
                ) from e
            rel = self._place_source(
                Path(docx), tpl_dir, folder, ".docx", copy_files
            )
            stem = Path(rel).stem
            base, i = stem, 1
            while stem in pj.templates:
                i += 1
                stem = f"{base} ({i})"
            placeholders = _scan_placeholders(folder / rel)
            default_counter = _first_bound_counter(pj)
            if placeholders and default_counter is None:
                raise StorageError(
                    "templates",
                    ["cannot add template: no data source to bind fields"],
                )
            pj.templates[stem] = TemplateDef(
                file=rel,
                fields={
                    ph: FieldDef(
                        source=FieldSource.TABLE,
                        value=ph,
                        counter_name=default_counter,
                    )
                    for ph in placeholders
                },
            )
            self.save(pj, name=config_path.stem)
            logger.info(f"attached template {stem} to {config_path.stem}")
            return stem

    def remove_template(self, ref: str | Path, name: str) -> None:
        """Detach a template; removing the last one is forbidden (FR-13)."""
        with self._lock:
            pj, config_path = self.resolve_project(ref)
            if name not in pj.templates:
                raise StorageError(
                    "templates", [f"template not found: {name}"]
                )
            if len(pj.templates) == 1:
                raise StorageError(
                    "templates",
                    ["cannot remove last template: project needs at least one"],
                )
            if name == pj.default_template:
                raise StorageError("templates", [STRINGS.STORAGE_DEFAULT_TPL])
            del pj.templates[name]
            self.save(pj, name=config_path.stem)
            logger.info(f"removed template {name} from {config_path.stem}")

    def set_default_template(self, ref: str | Path, name: str) -> None:
        """Choose the project's default template (021); used by window 1."""
        with self._lock:
            pj, config_path = self.resolve_project(ref)
            if name not in pj.templates:
                raise StorageError("templates", [f"template not found: {name}"])
            pj.default_template = name
            self.save(pj, name=config_path.stem)
            logger.info(f"default template of {config_path.stem} -> {name}")

    def set_result_dirs(self, ref: str | Path, dirs: list[str]) -> None:
        """Replace the extra result folders (025); each must be an absolute path."""
        with self._lock:
            pj, config_path = self.resolve_project(ref)
            pj.result_dirs = list(dirs)
            pj = validate_pj(pj.model_dump(mode="json"))
            self.save(pj, name=config_path.stem)
            logger.info(f"result copy folders of {config_path.stem}: {len(pj.result_dirs)}")

    def set_template_limit(self, ref: str | Path, name: str, limit: int) -> None:
        """Persist a per-template document limit (021); 0 = no limit."""
        with self._lock:
            pj, config_path = self.resolve_project(ref)
            if name not in pj.templates:
                raise StorageError("templates", [f"template not found: {name}"])
            pj.templates[name].limit = max(0, int(limit))
            self.save(pj, name=config_path.stem)
            logger.info(f"template {name} limit of {config_path.stem} -> {limit}")

    def set_source_mode(self, ref: str | Path, file: str, mode: str) -> None:
        """Persist one data source's mode (021): sequential/circular/constant."""
        modes = ("sequential", "circular", "constant")
        if mode not in modes:
            raise StorageError("data", [f"unknown mode: {mode}"])
        with self._lock:
            pj, config_path = self.resolve_project(ref)
            for ds in pj.data_sources:
                if ds.file == file:
                    ds.mode = mode
                    self.save(pj, name=config_path.stem)
                    logger.info(f"source {file} mode of {config_path.stem} -> {mode}")
                    return
            raise StorageError("data", [f"data source not found: {file}"])

    def add_source(
        self, ref: str | Path, xlsx: str | Path, copy_files: bool = True
    ) -> str:
        """Attach an xlsx data source (sequential); return its PJ path."""
        with self._lock:
            pj, config_path = self.resolve_project(ref)
            folder = project_folder(ref, pj, config_path, self)
            data_dir = folder / STRINGS.DATA_DIR
            try:
                data_dir.mkdir(parents=True, exist_ok=True)
            except OSError as e:
                logger.warning(f"cannot create data dir {data_dir}: {e}")
                raise StorageError(
                    "data", [f"cannot create folder {data_dir}: {e}"]
                ) from e
            rel = self._place_source(
                Path(xlsx), data_dir, folder, ".xlsx", copy_files
            )
            if any(ds.file == rel for ds in pj.data_sources):
                raise StorageError(
                    "data", [f"data source already attached: {rel}"]
                )
            pj.data_sources.append(DataSourceDef(file=rel))
            cname = _counter_name(set(pj.counters) | self._shared_names(), rel)
            pj.counters[cname] = CounterDef(table=rel)
            self.save(pj, name=config_path.stem)
            logger.info(
                f"attached data source {rel} (counter {cname}) to "
                f"{config_path.stem}"
            )
            return rel

    def remove_source(self, ref: str | Path, file: str) -> None:
        """Detach a data source and its counters (025). The xlsx file stays on disk.

        Refused for the last source, and while a template field still reads one
        of its counters: the user changes those fields first.
        """
        with self._lock:
            pj, config_path = self.resolve_project(ref)
            if not any(ds.file == file for ds in pj.data_sources):
                raise StorageError("data", [f"data source not found: {file}"])
            if len(pj.data_sources) == 1:
                raise StorageError("data", [STRINGS.STORAGE_SOURCE_LAST])
            counters = {name for name, c in pj.counters.items() if c.table == file}
            users = [
                f"{tname}: {fname}"
                for tname, tdef in pj.templates.items()
                for fname, fdef in tdef.fields.items()
                if (fdef.source == FieldSource.TABLE and fdef.counter_name in counters)
                or (fdef.source == FieldSource.COUNTER and fdef.value in counters)
            ]
            if users:
                raise StorageError(
                    "data",
                    [STRINGS.STORAGE_SOURCE_IN_USE.format(file=file, fields=", ".join(users))],
                )
            pj.data_sources = [ds for ds in pj.data_sources if ds.file != file]
            for name in counters:
                del pj.counters[name]
            self.save(pj, name=config_path.stem)
            logger.info(f"removed data source {file} from {config_path.stem}")

    @staticmethod
    def _remove_project_folder(folder: Path) -> None:
        """Remove the project's own dirs; the whole folder only if nothing else is in it.

        Own dirs: Данные, Шаблоны, Результат. Any other file or folder means
        the user keeps their data, so only the own dirs are removed.
        """
        own = (STRINGS.DATA_DIR, STRINGS.TPL_DIR, STRINGS.GEN_RESULT_DIR)
        extra = [
            p for p in folder.iterdir() if not (p.name in own and p.is_dir())
        ]
        if not extra:
            shutil.rmtree(folder)
            return
        logger.info(f"project folder {folder} has other content: own dirs only")
        for name in own:
            target = folder / name
            if target.is_dir():
                shutil.rmtree(target)

    def delete_project(self, ref: str | Path, delete_folder: bool = False) -> None:
        """Delete project config + recent entry; folder only on flag (FR-13)."""
        with self._lock:
            pj, config_path = self.resolve_project(ref)
            try:
                folder = project_folder(ref, pj, config_path, self)
            except StorageError:
                folder = None
            if delete_folder and folder == self.home_dir:
                raise StorageError(
                    "delete", ["refusing to delete home directory"]
                )
            # Remove the folder first: a locked file must not leave the user
            # with a deleted config and a half-deleted project.
            if delete_folder and folder is not None and folder.exists():
                try:
                    self._remove_project_folder(folder)
                except OSError as e:
                    logger.warning(f"rmtree {folder} failed: {e}")
                    raise StorageError(
                        "delete", [f"folder not removed, project kept: {e}"]
                    ) from e
            self.delete(config_path.stem)
            if folder is not None:
                self.remove_recent(str(folder), config_path.stem)
            logger.info(f"deleted project {config_path.stem}")


def resolve_project(ref: str | Path, store: ProjectStore) -> tuple:
    """Convenience function to resolve project reference."""
    return store.resolve_project(ref)
