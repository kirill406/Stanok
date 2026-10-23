# Copyright (C) 2026 Kirill Borovoy
# SPDX-License-Identifier: GPL-3.0-or-later
"""Excel table reading implementation."""

import logging
from datetime import date, datetime
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

import openpyxl
from openpyxl.utils import get_column_letter

from ..gui.strings import STRINGS

logger = logging.getLogger(__name__)


class TableReadError(Exception):
    def __init__(self, path: Path, cause: Exception):
        self.path = path
        self.cause = cause
        super().__init__(STRINGS.TBL_READ_FAILED.format(path=path, cause=cause))


class ExcelReader:
    def __init__(self) -> None:
        self.skipped_columns: dict[Path, list[str]] = {}

    def read(self, path: Path) -> list[dict[str, Any]]:
        self.skipped_columns.pop(path, None)
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            try:
                ws = wb.active
                rows = ws.iter_rows(values_only=True)
                headers, indices, skipped = self._read_headers(path, next(rows, []))
                if skipped:
                    self.skipped_columns[path] = skipped
                    logger.warning(
                        f"table {path}: skipped columns without headers: "
                        f"{', '.join(skipped)}"
                    )
                result = []
                for row in rows:
                    cells = list(row)
                    values = [
                        cells[i] if i < len(cells) else None for i in indices
                    ]
                    if all(v is None or v == "" for v in values):
                        # Пустая строка — dict с None значениями
                        result.append({h: None for h in headers})
                        continue
                    result.append(
                        {h: self._normalize(v) for h, v in zip(headers, values)}
                    )
                return result
            finally:
                wb.close()
        except TableReadError:
            raise
        except (OSError, ValueError, KeyError, BadZipFile) as e:
            logger.error(f"read table {path}: {e}", exc_info=True)
            raise TableReadError(path, e) from e
        except Exception as e:
            logger.critical(f"unexpected read table {path}: {e}", exc_info=True)
            raise TableReadError(path, e) from e

    @staticmethod
    def _read_headers(
        path: Path, raw_row: tuple
    ) -> tuple[list[str], list[int], list[str]]:
        raw = [str(h).strip() if h is not None else "" for h in raw_row]
        # Хвостовые пустые колонки — обрезать (обычное дело в живых файлах).
        last = max((i for i, h in enumerate(raw) if h), default=-1)
        if last < 0:
            raise TableReadError(path, ValueError(STRINGS.TBL_NO_HEADERS))
        headers: list[str] = []
        indices: list[int] = []
        skipped: list[str] = []
        for i, h in enumerate(raw[: last + 1]):
            if not h:
                skipped.append(get_column_letter(i + 1))
                continue
            headers.append(h)
            indices.append(i)
        # Дубли — тихая потеря данных, запрещены.
        seen, dups = set(), set()
        for h in headers:
            if h in seen:
                dups.add(h)
            seen.add(h)
        if dups:
            raise TableReadError(
                path,
                ValueError(STRINGS.TBL_DUP_HEADERS.format(headers=", ".join(sorted(dups)))),
            )
        return headers, indices, skipped

    def _normalize(self, v: Any) -> Any:
        if v is None:
            return None
        if isinstance(v, bool):
            return v
        if isinstance(v, int):
            return v
        if isinstance(v, float):
            return int(v) if v.is_integer() else v
        if isinstance(v, (datetime, date)):
            return v
        s = str(v).strip()
        return s if s else None
