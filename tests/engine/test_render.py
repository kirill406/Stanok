# Copyright (C) 2026 Kirill Borovoy
# SPDX-License-Identifier: GPL-3.0-or-later
"""Tests for render: FillingJSON + template → Document."""

from pathlib import Path

import pytest
from docx import Document

from stanok.engine.render import render
from stanok.engine.schema import FillingJSON, RenderError, validate_fj


FIXTURE_DIR = Path(__file__).parent.parent / "json" / "004-render"


def _fj(**fields) -> FillingJSON:
    return validate_fj(
        {"version": "0.0.0", "template": "Д", "fields": fields, "dist": "Д.docx"}
    )


def _texts(doc: Document) -> list[str]:
    out = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                out.extend(p.text for p in cell.paragraphs)
    return out


def test_basic_substitution_keeps_bold():
    doc = render(
        _fj(Номер="43", Дата="2026-10-05", ФИО="Иванов"),
        FIXTURE_DIR / "template_basic.docx",
    )
    assert doc.paragraphs[0].text == "Договор № 43 от 2026-10-05"
    bold_runs = [r.text for r in doc.paragraphs[0].runs if r.bold]
    assert "43" in "".join(bold_runs)


def test_split_runs_merged():
    doc = render(
        _fj(ФИО="Иванов", Город="Москва"), FIXTURE_DIR / "template_split.docx"
    )
    assert doc.paragraphs[0].text == "Иванов"
    assert "{{" not in doc.paragraphs[0].text
    assert doc.paragraphs[1].text == "Город: Москва"


def test_newline_becomes_break():
    doc = render(_fj(ФИО="Ива\nов", Город="X"), FIXTURE_DIR / "template_split.docx")
    xml = doc.paragraphs[0]._p.xml
    assert "<w:br" in xml
    assert "{{" not in doc.paragraphs[0].text


def test_table_cells_substituted():
    doc = render(
        _fj(ФИО="Иванов", Возраст=30, Номер="1", Итого="done"),
        FIXTURE_DIR / "template_table.docx",
    )
    texts = _texts(doc)
    assert "Иванов" in texts and "30" in texts and "done" in texts
    assert not any("{{" in t for t in texts)


def test_unknown_field_raises(tmp_path):
    with pytest.raises(RenderError, match="fields.Нетакого"):
        render(_fj(Номер="1"), _make_template(tmp_path, "Привет {{Нетакого}}"))


def test_broken_template_raises(tmp_path):
    bad = tmp_path / "broken.docx"
    bad.write_bytes(b"not a zip")
    with pytest.raises(RenderError, match="template"):
        render(_fj(Номер="1"), bad)


def test_missing_template_raises(tmp_path):
    with pytest.raises(RenderError, match="template"):
        render(_fj(Номер="1"), tmp_path / "nope.docx")


def test_unclosed_placeholder_raises(tmp_path):
    with pytest.raises(RenderError, match="template"):
        render(_fj(Номер="1"), _make_template(tmp_path, "Привет {{Номер"))


def _make_template(tmp_path: Path, text: str) -> Path:
    path = tmp_path / "t.docx"
    doc = Document()
    doc.add_paragraph(text)
    doc.save(path)
    return path


def test_none_and_bool_values(tmp_path):
    doc = render(
        _fj(А=None, Б=True, В=False), _make_template(tmp_path, "{{А}}|{{Б}}|{{В}}")
    )
    assert doc.paragraphs[0].text == "|True|False"


def test_multi_line_value(tmp_path):
    doc = render(_fj(А="x\ny\nz"), _make_template(tmp_path, "[{{А}}]"))
    assert doc.paragraphs[0]._p.xml.count("<w:br") == 2
    assert doc.paragraphs[0].text.replace("\n", "") == "[xyz]"


def test_empty_paragraph_untouched(tmp_path):
    path = tmp_path / "e.docx"
    doc = Document()
    doc.add_paragraph("")
    doc.add_paragraph("{{А}}")
    doc.save(path)
    out = render(_fj(А="1"), path)
    assert out.paragraphs[0].text == ""
    assert out.paragraphs[1].text == "1"


def test_datetime_value_isoformat(tmp_path):
    from datetime import datetime

    doc = render(
        _fj(А=datetime(2026, 10, 5, 14, 30)),
        _make_template(tmp_path, "{{А}}"),
    )
    assert doc.paragraphs[0].text == "2026-10-05T14:30:00"


def test_nested_table_substituted(tmp_path):
    path = tmp_path / "nested.docx"
    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    inner = outer.cell(0, 0).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "{{А}}"
    doc.save(path)
    out = render(_fj(А="глубоко"), path)
    assert inner.cell(0, 0).text == "{{А}}"  # исходный шаблон не тронут
    out_inner = out.tables[0].cell(0, 0).tables[0]
    assert out_inner.cell(0, 0).text == "глубоко"


def test_unexpected_error_wrapped(monkeypatch, tmp_path):
    import stanok.engine.render as render_mod

    def boom(*a, **k):
        raise RuntimeError("boom")

    monkeypatch.setattr(render_mod, "substitute_paragraph", boom)
    with pytest.raises(RenderError, match="template"):
        render(_fj(А="1"), _make_template(tmp_path, "{{А}}"))


def _header_docx(path, header_text="Договор {{Номер}}", footer_text="Стр. {{Номер}}"):
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_paragraph("Тело {{Имя}}")
    doc.sections[0].header.paragraphs[0].text = header_text
    doc.sections[0].footer.paragraphs[0].text = footer_text
    doc.save(path)
    return path


def test_header_footer_substituted(tmp_path):
    from stanok.engine.render import render
    from stanok.engine.schema import validate_fj

    path = _header_docx(tmp_path / "hdr.docx")
    fj = validate_fj({
        "version": "0.0.0",
        "template": "Д",
        "fields": {"Имя": "Иван", "Номер": "7"},
        "dist": "Д_1.docx",
    })
    doc = render(fj, path)
    assert doc.sections[0].header.paragraphs[0].text == "Договор 7"
    assert doc.sections[0].footer.paragraphs[0].text == "Стр. 7"
    assert doc.paragraphs[0].text == "Тело Иван"


def test_first_page_header_substituted(tmp_path):
    from docx import Document as DocxDocument

    from stanok.engine.render import render
    from stanok.engine.schema import validate_fj

    doc = DocxDocument()
    doc.add_paragraph("Тело")
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    section.first_page_header.paragraphs[0].text = "Первый {{Номер}}"
    path = tmp_path / "fp.docx"
    doc.save(path)
    fj = validate_fj({
        "version": "0.0.0", "template": "Д", "fields": {"Номер": "7"},
        "dist": "Д_1.docx",
    })
    out = render(fj, path)
    assert out.sections[0].first_page_header.paragraphs[0].text == "Первый 7"


def test_even_page_and_linked_sections(tmp_path):
    from docx import Document as DocxDocument
    from docx.enum.section import WD_SECTION

    from stanok.engine.render import render
    from stanok.engine.schema import validate_fj

    doc = DocxDocument()
    doc.add_paragraph("Тело")
    doc.settings.odd_and_even_pages_header_footer = True
    section = doc.sections[0]
    section.even_page_header.paragraphs[0].text = "Чёт {{Номер}}"
    section.header.paragraphs[0].text = "Шапка {{Номер}}"
    second = doc.add_section(WD_SECTION.NEW_PAGE)
    assert second.header.is_linked_to_previous
    path = tmp_path / "ev.docx"
    doc.save(path)
    fj = validate_fj({
        "version": "0.0.0", "template": "Д", "fields": {"Номер": "7"},
        "dist": "Д_1.docx",
    })
    out = render(fj, path)
    assert out.sections[0].even_page_header.paragraphs[0].text == "Чёт 7"
    assert out.sections[0].header.paragraphs[0].text == "Шапка 7"


def test_header_table_substituted(tmp_path):
    from docx import Document as DocxDocument
    from docx.shared import Mm

    from stanok.engine.render import render
    from stanok.engine.schema import validate_fj

    doc = DocxDocument()
    doc.add_paragraph("Тело")
    table = doc.sections[0].header.add_table(rows=1, cols=1, width=Mm(100))
    table.cell(0, 0).text = "Ячейка {{Номер}}"
    path = tmp_path / "ht.docx"
    doc.save(path)
    fj = validate_fj({
        "version": "0.0.0", "template": "Д", "fields": {"Номер": "7"},
        "dist": "Д_1.docx",
    })
    out = render(fj, path)
    assert out.sections[0].header.tables[0].cell(0, 0).text == "Ячейка 7"


def test_iter_paragraphs_no_headers_flag(tmp_path):
    from stanok.engine.xmlops import iter_paragraphs

    path = _header_docx(tmp_path / "nh.docx")
    from docx import Document as DocxDocument

    texts = [p.text for p in iter_paragraphs(DocxDocument(str(path)), headers=False)]
    assert texts == ["Тело {{Имя}}"]
