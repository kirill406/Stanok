# Copyright (C) 2026 Kirill Borovoy
# SPDX-License-Identifier: GPL-3.0-or-later
"""Offscreen tests for window 3 (FieldsDialog, 027): the template page."""

import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import json

import pytest
from PyQt5.QtCore import Qt
from PyQt5.QtWidgets import QApplication, QDialog, QMessageBox, QPushButton

import stanok.gui.field_editor as fe
import stanok.gui.fields_dialog as fd
from stanok.engine.schema import FieldDef, FieldSource
from stanok.gui.fields_dialog import FieldsDialog
from stanok.gui.strings import STRINGS
from stanok.services.generate import TemplateError
from stanok.services.storage import ProjectStore, StorageError
from tests.services.test_generate import make_project


@pytest.fixture(scope="session")
def qapp():
    app = QApplication.instance() or QApplication([])
    yield app


@pytest.fixture
def store(tmp_path):
    return ProjectStore(tmp_path / ".stanok")


@pytest.fixture
def messages(monkeypatch):
    """Capture warnings, information and critical boxes."""
    shown: list[str] = []
    for name in ("warning", "information", "critical"):
        monkeypatch.setattr(QMessageBox, name, lambda *a, _s=shown: _s.append(a[-1]))
    return shown


def _window(tmp_path, store):
    folder = make_project(tmp_path / "proj")
    store.migrate_from_project(folder)
    return FieldsDialog(str(folder), "Договор", store=store), folder


def _edit_field(monkeypatch, change, accept=True):
    """Stand in for the field editor: apply `change` to it, then press OK or Cancel."""

    def fake_exec(self):
        change(self)
        if accept:
            self._on_ok()
        return self.result()

    monkeypatch.setattr(fe.FieldEditDialog, "exec_", fake_exec)


def _set_constant(editor):
    editor.type_combo.setCurrentIndex(editor.type_combo.findData("constant"))
    editor._const_edit.setText("Москва")


def _unused_buttons(dlg):
    return [b.text() for b in dlg._unused.findChildren(QPushButton, "fieldButton")]


def test_fields_dialog_opens_maximized(qapp, tmp_path, store):
    dlg, _ = _window(tmp_path, store)

    assert dlg.windowState() & Qt.WindowMaximized
    dlg.close()


def test_fields_dialog_page_has_button_for_every_field_in_document(qapp, tmp_path, store):
    dlg, _ = _window(tmp_path, store)

    assert dlg._sheet.field_names == ["ФИО", "Сумма", "Номер"]
    dlg.close()


def test_fields_dialog_has_no_fields_or_counters_table(qapp, tmp_path, store):
    dlg, _ = _window(tmp_path, store)

    assert not hasattr(dlg, "table")
    assert not hasattr(dlg, "ctable")
    assert not hasattr(dlg, "counter_add_btn")
    dlg.close()


def test_fields_dialog_field_button_tooltip_shows_source(qapp, tmp_path, store):
    dlg, _ = _window(tmp_path, store)

    tip = dlg._sheet._buttons["ФИО"][0].toolTip()

    assert tip == STRINGS.FLD_DESC_TABLE.format(column="ФИО", counter="data")
    dlg.close()


def test_fields_dialog_field_missing_from_project_is_unset(qapp, tmp_path, store):
    dlg, _ = _window(tmp_path, store)
    del dlg._template_def.fields["Сумма"]
    dlg._sheet.refresh()

    button = dlg._sheet._buttons["Сумма"][0]
    assert button.property("state") == "unset"
    assert button.toolTip() == STRINGS.FLD_DESC_UNSET
    dlg.close()


def test_fields_dialog_unused_field_gets_its_own_button(qapp, tmp_path, store):
    dlg, _ = _window(tmp_path, store)
    dlg._template_def.fields["Лишнее"] = FieldDef(source=FieldSource.CONSTANT, value="x")

    dlg._refresh_unused()

    assert dlg._unused_names() == ["Лишнее"]
    assert _unused_buttons(dlg) == ["Лишнее"]
    dlg.close()


def test_fields_dialog_click_applies_edit_and_marks_dirty(qapp, tmp_path, store, monkeypatch):
    dlg, _ = _window(tmp_path, store)
    _edit_field(monkeypatch, _set_constant)

    dlg._open_field("ФИО")

    assert dlg._dirty is True
    assert not dlg.save_btn.isHidden()
    assert dlg._template_def.fields["ФИО"].source == FieldSource.CONSTANT
    assert dlg._sheet._buttons["ФИО"][0].toolTip() == STRINGS.FLD_DESC_CONSTANT.format(value="Москва")
    dlg.close()


def test_fields_dialog_cancelled_edit_is_not_dirty(qapp, tmp_path, store, monkeypatch):
    dlg, _ = _window(tmp_path, store)
    _edit_field(monkeypatch, _set_constant, accept=False)

    dlg._open_field("ФИО")

    assert dlg._dirty is False
    assert dlg._template_def.fields["ФИО"].source == FieldSource.TABLE
    dlg.close()


def test_fields_dialog_save_persists_edit(qapp, tmp_path, store, monkeypatch):
    dlg, _ = _window(tmp_path, store)
    _edit_field(monkeypatch, _set_constant)
    dlg._open_field("ФИО")

    dlg._on_save()

    assert dlg._dirty is False
    assert dlg.save_btn.isHidden()
    assert store.load("proj").templates["Договор"].fields["ФИО"].value == "Москва"
    dlg.close()


def test_fields_dialog_save_blocks_empty_filename(qapp, tmp_path, store, messages):
    dlg, _ = _window(tmp_path, store)
    dlg.fname_edit.setText("")

    dlg._on_save()

    assert messages == [STRINGS.FLD_EMPTY_VALUE.format(name="Договор")]
    assert dlg._dirty is True
    dlg.close()


def test_fields_dialog_filename_change_marks_dirty(qapp, tmp_path, store):
    dlg, _ = _window(tmp_path, store)
    assert dlg._dirty is False

    dlg.fname_edit.setText("Новое_{i}.docx")

    assert dlg._dirty is True
    dlg.close()


def test_fields_dialog_saved_project_keeps_filename_template(qapp, tmp_path, store):
    dlg, _ = _window(tmp_path, store)
    dlg.fname_edit.setText("Док_{ФИО}.docx")

    dlg._on_save()

    saved = json.loads((store.home_dir / "proj.stanok").read_text(encoding="utf-8"))
    assert saved["filename_template"] == "Док_{ФИО}.docx"
    dlg.close()


def test_fields_dialog_close_dirty_discard_closes(qapp, tmp_path, store, monkeypatch):
    dlg, _ = _window(tmp_path, store)
    dlg._mark_dirty()
    monkeypatch.setattr(fd, "ask_save_discard_cancel", lambda *a: QMessageBox.Discard)

    dlg.reject()

    assert dlg.result() == QDialog.Rejected
    dlg.close()


def test_fields_dialog_close_dirty_cancel_stays_open(qapp, tmp_path, store, monkeypatch):
    dlg, _ = _window(tmp_path, store)
    dlg._mark_dirty()
    monkeypatch.setattr(fd, "ask_save_discard_cancel", lambda *a: QMessageBox.Cancel)

    dlg.reject()

    assert dlg.result() == 0
    dlg.close()


def test_fields_dialog_close_dirty_save_writes_then_closes(qapp, tmp_path, store, monkeypatch):
    dlg, _ = _window(tmp_path, store)
    _edit_field(monkeypatch, _set_constant)
    dlg._open_field("ФИО")
    monkeypatch.setattr(fd, "ask_save_discard_cancel", lambda *a: QMessageBox.Save)

    dlg.reject()

    assert store.load("proj").templates["Договор"].fields["ФИО"].value == "Москва"
    assert dlg.result() == QDialog.Accepted
    dlg.close()


def test_fields_dialog_save_failure_stays_open(qapp, tmp_path, store, monkeypatch, messages):
    dlg, _ = _window(tmp_path, store)
    dlg._mark_dirty()

    def boom(*a, **k):
        raise RuntimeError("disk full")

    monkeypatch.setattr(dlg._store, "save", boom)
    monkeypatch.setattr(fd, "ask_save_discard_cancel", lambda *a: QMessageBox.Save)

    dlg.reject()

    assert messages == ["disk full"]
    assert dlg.result() == 0
    dlg.close()


def test_fields_dialog_missing_template_file_warns_and_lists_fields(
    qapp, tmp_path, store, messages
):
    folder = make_project(tmp_path / "proj")
    store.migrate_from_project(folder)
    (folder / "Шаблоны" / "tpl.docx").unlink()

    dlg = FieldsDialog(str(folder), "Договор", store=store)

    assert dlg._sheet is None
    assert messages and "файл шаблона не найден" in messages[0]
    assert dlg._unused_names() == ["ФИО", "Сумма", "Номер"]
    dlg.close()


def test_fields_dialog_missing_project_raises(qapp, tmp_path, store):
    with pytest.raises(StorageError):
        FieldsDialog("nope", "Договор", store=store)


def test_fields_dialog_missing_template_raises(qapp, tmp_path, store):
    folder = make_project(tmp_path / "proj")
    with pytest.raises(TemplateError, match="template not found"):
        FieldsDialog(str(folder), "Чужой", store=store)


def test_fields_dialog_insert_menu_offers_page_and_saved_fields(qapp, tmp_path, store):
    dlg, _ = _window(tmp_path, store)
    dlg._template_def.fields["Лишнее"] = FieldDef(source=FieldSource.CONSTANT, value="x")

    assert dlg._field_names() == ["ФИО", "Сумма", "Номер", "Лишнее"]
    dlg.close()


def _first_paragraph(dlg):
    return dlg._layout.body[0]


def test_fields_dialog_text_edit_is_kept_in_memory_until_save(qapp, tmp_path, store, messages):
    dlg, folder = _window(tmp_path, store)
    para = _first_paragraph(dlg)
    before = (folder / "Шаблоны" / "tpl.docx").read_bytes()

    dlg._on_text_edited(para, "Договор {{ФИО}} на сумму {{Сумма}} №{{Номер}} ДОП")

    assert dlg._dirty is True
    assert (folder / "Шаблоны" / "tpl.docx").read_bytes() == before
    assert "ДОП" in dlg._layout.body[0].pieces[-1].text
    dlg.close()


def test_fields_dialog_save_writes_template_and_keeps_backup(qapp, tmp_path, store):
    from docx import Document as DocxDocument

    dlg, folder = _window(tmp_path, store)
    dlg._on_text_edited(_first_paragraph(dlg), "новый текст {{ФИО}}")

    dlg._on_save()

    tpl = folder / "Шаблоны" / "tpl.docx"
    assert DocxDocument(str(tpl)).paragraphs[0].text == "новый текст {{ФИО}}"
    assert DocxDocument(str(tpl.with_name("tpl.docx.bak"))).paragraphs[0].text.startswith("Договор")
    assert dlg._dirty is False
    dlg.close()


def test_fields_dialog_text_edit_unchanged_text_not_dirty(qapp, tmp_path, store):
    dlg, _ = _window(tmp_path, store)
    para = _first_paragraph(dlg)

    dlg._on_text_edited(para, "Договор {{ФИО}} на сумму {{Сумма}} №{{Номер}}")

    assert dlg._dirty is False
    dlg.close()


def test_fields_dialog_new_placeholder_gets_unset_button(qapp, tmp_path, store):
    dlg, _ = _window(tmp_path, store)

    dlg._on_text_edited(_first_paragraph(dlg), "Договор {{ФИО}} {{Новое}}")

    assert "Новое" in dlg._sheet.field_names
    assert dlg._sheet._buttons["Новое"][0].property("state") == "unset"
    dlg.close()


def test_fields_dialog_removed_placeholder_moves_to_unused(qapp, tmp_path, store):
    dlg, _ = _window(tmp_path, store)

    dlg._on_text_edited(_first_paragraph(dlg), "Договор {{ФИО}}")

    assert "Сумма" in dlg._unused_names()
    dlg.close()


def test_fields_dialog_paragraph_with_hyperlink_is_refused(qapp, tmp_path, store, messages):
    from docx.oxml import OxmlElement

    dlg, _ = _window(tmp_path, store)
    para = _first_paragraph(dlg)
    para.element.append(OxmlElement("w:hyperlink"))

    dlg._begin_text_edit(para)

    assert messages == [STRINGS.TPL_EDIT_UNSUPPORTED]
    assert dlg._sheet._editing_para is None
    dlg.close()


def test_fields_dialog_double_click_opens_editor_with_paragraph_text(qapp, tmp_path, store):
    from PyQt5.QtWidgets import QPlainTextEdit

    dlg, _ = _window(tmp_path, store)

    dlg._begin_text_edit(_first_paragraph(dlg))

    editor = dlg._sheet.findChild(QPlainTextEdit, "paragraphEditor")
    assert editor.toPlainText() == "Договор {{ФИО}} на сумму {{Сумма}} №{{Номер}}"
    dlg.close()
