# Copyright (C) 2026 Kirill Borovoy
# SPDX-License-Identifier: GPL-3.0-or-later
"""Check dialog «Проверить» / «Закрыть» flow (024), headless."""

import pytest
from docx import Document
from openpyxl import Workbook
from PyQt5.QtWidgets import QApplication

from stanok.gui.project_check_dialog import ProjectCheckDialog
from stanok.gui.strings import STRINGS
from stanok.services.project_setup import create_skeleton
from stanok.services.storage import ProjectStore


@pytest.fixture(scope="session")
def qapp():
    app = QApplication.instance() or QApplication([])
    yield app


@pytest.fixture
def store(tmp_path):
    return ProjectStore(tmp_path / "home")


def _fill(folder, rows=(("Иван",),)):
    doc = Document()
    doc.add_paragraph("Договор {{ФИО}}")
    doc.save(folder / "Шаблоны" / "tpl.docx")
    wb = Workbook()
    wb.active.append(["ФИО"])
    for row in rows:
        wb.active.append(list(row))
    wb.save(folder / "Данные" / "data.xlsx")


def test_check_dialog_creates_project_when_ok(qapp, store, tmp_path):
    folder = create_skeleton(tmp_path, "клиент", store)
    _fill(folder)
    dlg = ProjectCheckDialog(store, folder, "клиент")
    dlg._on_check()
    assert dlg.result() == dlg.Accepted
    assert dlg.project_name == "клиент"
    assert store.config_exists("клиент")
    dlg.close()


def test_check_dialog_lists_errors_and_creates_nothing(qapp, store, tmp_path):
    folder = create_skeleton(tmp_path, "пусто", store)
    dlg = ProjectCheckDialog(store, folder, "пусто")
    dlg._on_check()
    assert dlg.table.rowCount() >= 2
    assert STRINGS.PCHK_NO_DOCX in [dlg.table.item(r, 2).text() for r in range(dlg.table.rowCount())]
    assert "Ошибок" in dlg.status.text()
    assert not store.config_exists("пусто")
    dlg.close()


def test_check_dialog_fix_and_recheck_passes(qapp, store, tmp_path):
    folder = create_skeleton(tmp_path, "исправлю", store)
    dlg = ProjectCheckDialog(store, folder, "исправлю")
    dlg._on_check()
    assert not store.config_exists("исправлю")
    _fill(folder)
    dlg._on_check()
    assert store.config_exists("исправлю")
    dlg.close()


def test_check_dialog_warning_needs_second_click(qapp, store, tmp_path):
    folder = create_skeleton(tmp_path, "пустой", store)
    _fill(folder, rows=())
    dlg = ProjectCheckDialog(store, folder, "пустой")
    dlg._on_check()
    assert not store.config_exists("пустой")
    assert dlg.check_btn.text() == STRINGS.PCHK_CREATE
    dlg._on_check()
    assert store.config_exists("пустой")
    dlg.close()


def test_check_dialog_create_dirs_button(qapp, store, tmp_path):
    folder = tmp_path / "голая"
    folder.mkdir()
    dlg = ProjectCheckDialog(store, folder, "голая")
    assert not dlg.dirs_btn.isHidden()
    dlg.dirs_btn.click()
    assert (folder / "Данные").is_dir()
    assert (folder / "Шаблоны").is_dir()
    dlg.close()


def test_check_dialog_name_editable_only_when_asked(qapp, store, tmp_path):
    folder = create_skeleton(tmp_path, "фикс", store)
    fixed = ProjectCheckDialog(store, folder, "фикс")
    assert fixed.name_edit.isReadOnly()
    fixed.close()
    free = ProjectCheckDialog(store, folder, "фикс", editable_name=True)
    assert not free.name_edit.isReadOnly()
    free.close()
