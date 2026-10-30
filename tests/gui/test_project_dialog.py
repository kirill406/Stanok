# Copyright (C) 2026 Kirill Borovoy
# SPDX-License-Identifier: GPL-3.0-or-later
"""Offscreen tests for Window 2 (ProjectDialog)."""

import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from PyQt5.QtCore import QEventLoop, QTimer
from PyQt5.QtWidgets import (
    QApplication,
    QHeaderView,
    QLabel,
    QMessageBox,
    QSpinBox,
)

from stanok.gui.project_dialog import ProjectDialog
from stanok.gui.strings import STRINGS
from stanok.services.storage import ProjectStore, StorageError
from tests.services.test_generate import make_project


@pytest.fixture(scope="session")
def qapp():
    app = QApplication.instance() or QApplication([])
    yield app


@pytest.fixture
def store(tmp_path):
    return ProjectStore(tmp_path / ".stanok")


def _await_run(dlg, timeout_ms=15000):
    loop = QEventLoop()
    poll = QTimer()
    poll.timeout.connect(lambda: dlg._progress_dialog is None and loop.quit())
    poll.start(50)
    QTimer.singleShot(timeout_ms, loop.quit)
    loop.exec_()
    poll.stop()


def test_open_lists_templates(qapp, tmp_path, store):
    folder = make_project(tmp_path / "proj")
    dlg = ProjectDialog(str(folder), store=store)
    assert dlg.table.rowCount() == 1
    assert dlg.table.item(0, 0).text() == "Договор"
    assert dlg.table.cellWidget(0, 2).value() == 100  # make_project's large test limit
    assert "proj" in dlg.windowTitle()
    dlg.close()


def test_template_columns_order(qapp, tmp_path, store):
    folder = make_project(tmp_path / "proj")
    dlg = ProjectDialog(str(folder), store=store)
    headers = [dlg.table.horizontalHeaderItem(i).text() for i in range(6)]
    assert headers == [
        STRINGS.PROJ_TPL_COL,
        STRINGS.PROJ_DEFAULT_COL,
        STRINGS.PROJ_COUNT_COL,
        STRINGS.PROJ_RUN,
        STRINGS.PROJ_EDIT,
        STRINGS.TPL_REMOVE,
    ]
    assert dlg.table.cellWidget(0, 4).text() == STRINGS.PROJ_EDIT
    dlg.close()


def test_template_columns_all_interactive(qapp, tmp_path, store):
    folder = make_project(tmp_path / "proj")
    dlg = ProjectDialog(str(folder), store=store)
    header = dlg.table.horizontalHeader()
    assert not header.stretchLastSection()
    assert all(
        header.sectionResizeMode(col) == QHeaderView.Interactive
        for col in range(6)
    )
    dlg.close()


def test_long_template_name_grows_row(qapp, tmp_path, store):
    import json

    folder = make_project(tmp_path / "proj")
    path = folder / "project.stanok"
    data = json.loads(path.read_text())
    name = "Очень длинное название шаблона " * 4
    data["templates"][name] = json.loads(json.dumps(data["templates"]["Договор"]))
    path.write_text(json.dumps(data, ensure_ascii=False))
    dlg = ProjectDialog(str(folder), store=store)
    row = dlg.table.rowCount() - 1
    dlg.table.setColumnWidth(0, 80)
    qapp.processEvents()
    assert dlg.table.rowHeight(row) > 42
    dlg.close()


def test_sources_types_hint_shown(qapp, tmp_path, store):
    folder = make_project(tmp_path / "proj")
    dlg = ProjectDialog(str(folder), store=store)
    texts = [label.text() for label in dlg.findChildren(QLabel)]
    assert STRINGS.PROJ_SRC_TYPES_HINT in texts
    dlg.close()


def test_new_project_highlights_sources_and_edit(qapp, tmp_path, store):
    folder = make_project(tmp_path / "proj")
    dlg = ProjectDialog(str(folder), store=store, highlight_new=True)
    assert len(dlg._source_blinks) == 1 and dlg._source_blinks[0].running
    assert set(dlg._edit_blinks) == {"Договор"}
    assert dlg._edit_blinks["Договор"].running
    dlg.close()
    assert dlg._source_blinks == []
    assert dlg._edit_blinks == {}


def test_source_hint_stops_when_popup_opens(qapp, tmp_path, store):
    folder = make_project(tmp_path / "proj")
    dlg = ProjectDialog(str(folder), store=store, highlight_new=True)
    combo = dlg.sources_table.cellWidget(0, 1)
    assert "background-color: white" in combo.view().styleSheet()
    assert len(dlg._source_blinks) == 1
    combo.showPopup()
    combo.hidePopup()
    assert dlg._source_blinks == []
    assert combo.styleSheet() == ""
    dlg.close()


def test_existing_project_has_no_hints(qapp, tmp_path, store):
    folder = make_project(tmp_path / "proj")
    dlg = ProjectDialog(str(folder), store=store)
    assert dlg._source_blinks == []
    assert dlg._edit_blinks == {}
    dlg.close()


def test_source_hint_stops_after_mode_change(qapp, tmp_path, store):
    folder = make_project(tmp_path / "proj")
    dlg = ProjectDialog(str(folder), store=store, highlight_new=True)
    combo = dlg.sources_table.cellWidget(0, 1)
    combo.setCurrentIndex(combo.findData("circular"))
    assert dlg._source_blinks == []
    assert "Договор" in dlg._edit_blinks
    dlg.close()


def test_edit_click_stops_its_hint(qapp, tmp_path, store, monkeypatch):
    monkeypatch.setattr(
        "stanok.gui.project_dialog.FieldsDialog",
        lambda ref, name, st, parent: type("Fake", (), {"exec_": lambda self: None})(),
    )
    folder = make_project(tmp_path / "proj")
    dlg = ProjectDialog(str(folder), store=store, highlight_new=True)
    dlg.table.cellWidget(0, 4).click()
    assert dlg._edit_blinks == {}
    dlg.close()


def test_run_template_with_limit(qapp, tmp_path, store, monkeypatch):
    infos = []
    monkeypatch.setattr(QMessageBox, "information", lambda *a: infos.append(a[-1]))
    folder = make_project(
        tmp_path / "proj", rows=(("А", 1), ("Б", 2), ("В", 3))
    )
    dlg = ProjectDialog(str(folder), store=store)
    dlg.table.cellWidget(0, 2).setValue(1)
    dlg._on_run_template("Договор")
    _await_run(dlg)
    assert infos and "создано 1" in infos[0]
    assert len(list((folder / "Результат").glob("*.docx"))) == 1
    dlg.close()


def test_edit_button_opens_real_dialog(qapp, tmp_path, store, monkeypatch):
    exec_calls = []
    monkeypatch.setattr(
        "stanok.gui.project_dialog.FieldsDialog",
        lambda ref, name, st, parent: type(
            "Fake", (), {"exec_": lambda self: exec_calls.append((ref, name))}
        )(),
    )
    folder = make_project(tmp_path / "proj")
    dlg = ProjectDialog(str(folder), store=store)
    dlg.table.cellWidget(0, 4).click()
    assert exec_calls == [(str(folder), "Договор")]
    dlg.close()


def test_run_row_errors_shown(qapp, tmp_path, store, monkeypatch):
    infos = []
    monkeypatch.setattr(QMessageBox, "information", lambda *a: infos.append(a[-1]))
    import stanok.services.generate as gen

    real_render = gen.render
    calls = {"n": 0}

    def flaky(fj, template_path):
        calls["n"] += 1
        if calls["n"] == 2:
            raise RuntimeError("boom")
        return real_render(fj, template_path)

    monkeypatch.setattr(gen, "render", flaky)
    folder = make_project(tmp_path / "proj")
    dlg = ProjectDialog(str(folder), store=store)
    dlg._on_run_template("Договор")
    _await_run(dlg)
    assert infos and "boom" in infos[0]
    dlg.close()


def test_double_run_ignored(qapp, tmp_path, store, monkeypatch):
    monkeypatch.setattr(QMessageBox, "information", lambda *a: None)
    folder = make_project(tmp_path / "proj")
    dlg = ProjectDialog(str(folder), store=store)
    dlg._on_run_template("Договор")
    first_worker = dlg._worker
    dlg._on_run_template("Договор")
    assert dlg._worker is first_worker
    _await_run(dlg)
    dlg.close()


def test_broken_template_shows_error(qapp, tmp_path, store, monkeypatch):
    shown = []
    monkeypatch.setattr(QMessageBox, "critical", lambda *a: shown.append(a[-1]))
    monkeypatch.setattr(QMessageBox, "information", lambda *a: None)
    folder = make_project(tmp_path / "proj")
    (folder / "Шаблоны" / "tpl.docx").unlink()
    dlg = ProjectDialog(str(folder), store=store)
    dlg._on_run_template("Договор")
    _await_run(dlg)
    assert shown
    assert dlg.isEnabled()
    dlg.close()


def test_open_missing_project_raises(qapp, tmp_path, store):
    with pytest.raises(StorageError):
        ProjectDialog("nope", store=store)


def test_source_mode_change_persists(qapp, tmp_path, store):
    import json

    folder = make_project(tmp_path / "proj")
    dlg = ProjectDialog(str(folder), store=store)
    assert dlg.sources_table.rowCount() == 1
    assert dlg.sources_table.item(0, 0).text() == "Данные/data.xlsx"
    combo = dlg.sources_table.cellWidget(0, 1)
    assert combo.currentData() == "sequential"
    combo.setCurrentIndex(combo.findData("circular"))
    saved = json.loads((store.home_dir / "proj.stanok").read_text())
    assert saved["data_sources"][0]["mode"] == "circular"
    dlg.close()


def _extra_docx(path, text="Акт {{Номер}}"):
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_paragraph(text)
    doc.save(path)
    return path


def test_add_template_reloads_table(qapp, tmp_path, store, monkeypatch):
    from PyQt5.QtWidgets import QFileDialog

    folder = make_project(tmp_path / "proj")
    extra = _extra_docx(tmp_path / "extra.docx")
    monkeypatch.setattr(
        QFileDialog, "getOpenFileName", lambda *a, **k: (str(extra), "")
    )
    dlg = ProjectDialog(str(folder), store=store)
    assert dlg.table.rowCount() == 1
    dlg._on_add_template()
    assert dlg.table.rowCount() == 2
    assert dlg.table.item(1, 0).text() == "extra"
    dlg.close()


def test_remove_template_with_confirm(qapp, tmp_path, store, monkeypatch):
    folder = make_project(tmp_path / "proj")
    extra = _extra_docx(tmp_path / "extra.docx")
    store.add_template(folder, extra)
    monkeypatch.setattr("stanok.gui.project_dialog.ask_yes_no", lambda *a, **k: True)
    infos = []
    monkeypatch.setattr(QMessageBox, "information", lambda *a: infos.append(a[-1]))
    dlg = ProjectDialog(str(folder), store=store)
    assert dlg.table.rowCount() == 2
    dlg._on_remove_template("extra")
    assert dlg.table.rowCount() == 1
    dlg._on_remove_template("Договор")
    assert infos == [STRINGS.TPL_LAST_KEPT]
    assert dlg.table.rowCount() == 1
    dlg.close()


def test_remove_default_template_locked(qapp, tmp_path, store, monkeypatch):
    folder = make_project(tmp_path / "proj")
    extra = _extra_docx(tmp_path / "extra.docx")
    store.add_template(folder, extra)
    infos = []
    monkeypatch.setattr(QMessageBox, "information", lambda *a: infos.append(a[-1]))
    dlg = ProjectDialog(str(folder), store=store)
    dlg._on_remove_template("Договор")
    assert infos == [STRINGS.TPL_DEFAULT_LOCKED]
    assert dlg.table.rowCount() == 2
    dlg.close()


def test_default_template_radio_changes(qapp, tmp_path, store):
    import json

    folder = make_project(tmp_path / "proj")
    data = json.loads((folder / "project.stanok").read_text())
    data["templates"]["Второй"] = json.loads(
        json.dumps(data["templates"]["Договор"])
    )
    (folder / "project.stanok").write_text(json.dumps(data, ensure_ascii=False))
    dlg = ProjectDialog(str(folder), store=store)
    assert dlg.table.rowCount() == 2
    dlg.table.cellWidget(1, 1).setChecked(True)
    pj, _ = store.resolve_project(str(folder))
    assert pj.default_template == "Второй"
    dlg.close()


def test_source_mode_change_error_warns(qapp, tmp_path, store, monkeypatch):
    folder = make_project(tmp_path / "proj")
    warnings = []
    monkeypatch.setattr(
        QMessageBox, "warning", lambda *a: warnings.append(a[-1])
    )
    dlg = ProjectDialog(str(folder), store=store)

    def boom(*a, **k):
        raise RuntimeError("locked")

    monkeypatch.setattr(store, "set_source_mode", boom)
    combo = dlg.sources_table.cellWidget(0, 1)
    combo.setCurrentIndex(combo.findData("circular"))
    assert warnings
    dlg.close()


def test_run_template_limit_save_error(qapp, tmp_path, store, monkeypatch):
    folder = make_project(tmp_path / "proj")
    warnings = []
    monkeypatch.setattr(
        QMessageBox, "warning", lambda *a: warnings.append(a[-1])
    )
    dlg = ProjectDialog(str(folder), store=store)

    def boom(*a, **k):
        raise RuntimeError("locked")

    monkeypatch.setattr(store, "set_template_limit", boom)
    dlg._on_run_template("Договор")
    assert warnings
    assert not dlg._running
    dlg.close()


def test_template_limit_saved_on_change(qapp, tmp_path, store):
    import json

    folder = make_project(tmp_path / "proj")
    dlg = ProjectDialog(str(folder), store=store)
    dlg.table.cellWidget(0, 2).setValue(3)
    saved = json.loads((store.home_dir / "proj.stanok").read_text())
    assert saved["templates"]["Договор"]["limit"] == 3
    dlg.close()


def test_template_limit_change_error_warns(qapp, tmp_path, store, monkeypatch):
    folder = make_project(tmp_path / "proj")
    warnings = []
    monkeypatch.setattr(
        QMessageBox, "warning", lambda *a: warnings.append(a[-1])
    )
    dlg = ProjectDialog(str(folder), store=store)

    def boom(*a, **k):
        raise RuntimeError("locked")

    monkeypatch.setattr(store, "set_template_limit", boom)
    dlg.table.cellWidget(0, 2).setValue(4)
    assert warnings
    dlg.close()


def test_settings_button_opens_edit_dialog(qapp, tmp_path, store, monkeypatch):
    import stanok.gui.create_dialog as cd

    opened = []

    class Fake:
        def __init__(self, *a, **k):
            opened.append(k.get("edit_ref"))

        def exec_(self):
            return True

    monkeypatch.setattr(cd, "CreateDialog", Fake)
    folder = make_project(tmp_path / "proj")
    dlg = ProjectDialog(str(folder), store=store)
    dlg._on_settings()
    assert opened == [str(folder)]
    dlg.close()


def test_settings_button_builds_real_create_dialog(qapp, tmp_path, store, monkeypatch):
    import stanok.gui.create_dialog as cd

    folder = make_project(tmp_path / "proj")
    monkeypatch.setattr(cd.CreateDialog, "exec_", lambda self: 0)
    dlg = ProjectDialog(str(folder), store=store)
    dlg._on_settings()
    dlg.close()


def test_delete_project_closes_dialog(qapp, tmp_path, store, monkeypatch):
    folder = make_project(tmp_path / "proj")
    store.add_recent(str(folder), "proj")
    monkeypatch.setattr(QMessageBox, "exec_", lambda self: QMessageBox.Yes)
    infos = []
    monkeypatch.setattr(QMessageBox, "information", lambda *a: infos.append(a[-1]))
    dlg = ProjectDialog(str(folder), store=store)
    dlg._on_delete_project()
    assert dlg.result()
    assert not (store.home_dir / "proj.stanok").exists()
    assert infos and "удалён" in infos[0]
    dlg.close()


def test_delete_project_confirm_buttons_russian(qapp, tmp_path, store, monkeypatch):
    folder = make_project(tmp_path / "proj")
    seen = {}

    def fake_exec(box):
        seen["yes"] = box.button(QMessageBox.Yes).text()
        seen["no"] = box.button(QMessageBox.No).text()
        return QMessageBox.No

    monkeypatch.setattr(QMessageBox, "exec_", fake_exec)
    dlg = ProjectDialog(str(folder), store=store)
    dlg._on_delete_project()
    assert seen == {"yes": "Да", "no": "Нет"}
    dlg.close()


def test_manage_slots_guards_and_errors(qapp, tmp_path, store, monkeypatch):
    from PyQt5.QtWidgets import QFileDialog

    folder = make_project(tmp_path / "proj")
    warnings = []
    monkeypatch.setattr(
        QMessageBox, "warning", lambda *a: warnings.append(a[-1])
    )
    monkeypatch.setattr(
        "stanok.gui.project_dialog.ask_yes_no", lambda *a, **k: False
    )
    monkeypatch.setattr(
        QFileDialog, "getOpenFileName", lambda *a, **k: ("", "")
    )
    monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: None)
    dlg = ProjectDialog(str(folder), store=store)

    dlg._running = True

    dlg._running = True
    dlg._on_add_template()
    dlg._on_remove_template("Договор")
    dlg._on_delete_project()
    assert dlg.table.rowCount() == 1
    dlg._running = False

    dlg._on_add_template()
    assert dlg.table.rowCount() == 1

    monkeypatch.setattr(
        QFileDialog, "getOpenFileName", lambda *a, **k: (str(tmp_path / "gone.docx"), "")
    )
    dlg._on_add_template()
    assert warnings and dlg.table.rowCount() == 1

    extra = _extra_docx(tmp_path / "extra2.docx")
    store.add_template(folder, extra)
    dlg._reload_templates()
    assert dlg.table.rowCount() == 2
    dlg._on_remove_template("extra2")
    assert dlg.table.rowCount() == 2

    monkeypatch.setattr(
        "stanok.gui.project_dialog.ask_yes_no", lambda *a, **k: True
    )
    dlg._on_remove_template("extra2")
    assert dlg.table.rowCount() == 1

    def boom(*a, **k):
        raise RuntimeError("disk gone")

    monkeypatch.setattr(store, "remove_template", boom)
    extra = _extra_docx(tmp_path / "extra.docx")
    store.add_template(folder, extra)
    dlg._reload_templates()
    assert dlg.table.rowCount() == 2
    dlg._on_remove_template("extra")
    assert len(warnings) == 2
    dlg.close()


def test_delete_project_errors(qapp, tmp_path, store, monkeypatch):
    folder = make_project(tmp_path / "proj")
    warnings = []
    monkeypatch.setattr(
        QMessageBox, "warning", lambda *a: warnings.append(a[-1])
    )
    monkeypatch.setattr(QMessageBox, "exec_", lambda self: QMessageBox.No)
    dlg = ProjectDialog(str(folder), store=store)
    dlg._on_delete_project()
    assert not dlg.result()
    assert (store.home_dir / "proj.stanok").exists()

    monkeypatch.setattr(QMessageBox, "exec_", lambda self: QMessageBox.Yes)

    def boom(*a, **k):
        raise RuntimeError("locked")

    monkeypatch.setattr(store, "delete_project", boom)
    dlg._on_delete_project()
    assert warnings
    assert (store.home_dir / "proj.stanok").exists()
    dlg.close()


def test_fields_dialog_broken_config_warns(qapp, tmp_path, store, monkeypatch):
    folder = make_project(tmp_path / "proj")
    warnings = []
    monkeypatch.setattr(
        QMessageBox, "warning", lambda *a: warnings.append(a[-1])
    )
    dlg = ProjectDialog(str(folder), store=store)
    (store.home_dir / "proj.stanok").write_text("{broken")
    dlg._open_fields_dialog("Договор")
    assert warnings
    dlg.close()


def test_fields_dialog_missing_template_warns(qapp, tmp_path, store, monkeypatch):
    import json

    folder = make_project(tmp_path / "proj")
    warnings = []
    monkeypatch.setattr(
        QMessageBox, "warning", lambda *a: warnings.append(a[-1])
    )
    dlg = ProjectDialog(str(folder), store=store)
    data = json.loads((store.home_dir / "proj.stanok").read_text())
    data["templates"] = {"Чужой": data["templates"]["Договор"]}
    data["default_template"] = "Чужой"
    (store.home_dir / "proj.stanok").write_text(json.dumps(data))
    dlg._open_fields_dialog("Договор")
    assert warnings and "template not found" in warnings[0]
    dlg.close()


def test_reload_broken_config_keeps_rows(qapp, tmp_path, store, monkeypatch):
    folder = make_project(tmp_path / "proj")
    warnings = []
    monkeypatch.setattr(
        QMessageBox, "warning", lambda *a: warnings.append(a[-1])
    )
    dlg = ProjectDialog(str(folder), store=store)
    assert dlg.table.rowCount() == 1
    (store.home_dir / "proj.stanok").write_text("{broken")
    dlg._reload_templates()
    assert warnings
    assert dlg.table.rowCount() == 1
    assert dlg.table.item(0, 0).text() == "Договор"
    dlg.close()


def test_run_stale_template_shows_error(qapp, tmp_path, store, monkeypatch):
    folder = make_project(tmp_path / "proj", rows=(("Иван", 100),))
    shown = []
    monkeypatch.setattr(QMessageBox, "critical", lambda *a: shown.append(a[-1]))
    monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: None)
    dlg = ProjectDialog(str(folder), store=store)
    (folder / "Шаблоны" / "tpl.docx").unlink()
    dlg._on_run_template("Договор")
    _await_run(dlg)
    assert shown and "template file not found" in shown[0]
    dlg.close()


def test_run_stale_source_shows_error(qapp, tmp_path, store, monkeypatch):
    folder = make_project(tmp_path / "proj", rows=(("Иван", 100),))
    shown = []
    monkeypatch.setattr(QMessageBox, "critical", lambda *a: shown.append(a[-1]))
    monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: None)
    dlg = ProjectDialog(str(folder), store=store)
    (folder / "Данные" / "data.xlsx").unlink()
    dlg._on_run_template("Договор")
    _await_run(dlg)
    assert shown
    dlg.close()


def test_close_button_russian(qapp, tmp_path, store):
    from PyQt5.QtWidgets import QPushButton

    folder = make_project(tmp_path / "proj")
    dlg = ProjectDialog(str(folder), store=store)
    texts = [
        b.text()
        for b in dlg.findChildren(QPushButton)
        if b.text() == STRINGS.FLD_CLOSE
    ]
    assert len(texts) == 1
    dlg.close()


def test_open_syncs_new_fields(qapp, tmp_path, store):
    from docx import Document as DocxDocument

    from stanok.gui.fields_dialog import FieldsDialog
    from tests.services.test_generate import make_project

    folder = make_project(tmp_path / "proj")
    tpl = folder / "Шаблоны" / "tpl.docx"
    doc = DocxDocument(str(tpl))
    doc.add_paragraph("Штрихкод {{Штрихкод}}")
    doc.save(str(tpl))
    dlg = ProjectDialog(str(folder), store=store)
    dlg.close()
    fld = FieldsDialog(str(folder), "Договор", store=store)
    assert "Штрихкод" in fld._sheet.field_names
    fld.close()


def test_add_source_reloads_table(qapp, tmp_path, store, monkeypatch):
    from PyQt5.QtWidgets import QFileDialog

    from tests.services.test_storage import _src_files

    folder = make_project(tmp_path / "proj")
    xlsx, _ = _src_files(tmp_path, "extra")
    monkeypatch.setattr(
        QFileDialog, "getOpenFileName", lambda *a, **k: (str(xlsx), "")
    )
    dlg = ProjectDialog(str(folder), store=store)
    assert dlg.sources_table.rowCount() == 1
    dlg._on_add_source()
    assert dlg.sources_table.rowCount() == 2
    from PyQt5.QtWidgets import QLabel

    assert any(
        w.text() == STRINGS.PROJ_SOURCES_TITLE for w in dlg.findChildren(QLabel)
    )
    assert dlg.add_src_btn.text() == STRINGS.PROJ_SRC_ADD
    dlg.close()


def test_add_source_guards(qapp, tmp_path, store, monkeypatch):
    from PyQt5.QtWidgets import QFileDialog

    folder = make_project(tmp_path / "proj")
    warnings = []
    monkeypatch.setattr(
        QMessageBox, "warning", lambda *a: warnings.append(a[-1])
    )
    dlg = ProjectDialog(str(folder), store=store)

    dlg._running = True
    dlg._on_add_source()
    dlg._running = False

    monkeypatch.setattr(
        QFileDialog, "getOpenFileName", lambda *a, **k: ("", "")
    )
    dlg._on_add_source()
    assert dlg.sources_table.rowCount() == 1

    bad = tmp_path / "notes.txt"
    bad.write_text("hi")
    monkeypatch.setattr(
        QFileDialog, "getOpenFileName", lambda *a, **k: (str(bad), "")
    )
    dlg._on_add_source()
    assert warnings and "wrong file type" in warnings[0]
    dlg.close()


def test_remove_source_button_detaches_after_confirm(qapp, tmp_path, store, monkeypatch):
    from openpyxl import Workbook

    folder = make_project(tmp_path / "proj")
    store.migrate_from_project(folder)
    extra = tmp_path / "extra.xlsx"
    wb = Workbook()
    wb.active.append(["Город"])
    wb.save(extra)
    store.add_source("proj", extra)
    dlg = ProjectDialog(str(folder), store=store)
    assert dlg.sources_table.columnCount() == 3
    rows = {dlg.sources_table.item(r, 0).text(): r for r in range(dlg.sources_table.rowCount())}
    remove_btn = dlg.sources_table.cellWidget(rows["Данные/extra.xlsx"], 2)
    monkeypatch.setattr("stanok.gui.project_dialog.ask_yes_no", lambda *a, **k: True)
    remove_btn.click()
    assert [ds.file for ds in store.load("proj").data_sources] == ["Данные/data.xlsx"]
    assert dlg.sources_table.rowCount() == 1
    dlg.close()
