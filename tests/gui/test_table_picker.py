# Copyright (C) 2026 Kirill Borovoy
# SPDX-License-Identifier: GPL-3.0-or-later
"""Tests for shared table-picker helpers (Window 3)."""

import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from openpyxl import Workbook
from PyQt5.QtWidgets import QApplication, QDialog, QMessageBox

from stanok.gui import table_picker
from stanok.gui.strings import STRINGS


@pytest.fixture(scope="session")
def qapp():
    app = QApplication.instance() or QApplication([])
    yield app


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        (None, None),
        (True, None),
        (0, 0),
        (5, 5),
        (-1, None),
        (2.0, 2),
        (2.5, None),
        (-2.0, None),
        ("7", 7),
        (" 8 ", 8),
        ("-3", None),
        ("abc", None),
        ("", None),
    ],
)
def test_coerce_number(raw, expected):
    assert table_picker.coerce_number(raw) == expected


def test_load_grid_reads_headers_and_rows(qapp, tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["ФИО", "Сумма"])
    ws.append(["Иван", 100])
    ws.append([None, None])
    ws.append(["Петр", 200])
    path = tmp_path / "data.xlsx"
    wb.save(path)
    grid = table_picker.load_grid(tmp_path, "data.xlsx")
    assert grid.headers == ["ФИО", "Сумма"]
    assert grid.rows == [["Иван", 100], [None, None], ["Петр", 200]]


def test_eye_icon_is_not_null(qapp):
    icon = table_picker.eye_icon()
    assert not icon.isNull()


def test_open_picker_returns_none_when_unavailable(monkeypatch):
    monkeypatch.setattr(table_picker, "picker_available", lambda: False)
    assert table_picker.open_picker(None, ".", "data.xlsx", 0) is None


def test_value_picker_rejects_non_number(qapp, tmp_path, monkeypatch):
    wb = Workbook()
    ws = wb.active
    ws.append(["Сумма"])
    ws.append([100])
    wb.save(tmp_path / "data.xlsx")
    warnings = []
    monkeypatch.setattr(
        QMessageBox,
        "warning",
        lambda *a, **k: warnings.append(a),
    )
    dlg = table_picker.ValuePickerDialog(tmp_path, "data.xlsx", 1)
    dlg.choose("abc")
    assert dlg.value is None
    assert warnings and warnings[0][2] == STRINGS.FLD_PICK_NOT_NUMBER


def test_value_picker_load_and_choose(qapp, tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Сумма"])
    ws.append([100])
    wb.save(tmp_path / "data.xlsx")
    dlg = table_picker.ValuePickerDialog(tmp_path, "data.xlsx", 1)
    grid = dlg.load()
    assert grid is not None
    assert grid.headers == ["Сумма"]
    assert grid.rows == [[100]]
    dlg.choose(100)
    assert dlg.value == 100
    assert dlg.result() == QDialog.Accepted


def test_value_picker_load_reports_error(qapp, tmp_path, monkeypatch):
    warnings = []
    monkeypatch.setattr(
        QMessageBox, "warning", lambda *a, **k: warnings.append(a)
    )
    dlg = table_picker.ValuePickerDialog(tmp_path, "missing.xlsx", 0)
    assert dlg.load() is None
    assert warnings


def test_open_picker_uses_variant_module(monkeypatch):
    import sys
    import types

    class FakeDialog:
        def __init__(self, folder, rel, current, parent):
            self.value = 9

        def exec_(self):
            return QDialog.Accepted

    module = types.ModuleType("fake_picker")
    module.PickerDialog = FakeDialog
    monkeypatch.setattr(table_picker, "picker_available", lambda: True)
    monkeypatch.setitem(sys.modules, table_picker.PICKER_MODULE, module)
    assert table_picker.open_picker(None, ".", "data.xlsx", 0) == 9

