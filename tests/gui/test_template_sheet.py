# Copyright (C) 2026 Kirill Borovoy
# SPDX-License-Identifier: GPL-3.0-or-later
"""Offscreen tests for the template page (027): field buttons, lines, layout."""

import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from docx import Document
from docx.enum.text import WD_BREAK
from PyQt5.QtCore import QRect
from PyQt5.QtCore import Qt
from PyQt5.QtWidgets import QApplication, QLabel, QWidget

import pytest

from stanok.gui.strings import STRINGS
from stanok.gui.template_sheet import FlowWidget, TemplateSheet
from stanok.services.template_layout import load_template_layout


@pytest.fixture(scope="session")
def qapp():
    app = QApplication.instance() or QApplication([])
    yield app


def _layout(tmp_path, build):
    doc = Document()
    build(doc)
    path = tmp_path / "tpl.docx"
    doc.save(path)
    return load_template_layout(path)


def _sheet(qapp, layout, configured=lambda name: True):
    sheet = TemplateSheet(
        layout,
        describe=lambda name: f"описание {name}",
        is_configured=configured,
    )
    sheet.resize(900, 700)
    sheet.show()  # geometry is computed on show
    qapp.processEvents()
    return sheet


def test_template_sheet_field_becomes_button_with_field_name(qapp, tmp_path):
    layout = _layout(tmp_path, lambda d: d.add_paragraph("Договор {{ФИО}} на {{Сумма}}"))
    sheet = _sheet(qapp, layout)

    assert sheet.field_names == ["ФИО", "Сумма"]
    assert [b.text() for b in sheet._buttons["ФИО"]] == ["ФИО"]


def test_template_sheet_repeated_field_has_button_per_place(qapp, tmp_path):
    def build(doc):
        doc.add_paragraph("{{Имя}}")
        doc.add_paragraph("снова {{Имя}}")

    sheet = _sheet(qapp, _layout(tmp_path, build))

    assert sheet.field_names == ["Имя"]
    assert len(sheet._buttons["Имя"]) == 2


def test_template_sheet_click_on_field_emits_its_name(qapp, tmp_path):
    sheet = _sheet(qapp, _layout(tmp_path, lambda d: d.add_paragraph("{{Город}}")))
    clicked: list[str] = []
    sheet.fieldClicked.connect(clicked.append)

    sheet._buttons["Город"][0].click()

    assert clicked == ["Город"]


def test_template_sheet_unconfigured_field_has_unset_state(qapp, tmp_path):
    def build(doc):
        doc.add_paragraph("{{Готово}} {{Нет}}")

    sheet = _sheet(qapp, _layout(tmp_path, build), configured=lambda n: n == "Готово")

    assert sheet._buttons["Готово"][0].property("state") == "set"
    assert sheet._buttons["Нет"][0].property("state") == "unset"


def test_template_sheet_refresh_follows_configuration(qapp, tmp_path):
    state = {"ready": False}
    sheet = _sheet(
        qapp,
        _layout(tmp_path, lambda d: d.add_paragraph("{{Поле}}")),
        configured=lambda n: state["ready"],
    )
    button = sheet._buttons["Поле"][0]
    assert button.property("state") == "unset"

    state["ready"] = True
    sheet.refresh()

    assert button.property("state") == "set"
    assert button.toolTip() == "описание Поле"


def test_template_sheet_line_break_starts_new_line(qapp, tmp_path):
    def build(doc):
        para = doc.add_paragraph("первая")
        para.runs[0].add_break()
        para.add_run("вторая")

    sheet = _sheet(qapp, _layout(tmp_path, build))

    labels = [w.text() for w in sheet.findChildren(QLabel) if w.text() in ("первая", "вторая")]
    assert labels == ["первая", "вторая"]
    first = next(w for w in sheet.findChildren(QLabel) if w.text() == "первая")
    second = next(w for w in sheet.findChildren(QLabel) if w.text() == "вторая")
    assert second.mapTo(sheet, second.rect().topLeft()).y() > first.mapTo(sheet, first.rect().topLeft()).y()


def test_template_sheet_page_break_draws_separator(qapp, tmp_path):
    def build(doc):
        para = doc.add_paragraph("до")
        para.runs[0].add_break(WD_BREAK.PAGE)
        para.add_run("после")

    sheet = _sheet(qapp, _layout(tmp_path, build))

    separators = sheet.findChildren(QLabel, "pageBreak")
    assert [s.text() for s in separators] == [STRINGS.TPL_PAGE_BREAK]


def test_template_sheet_header_and_footer_are_captioned(qapp, tmp_path):
    def build(doc):
        section = doc.sections[0]
        section.header.paragraphs[0].text = "шапка {{Орг}}"
        section.footer.paragraphs[0].text = "подвал"
        doc.add_paragraph("тело")

    sheet = _sheet(qapp, _layout(tmp_path, build))

    captions = [c.text() for c in sheet.findChildren(QLabel, "partCaption")]
    assert captions == [STRINGS.TPL_PART_HEADER, STRINGS.TPL_PART_FOOTER]
    assert sheet.field_names == ["Орг"]


def test_template_sheet_table_cells_hold_their_fields(qapp, tmp_path):
    def build(doc):
        table = doc.add_table(rows=1, cols=2)
        table.style = doc.styles["Table Grid"]
        table.cell(0, 0).text = "Товар"
        table.cell(0, 1).text = "{{Цена}}"

    sheet = _sheet(qapp, _layout(tmp_path, build))

    assert sheet.field_names == ["Цена"]
    assert sheet.findChildren(QWidget, "tableBlock")


def test_template_sheet_empty_paragraph_keeps_a_line(qapp, tmp_path):
    def build(doc):
        doc.add_paragraph("до")
        doc.add_paragraph()
        doc.add_paragraph("после")

    sheet = _sheet(qapp, _layout(tmp_path, build))

    texts = [w.text() for w in sheet.findChildren(QLabel)]
    assert "до" in texts and "после" in texts
    assert " " in texts  # the empty paragraph's placeholder label


def test_flow_widget_wraps_words_when_narrow(qapp):
    host = FlowWidget()
    for i in range(6):
        QLabel(f"слово{i}", host)

    wide = host.heightForWidth(2000)
    narrow = host.heightForWidth(60)

    assert narrow > wide


def test_flow_widget_right_alignment_puts_last_item_at_edge(qapp):
    host = FlowWidget(align="right")
    label = QLabel("конец", host)

    host.resize(400, 40)
    host.relayout()

    assert label.geometry().right() == 399


def test_flow_widget_center_alignment_centers_the_line(qapp):
    host = FlowWidget(align="center")
    label = QLabel("х", host)

    host.resize(400, 40)
    host.relayout()

    assert abs(label.geometry().center().x() - 200) <= 2


def test_flow_widget_first_indent_applies_to_first_line_only(qapp):
    host = FlowWidget(first_indent=50)
    first = QLabel("а", host)

    host.resize(400, 40)
    host.relayout()

    assert first.geometry().x() == 50


def test_flow_widget_removed_child_is_not_placed(qapp):
    host = FlowWidget()
    kept = QLabel("оставить", host)
    gone = QLabel("убрать", host)
    gone.setParent(None)
    gone.deleteLater()

    host.resize(400, 40)
    host.relayout()

    assert kept.geometry().x() == 0
    assert host._items() == [kept]


def test_template_sheet_double_click_elsewhere_commits_open_edit(qapp, tmp_path):
    sheet = _sheet(qapp, _layout(tmp_path, lambda d: d.add_paragraph("один")))
    edited, asked = [], []
    sheet.paragraphEdited.connect(lambda p, t: edited.append(t))
    sheet.paragraphActivated.connect(asked.append)
    sheet.begin_edit(sheet._layout_model.body[0], "один, правка")

    sheet._activate(sheet._layout_model.body[0])

    assert edited == ["один, правка"]
    assert asked == []
