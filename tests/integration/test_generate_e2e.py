# Copyright (C) 2026 Kirill Borovoy
# SPDX-License-Identifier: GPL-3.0-or-later
"""E2E: Excel + PJ → docx files in Результат/ with substituted placeholders."""

from docx import Document as DocxDocument

from stanok.services.generate import GenerateCommand, generate_documents
from stanok.services.storage import ProjectStore
from tests.services.test_generate import make_project


def test_generate_e2e_excel_to_docx(tmp_path):
    folder = make_project(tmp_path / "e2e")
    store = ProjectStore(tmp_path / ".stanok")
    report = generate_documents(GenerateCommand(project_ref=folder), store=store)

    assert report.created == 2
    assert report.errors == []
    for path in report.output_paths:
        text = "\n".join(p.text for p in DocxDocument(str(path)).paragraphs)
        assert "{{" not in text
        assert "Договор" in text
    names = sorted(p.name for p in report.output_paths)
    assert names == ["Договор_Иван_1.docx", "Договор_Петр_2.docx"]
