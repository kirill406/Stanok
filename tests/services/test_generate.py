# Copyright (C) 2026 Kirill Borovoy
# SPDX-License-Identifier: GPL-3.0-or-later
"""Tests for generate_documents: orchestration tables → resolve → render (021)."""

import json
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook

from stanok.engine.schema import ResolveError, StorageError, validate_pj
from stanok.services.generate import (
    GenerateCommand,
    TemplateError,
    _unique_path,
    generate_documents,
)
from stanok.services.storage import ProjectStore

DATA = "Данные/data.xlsx"
TPL = "Шаблоны/tpl.docx"


def make_project(
    root: Path,
    *,
    headers=("ФИО", "Сумма"),
    rows=(("Иван", 100), ("Петр", 200)),
    mode="sequential",
    pos=0,
    with_counter=True,
    limit=0,
    resume=1,
) -> Path:
    """Build minimal project folder: Данные/data.xlsx + Шаблоны/tpl.docx + legacy PJ."""
    data_dir = root / "Данные"
    tpl_dir = root / "Шаблоны"
    data_dir.mkdir(parents=True)
    tpl_dir.mkdir(parents=True)

    wb = Workbook()
    ws = wb.active
    ws.append(list(headers))
    for r in rows:
        ws.append(list(r))
    wb.save(data_dir / "data.xlsx")

    fields = {
        "ФИО": {"source": "table", "value": "ФИО", "counter_name": "data"},
        "Сумма": {"source": "table", "value": "Сумма", "counter_name": "data"},
    }
    counters = {
        "data": {
            "pos": pos,
            "format": "1",
            "table": DATA,
            "resume": resume,
        }
    }
    doc = DocxDocument()
    text = "Договор {{ФИО}} на сумму {{Сумма}} №{{Номер}}"
    if with_counter:
        fields["Номер"] = {"source": "counter", "value": "num"}
        counters["num"] = {"pos": 0, "format": "1"}
    else:
        text = "Договор {{ФИО}} на сумму {{Сумма}}"
    doc.add_paragraph(text)
    doc.save(tpl_dir / "tpl.docx")

    pj = {
        "version": "0.0.0",
        "templates": {"Договор": {"file": TPL, "fields": fields, "limit": limit}},
        "default_template": "Договор",
        "data_sources": [{"file": DATA, "mode": mode}],
        "counters": counters,
        "filename_template": "Договор_{ФИО}_{i}.docx",
    }
    (root / "project.stanok").write_text(json.dumps(pj, ensure_ascii=False, indent=2))
    return root


@pytest.fixture
def store(tmp_path):
    return ProjectStore(tmp_path / ".stanok")


def _load_saved(store, folder_name):
    return validate_pj(
        json.loads((store.home_dir / f"{folder_name}.stanok").read_text())
    )


def _doc_text(path: Path) -> str:
    doc = DocxDocument(str(path))
    return "\n".join(p.text for p in doc.paragraphs)


def _write_pj(folder: Path, pj: dict) -> None:
    (folder / "project.stanok").write_text(
        json.dumps(pj, ensure_ascii=False, indent=2)
    )


def test_full_sequential_run(tmp_path, store):
    folder = make_project(tmp_path / "proj")
    report = generate_documents(GenerateCommand(project_ref=folder), store=store)
    assert report.created == 2
    assert report.errors == []
    assert len(report.output_paths) == 2
    assert all(p.exists() for p in report.output_paths)
    saved = _load_saved(store, "proj")
    assert saved.counters["num"].pos == 2
    # Two table fields share one counter → one bump per document, not two.
    assert saved.counters["data"].pos == 2
    assert saved.counters["data"].pos == 2


def test_limit_truncates(tmp_path, store):
    folder = make_project(
        tmp_path / "proj", rows=(("А", 1), ("Б", 2), ("В", 3)), limit=2
    )
    report = generate_documents(GenerateCommand(project_ref=folder), store=store)
    assert report.created == 2
    assert (report.planned, report.limit) == (2, 2)
    saved = _load_saved(store, "proj")
    assert saved.counters["num"].pos == 2
    assert saved.counters["data"].pos == 2


def test_limit_reports_data_shortfall(tmp_path, store):
    """Limit is a cap: sequential stops at the last row (2 docs out of 3)."""
    folder = make_project(
        tmp_path / "proj", rows=(("А", 1), ("Б", 2)), limit=3
    )
    report = generate_documents(GenerateCommand(project_ref=folder), store=store)
    assert report.created == 2
    assert (report.planned, report.limit) == (2, 3)


def test_pos_skips(tmp_path, store):
    folder = make_project(tmp_path / "proj", pos=1)
    report = generate_documents(GenerateCommand(project_ref=folder), store=store)
    assert report.created == 1
    assert "Петр" in report.output_paths[0].name


def test_resume_0_keeps_cursor(tmp_path, store):
    """resume=0 → pos stays; the numeric counter still advances (021)."""
    folder = make_project(tmp_path / "proj", resume=0)
    report = generate_documents(GenerateCommand(project_ref=folder), store=store)
    assert report.created == 2
    saved = _load_saved(store, "proj")
    assert saved.counters["data"].pos == 0
    assert saved.counters["num"].pos == 2


def test_empty_source_returns_empty_report(tmp_path, store):
    folder = make_project(tmp_path / "proj", rows=())
    report = generate_documents(GenerateCommand(project_ref=folder), store=store)
    assert report.created == 0
    assert report.errors == []


def test_empty_rows_still_generate(tmp_path, store):
    folder = make_project(
        tmp_path / "proj", rows=(("Иван", 100), (None, None), ("Петр", 200))
    )
    report = generate_documents(GenerateCommand(project_ref=folder), store=store)
    assert report.created == 3
    assert report.errors == []
    assert report.output_paths[1].exists()
    assert _doc_text(report.output_paths[1]) == "Договор  на сумму  №1"


def test_error_continues_and_advances(monkeypatch, tmp_path, store):
    folder = make_project(tmp_path / "proj")
    import stanok.services.generate as gen

    real_render = gen.render
    calls = {"n": 0}

    def flaky(fj, template_path):
        calls["n"] += 1
        if calls["n"] == 2:
            raise RuntimeError("boom")
        return real_render(fj, template_path)

    monkeypatch.setattr(gen, "render", flaky)
    report = generate_documents(GenerateCommand(project_ref=folder), store=store)
    assert report.created == 1
    assert len(report.errors) == 1
    assert report.errors[0][0] == 2
    assert "boom" in report.errors[0][1]
    saved = _load_saved(store, "proj")
    assert saved.counters["num"].pos == 2
    assert saved.counters["data"].pos == 2


def test_all_failed_still_persists_counters(monkeypatch, tmp_path, store):
    """All docs error (created=0) → counters/cursors still advance (021)."""
    folder = make_project(tmp_path / "proj")
    import stanok.services.generate as gen

    def boom(fj, template_path):
        raise RuntimeError("boom")

    monkeypatch.setattr(gen, "render", boom)
    report = generate_documents(GenerateCommand(project_ref=folder), store=store)
    assert report.created == 0
    assert len(report.errors) == 2
    saved = _load_saved(store, "proj")
    assert saved.counters["num"].pos == 2
    assert saved.counters["data"].pos == 2


def test_missing_template_file(tmp_path, store):
    folder = make_project(tmp_path / "proj")
    (folder / TPL).unlink()
    with pytest.raises(TemplateError, match="template file not found"):
        generate_documents(GenerateCommand(project_ref=folder), store=store)


def test_missing_project(store, tmp_path):
    with pytest.raises(StorageError):
        generate_documents(GenerateCommand(project_ref="nope"), store=store)


def test_unknown_template_name(tmp_path, store):
    folder = make_project(tmp_path / "proj")
    with pytest.raises(TemplateError, match="template not found"):
        generate_documents(
            GenerateCommand(project_ref=folder, template="Чужой"), store=store
        )


def test_unique_names_on_collision(tmp_path, store):
    folder = make_project(tmp_path / "proj", rows=(("Иван", 100),))
    out = folder / "Результат"
    out.mkdir()
    (out / "Договор_Иван_1.docx").write_bytes(b"old")
    report = generate_documents(GenerateCommand(project_ref=folder), store=store)
    assert report.created == 1
    assert report.output_paths[0].name == "Договор_Иван_1 (1).docx"
    assert (out / "Договор_Иван_1.docx").read_bytes() == b"old"


def test_constant_mode_cursor_stays(tmp_path, store):
    folder = make_project(
        tmp_path / "proj",
        rows=(("Иван", 100), ("Петр", 200), ("Сидор", 300)),
        mode="constant",
        limit=2,
    )
    report = generate_documents(GenerateCommand(project_ref=folder), store=store)
    assert report.created == 2
    saved = _load_saved(store, "proj")
    assert saved.counters["data"].pos == 0
    assert saved.counters["num"].pos == 2


def test_circular_cycles_and_wraps(tmp_path, store):
    folder = make_project(
        tmp_path / "proj", rows=(("А", 1), ("Б", 2)), mode="circular", limit=5
    )
    report = generate_documents(GenerateCommand(project_ref=folder), store=store)
    assert report.created == 5
    saved = _load_saved(store, "proj")
    assert saved.counters["num"].pos == 5
    assert saved.counters["data"].pos == 5  # absolute count; row wraps at read


def test_config_name_ref_via_recent(tmp_path, store):
    folder = make_project(tmp_path / "proj", rows=(("Иван", 100),))
    data = json.loads((folder / "project.stanok").read_text())
    data["counters"]["data"]["resume"] = 0
    _write_pj(folder, data)
    report = generate_documents(GenerateCommand(project_ref=folder), store=store)
    assert report.created == 1
    store.add_recent(str(folder), "proj")
    report2 = generate_documents(GenerateCommand(project_ref="proj"), store=store)
    assert report2.created == 1


def test_config_name_without_folder_fails(tmp_path, store):
    pj = validate_pj(
        {
            "version": "0.0.0",
            "templates": {"t": {"file": "Шаблоны/t.docx", "fields": {}}},
            "default_template": "t",
            "data_sources": [{"file": DATA, "mode": "sequential"}],
            "counters": {"data": {"format": "1", "table": DATA}},
            "filename_template": "nofolder",
        }
    )
    store.save(pj, name="nofolder")
    with pytest.raises(StorageError, match="project folder not found"):
        generate_documents(GenerateCommand(project_ref="nofolder"), store=store)


def test_unique_path_helper(tmp_path):
    base = tmp_path / "a.docx"
    assert _unique_path(base) == base
    base.write_bytes(b"x")
    assert _unique_path(base) == tmp_path / "a (1).docx"
    (tmp_path / "a (1).docx").write_bytes(b"x")
    assert _unique_path(base) == tmp_path / "a (2).docx"


def test_no_templates_fails(tmp_path, store):
    folder = make_project(tmp_path / "proj")
    data = json.loads((folder / "project.stanok").read_text())
    data["templates"] = {}
    data["default_template"] = ""
    _write_pj(folder, data)
    with pytest.raises(TemplateError, match="no templates"):
        generate_documents(GenerateCommand(project_ref=folder), store=store)


def test_named_template_pick(tmp_path, store):
    folder = make_project(tmp_path / "proj")
    data = json.loads((folder / "project.stanok").read_text())
    data["templates"]["Второй"] = data["templates"]["Договор"]
    _write_pj(folder, data)
    report = generate_documents(
        GenerateCommand(project_ref=folder, template="Второй"), store=store
    )
    assert report.created == 2


def test_default_template_used(tmp_path, store):
    folder = make_project(tmp_path / "proj", rows=(("А", 1),))
    data = json.loads((folder / "project.stanok").read_text())
    data["templates"]["Второй"] = json.loads(json.dumps(data["templates"]["Договор"]))
    data["templates"]["Второй"]["limit"] = 1
    data["templates"]["Договор"]["limit"] = 0
    data["default_template"] = "Второй"
    _write_pj(folder, data)
    report = generate_documents(GenerateCommand(project_ref=folder), store=store)
    assert report.created == 1


def test_missing_source_file_fails(tmp_path, store):
    folder = make_project(tmp_path / "proj")
    (folder / DATA).unlink()
    with pytest.raises(TemplateError, match="not readable"):
        generate_documents(GenerateCommand(project_ref=folder), store=store)


def test_missing_column_fails_fast(tmp_path, store):
    folder = make_project(tmp_path / "proj", rows=(("Иван", 100),))
    data = json.loads((folder / "project.stanok").read_text())
    data["templates"]["Договор"]["fields"]["Инфа"] = {
        "source": "table",
        "value": "инфа",
        "counter_name": "data",
    }
    _write_pj(folder, data)
    with pytest.raises(TemplateError, match="column 'инфа'.*data.xlsx"):
        generate_documents(GenerateCommand(project_ref=folder), store=store)
    assert not (folder / "Результат").exists()


def test_two_sources_composition(tmp_path, store):
    folder = make_project(tmp_path / "proj")
    cdir = folder / "Данные"
    wb = Workbook()
    ws = wb.active
    ws.append(["Город"])
    for c in ("Москва", "Тверь", "Казань"):
        ws.append([c])
    wb.save(cdir / "cities.xlsx")

    doc = DocxDocument()
    doc.add_paragraph("{{ФИО}} — {{Город}}")
    doc.save(folder / TPL)

    data = json.loads((folder / "project.stanok").read_text())
    data["templates"]["Договор"]["fields"] = {
        "ФИО": {"source": "table", "value": "ФИО", "counter_name": "data"},
        "Город": {"source": "table", "value": "Город", "counter_name": "cities"},
    }
    data["data_sources"].append({"file": "Данные/cities.xlsx", "mode": "sequential"})
    data["counters"]["cities"] = {
        "pos": 0, "format": "1", "table": "Данные/cities.xlsx", "resume": 1
    }
    _write_pj(folder, data)

    report = generate_documents(GenerateCommand(project_ref=folder), store=store)
    assert report.created == 2  # min(2, 3)
    texts = sorted(_doc_text(p) for p in report.output_paths)
    assert texts == ["Иван — Москва", "Петр — Тверь"]
    saved = _load_saved(store, "proj")
    assert saved.counters["data"].pos == 2
    assert saved.counters["cities"].pos == 2


def test_save_document_escape_rejected(tmp_path):
    from stanok.services.generate import _save_document

    with pytest.raises(TemplateError, match="escapes"):
        _save_document(DocxDocument(), tmp_path, "../evil.docx", 0)


def test_progress_callback_error_ignored(tmp_path, store):
    folder = make_project(tmp_path / "proj", rows=(("А", 1),))

    def bad_progress(created, total):
        raise RuntimeError("callback boom")

    report = generate_documents(
        GenerateCommand(project_ref=folder), store=store, progress=bad_progress
    )
    assert report.created == 1


def test_dist_escape_fails_fast(tmp_path, store):
    folder = make_project(tmp_path / "proj", rows=(("Иван", 100),))
    data = json.loads((folder / "project.stanok").read_text())
    data["filename_template"] = "../evil_{i}.docx"
    _write_pj(folder, data)
    with pytest.raises(ResolveError, match="dist"):
        generate_documents(GenerateCommand(project_ref=folder), store=store)
    assert not (folder.parent / "evil_1.docx").exists()


def test_dist_from_counter_and_today(tmp_path, store):
    import re

    folder = make_project(tmp_path / "proj", rows=(("Иван", 100),))
    data = json.loads((folder / "project.stanok").read_text())
    data["templates"]["Договор"]["fields"]["Дата"] = {"source": "today"}
    data["filename_template"] = "Д_{Номер}_{Дата}.docx"
    _write_pj(folder, data)
    report = generate_documents(GenerateCommand(project_ref=folder), store=store)
    assert report.created == 1
    name = report.output_paths[0].name
    assert re.fullmatch(r"Д_0_\d{4}-\d{2}-\d{2}\.docx", name), name


def test_progress_callback_cancel(tmp_path, store):
    folder = make_project(
        tmp_path / "proj", rows=(("А", 1), ("Б", 2), ("В", 3))
    )
    seen = []

    def progress(created, total):
        seen.append((created, total))
        return created >= 1

    report = generate_documents(
        GenerateCommand(project_ref=folder), store=store, progress=progress
    )
    assert report.created == 1
    assert seen[0] == (1, 3)
    saved = _load_saved(store, "proj")
    assert saved.counters["data"].pos == 1
    assert saved.counters["num"].pos == 1


def test_progress_callback_no_cancel(tmp_path, store):
    folder = make_project(tmp_path / "proj", rows=(("А", 1), ("Б", 2)))
    calls = []
    report = generate_documents(
        GenerateCommand(project_ref=folder),
        store=store,
        progress=lambda c, t: calls.append((c, t)) or False,
    )
    assert report.created == 2
    assert calls == [(1, 2), (2, 2)]


def test_dist_without_extension_gets_docx(tmp_path, store):
    folder = make_project(tmp_path / "proj", rows=(("Иван", 100),))
    data = json.loads((folder / "project.stanok").read_text())
    data["filename_template"] = "Договор_{ФИО}_{i}"
    _write_pj(folder, data)
    report = generate_documents(GenerateCommand(project_ref=folder), store=store)
    assert report.created == 1
    assert report.output_paths[0].suffix == ".docx"
    assert report.output_paths[0].name == "Договор_Иван_1.docx"
