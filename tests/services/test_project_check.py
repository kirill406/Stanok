# Copyright (C) 2026 Kirill Borovoy
# SPDX-License-Identifier: GPL-3.0-or-later
"""Structure check before project creation (024)."""

from docx import Document
from openpyxl import Workbook

import pytest

from stanok.gui.strings import STRINGS
from stanok.services.project_check import check_project_structure
from stanok.services.storage import ProjectStore


@pytest.fixture
def store(tmp_path):
    return ProjectStore(tmp_path / "home")


def _folder(tmp_path, name="proj"):
    root = tmp_path / name
    (root / "Данные").mkdir(parents=True)
    (root / "Шаблоны").mkdir()
    (root / "Результат").mkdir()
    return root


def _docx(path, text="Договор {{ФИО}} на {{Сумма}}"):
    doc = Document()
    doc.add_paragraph(text)
    doc.save(path)


def _xlsx(path, headers=("ФИО", "Сумма"), rows=(("Иван", 1),)):
    wb = Workbook()
    ws = wb.active
    if headers is not None:
        ws.append(list(headers))
    for row in rows:
        ws.append(list(row))
    wb.save(path)


def test_check_ok_structure(store, tmp_path):
    root = _folder(tmp_path)
    _docx(root / "Шаблоны" / "tpl.docx")
    _xlsx(root / "Данные" / "data.xlsx")
    check = check_project_structure(root, "новый", store)
    assert check.ok
    assert check.docx == ["Шаблоны/tpl.docx"]
    assert check.xlsx == ["Данные/data.xlsx"]
    assert check.missing_dirs == []


def test_check_missing_folder_is_error(store, tmp_path):
    root = tmp_path / "nope"
    check = check_project_structure(root, "x", store)
    assert not check.ok
    assert check.errors[0].file == ""


def test_check_missing_subfolders_listed(store, tmp_path):
    root = tmp_path / "proj"
    root.mkdir()
    check = check_project_structure(root, "x", store)
    assert not check.ok
    assert check.missing_dirs == ["Данные", "Шаблоны"]
    assert "Данные" in check.errors[0].message


def test_check_result_folder_not_required(store, tmp_path):
    root = _folder(tmp_path)
    _docx(root / "Шаблоны" / "tpl.docx")
    _xlsx(root / "Данные" / "data.xlsx")
    (root / "Результат").rmdir()
    assert check_project_structure(root, "x", store).ok


def test_check_no_docx_is_error(store, tmp_path):
    root = _folder(tmp_path)
    _xlsx(root / "Данные" / "data.xlsx")
    check = check_project_structure(root, "x", store)
    assert [i.message for i in check.errors] == [STRINGS.PCHK_NO_DOCX]


def test_check_no_xlsx_is_error(store, tmp_path):
    root = _folder(tmp_path)
    _docx(root / "Шаблоны" / "tpl.docx")
    check = check_project_structure(root, "x", store)
    assert not check.ok
    assert check.errors[0].file == "Данные"


def test_check_template_without_fields_is_error(store, tmp_path):
    root = _folder(tmp_path)
    _docx(root / "Шаблоны" / "plain.docx", text="Без полей")
    _xlsx(root / "Данные" / "data.xlsx")
    check = check_project_structure(root, "x", store)
    assert [i.file for i in check.errors] == ["Шаблоны/plain.docx"]


def test_check_broken_docx_is_error(store, tmp_path):
    root = _folder(tmp_path)
    (root / "Шаблоны" / "bad.docx").write_bytes(b"not a zip")
    _xlsx(root / "Данные" / "data.xlsx")
    check = check_project_structure(root, "x", store)
    assert [i.file for i in check.errors] == ["Шаблоны/bad.docx"]


def test_check_empty_field_is_warning(store, tmp_path):
    root = _folder(tmp_path)
    _docx(root / "Шаблоны" / "tpl.docx", text="Договор {{ФИО}} и {{ }}")
    _xlsx(root / "Данные" / "data.xlsx")
    check = check_project_structure(root, "x", store)
    assert check.ok
    assert [i.file for i in check.warnings] == ["Шаблоны/tpl.docx"]


def test_check_table_without_headers_is_error(store, tmp_path):
    root = _folder(tmp_path)
    _docx(root / "Шаблоны" / "tpl.docx")
    _xlsx(root / "Данные" / "empty.xlsx", headers=None, rows=())
    check = check_project_structure(root, "x", store)
    assert [i.file for i in check.errors] == ["Данные/empty.xlsx"]


def test_check_duplicate_headers_is_error(store, tmp_path):
    root = _folder(tmp_path)
    _docx(root / "Шаблоны" / "tpl.docx")
    _xlsx(root / "Данные" / "dup.xlsx", headers=("Имя", "Имя"), rows=())
    check = check_project_structure(root, "x", store)
    assert [i.file for i in check.errors] == ["Данные/dup.xlsx"]
    assert "Имя" in check.errors[0].message


def test_check_table_without_rows_is_warning(store, tmp_path):
    root = _folder(tmp_path)
    _docx(root / "Шаблоны" / "tpl.docx")
    _xlsx(root / "Данные" / "data.xlsx", rows=())
    check = check_project_structure(root, "x", store)
    assert check.ok
    assert [i.file for i in check.warnings] == ["Данные/data.xlsx"]


def test_check_does_not_match_fields_to_columns(store, tmp_path):
    root = _folder(tmp_path)
    _docx(root / "Шаблоны" / "tpl.docx", text="Город {{Город}}")
    _xlsx(root / "Данные" / "data.xlsx", headers=("ФИО",))
    check = check_project_structure(root, "x", store)
    assert check.ok
    assert check.warnings == []


def test_check_name_empty_is_error(store, tmp_path):
    root = _folder(tmp_path)
    _docx(root / "Шаблоны" / "tpl.docx")
    _xlsx(root / "Данные" / "data.xlsx")
    check = check_project_structure(root, "  ", store)
    assert not check.ok


@pytest.mark.parametrize("name", ["a/b", "a\\b", ".", ".."])
def test_check_name_with_bad_chars_is_error(store, tmp_path, name):
    root = _folder(tmp_path)
    _docx(root / "Шаблоны" / "tpl.docx")
    _xlsx(root / "Данные" / "data.xlsx")
    assert not check_project_structure(root, name, store).ok


def test_check_name_taken_is_error(store, tmp_path):
    root = _folder(tmp_path)
    _docx(root / "Шаблоны" / "tpl.docx")
    _xlsx(root / "Данные" / "data.xlsx")
    (store.home_dir / "занято.stanok").write_text("{}", encoding="utf-8")
    check = check_project_structure(root, "занято", store)
    assert not check.ok
    assert check.errors[0].file == ""
