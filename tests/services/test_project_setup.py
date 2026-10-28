# Copyright (C) 2026 Kirill Borovoy
# SPDX-License-Identifier: GPL-3.0-or-later
"""Folder preparation and creation from a passed check (024)."""

import pytest
from docx import Document
from openpyxl import Workbook

from stanok.engine.schema import StorageError
from stanok.services.project_check import check_project_structure
from stanok.services.project_setup import (
    clone_structure,
    create_project_from_check,
    create_skeleton,
    ensure_subfolders,
)
from stanok.services.storage import ProjectStore



@pytest.fixture
def store(tmp_path):
    return ProjectStore(tmp_path / "home")


def test_create_skeleton_makes_subfolders_without_config(store, tmp_path):
    folder = create_skeleton(tmp_path, "новый", store)
    assert folder == tmp_path / "новый"
    for sub in ("Данные", "Шаблоны", "Результат"):
        assert (folder / sub).is_dir()
    assert not store.config_exists("новый")


def test_create_skeleton_refuses_taken_name(store, tmp_path):
    (store.home_dir / "занято.stanok").write_text("{}", encoding="utf-8")
    with pytest.raises(StorageError, match="уже существует"):
        create_skeleton(tmp_path, "занято", store)
    assert not (tmp_path / "занято").exists()


def test_create_skeleton_refuses_non_empty_folder(store, tmp_path):
    busy = tmp_path / "busy"
    busy.mkdir()
    (busy / "чужой.txt").write_text("x", encoding="utf-8")
    with pytest.raises(StorageError, match="не пустая"):
        create_skeleton(tmp_path, "busy", store)
    assert (busy / "чужой.txt").exists()


def test_ensure_subfolders_creates_only_missing(store, tmp_path):
    folder = tmp_path / "p"
    (folder / "Данные").mkdir(parents=True)
    created = ensure_subfolders(folder)
    assert created == ["Шаблоны"]
    assert (folder / "Шаблоны").is_dir()


def _make_source(store, tmp_path, name="источник"):
    """A stored project created through the new flow."""
    folder = create_skeleton(tmp_path / "src", name, store)
    doc = Document()
    doc.add_paragraph("Договор {{ФИО}}")
    doc.save(folder / "Шаблоны" / "tpl.docx")
    wb = Workbook()
    wb.active.append(["ФИО"])
    wb.active.append(["Иван"])
    wb.save(folder / "Данные" / "data.xlsx")
    check = check_project_structure(folder, name, store)
    create_project_from_check(store, folder, name, check)
    return folder


def test_clone_structure_copies_templates_and_data(store, tmp_path):
    _make_source(store, tmp_path)
    folder = clone_structure("источник", tmp_path / "out", "копия", store)
    assert (folder / "Данные" / "data.xlsx").is_file()
    assert (folder / "Шаблоны" / "tpl.docx").is_file()
    assert (folder / "Результат").is_dir()
    assert not store.config_exists("копия")


def test_clone_structure_missing_source_raises(store, tmp_path):
    with pytest.raises(StorageError):
        clone_structure("нет", tmp_path, "копия", store)


def test_create_project_from_check_writes_config(store, tmp_path):
    folder = create_skeleton(tmp_path, "клиент", store)
    doc = Document()
    doc.add_paragraph("Договор {{ФИО}}")
    doc.save(folder / "Шаблоны" / "tpl.docx")
    wb = Workbook()
    ws = wb.active
    ws.append(["ФИО"])
    ws.append(["Иван"])
    wb.save(folder / "Данные" / "data.xlsx")

    check = check_project_structure(folder, "клиент", store)
    assert check.ok
    pj = create_project_from_check(store, folder, "клиент", check)
    assert store.config_exists("клиент")
    assert "tpl" in pj.templates
    assert pj.folder == str(folder.resolve())
    assert (folder / "Данные" / "data.xlsx").is_file()


def test_create_project_from_check_refuses_errors(store, tmp_path):
    folder = create_skeleton(tmp_path, "пусто", store)
    check = check_project_structure(folder, "пусто", store)
    assert not check.ok
    with pytest.raises(StorageError):
        create_project_from_check(store, folder, "пусто", check)
    assert not store.config_exists("пусто")
