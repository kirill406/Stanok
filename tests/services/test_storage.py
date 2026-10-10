# Copyright (C) 2026 Kirill Borovoy
# SPDX-License-Identifier: GPL-3.0-or-later
"""Tests for ProjectStore: atomic save/load, Home migration, recent projects."""

import json
import threading
from pathlib import Path

import pytest

from stanok.engine.schema import validate_pj
from stanok.services.storage import ProjectStore, resolve_project, StorageError


@pytest.fixture
def temp_home(tmp_path):
    """Create a temporary home directory for testing."""
    return tmp_path / ".stanok"


@pytest.fixture
def store(temp_home):
    """Create a ProjectStore with temporary home."""
    return ProjectStore(temp_home)


@pytest.fixture
def sample_pj():
    """Create a sample ProjectJSON for testing."""
    from stanok.engine.schema import validate_pj
    return validate_pj({
        "version": "0.0.0",
        "templates": {
            "template1": {
                "file": "templates/template1.docx",
                "fields": {
                    "field1": {"source": "constant", "value": "value1"},
                },
            }
        },
        "data_sources": [
            {"file": "data.xlsx", "mode": "sequential", "start_row": 0}
        ],
        "counters": {"counter1": {"last": 0, "format": "plain"}},
        "filename_template": "output_{template}_{i}.docx",
    })


def test_save_and_load(store, sample_pj):
    """Test save and load round-trip."""
    path = store.save(sample_pj)
    loaded = store.load(path.stem)
    assert loaded.version == sample_pj.version
    assert loaded.templates == sample_pj.templates


def test_atomic_write_on_error(store, tmp_path, monkeypatch):
    """Test that failed write leaves no .tmp and keeps original."""
    target = store._get_config_path("atom")
    target.write_text("{}")

    def boom(*a, **k):
        raise OSError("disk full")

    monkeypatch.setattr("stanok.services.storage.json.dump", boom)
    with pytest.raises(OSError, match="disk full"):
        store._atomic_write_json(target, {"test": "data"})
    assert target.read_text() == "{}"
    assert not target.with_suffix(target.suffix + ".tmp").exists()


def test_load_nonexistent(store):
    """Test loading non-existent project raises error."""
    with pytest.raises(StorageError, match="project not found"):
        store.load("nonexistent")


def test_delete_project(store, sample_pj):
    """Test deleting a project."""
    store.save(sample_pj)
    store.delete(sample_pj.filename_template.split("{")[0].strip() or "project")
    with pytest.raises(StorageError):
        store.load("project")


def test_list_projects(store, sample_pj):
    """Test listing projects."""
    store.save(sample_pj)
    projects = store.list()
    assert len(projects) == 1


def test_migrate_from_project(store, tmp_path):
    """Test migrating project config from project folder."""
    project_path = tmp_path / "legacy_proj"
    project_path.mkdir()
    legacy = project_path / "project.stanok"
    legacy.write_text(json.dumps({
        "version": "0.0.0",
        "templates": {"t1": {"file": "t.docx", "fields": {}}},
        "data_sources": [],
        "counters": {},
        "filename_template": "out.docx",
    }, indent=2))

    pj = store.migrate_from_project(project_path)
    assert pj.version == "0.0.0"
    # Saved under folder name, legacy removed
    assert (store.home_dir / "legacy_proj.stanok").exists()
    assert not legacy.exists()
    assert len(store.list()) == 1


def test_resolve_by_path(store, sample_pj, tmp_path):
    """Test resolving project by folder path."""
    # Create a project folder with legacy config
    project_dir = tmp_path / "my_project"
    project_dir.mkdir()
    legacy = project_dir / "project.stanok"
    legacy.write_text(json.dumps({
        "version": "0.0.0",
        "templates": {"t1": {"file": "t.docx", "fields": {}}},
        "data_sources": [],
        "counters": {},
        "filename_template": "out.docx",
    }))
    
    pj, path = resolve_project(project_dir, ProjectStore(tmp_path / ".stanok"))
    assert pj.version == "0.0.0"


def test_resolve_by_name(store, sample_pj):
    """Test resolving by config name in Home."""
    store.save(sample_pj)
    pj, path = resolve_project(sample_pj.filename_template.split("{")[0].strip() or "project", 
                                store)
    assert pj.version == "0.0.0"


def test_resolve_not_found(store):
    """Test resolving non-existent project raises error."""
    with pytest.raises(StorageError, match="not found"):
        resolve_project("nonexistent", store)


def test_recent_projects(store):
    """Test recent projects management."""
    # Add recent projects
    store.add_recent("folder1", "config1")
    store.add_recent("folder2", "config2")
    store.add_recent("folder1", "config1")  # duplicate
    
    recent = store.get_recent()
    assert len(store.get_recent()) == 2
    # Duplicate should be deduplicated, most recent first
    assert store.get_recent()[0].folder == "folder1"
    assert len(store.get_recent()) == 2


def test_recent_limit_10(store):
    """Test recent projects limit of 10."""
    for i in range(15):
        store.add_recent(f"folder{i}", f"config{i}")
    recent = store.get_recent()
    assert len(recent) == 10
    # Most recent should be first
    assert recent[0].config == "config14"


def test_recent_dedup(store):
    """Test deduplication of recent projects."""
    store.add_recent("f1", "c1")
    store.add_recent("f2", "c2")
    store.add_recent("f1", "c1")  # duplicate
    
    recent = store.get_recent()
    assert len(recent) == 2
    # Most recent should be the duplicate
    assert recent[0].config == "c1"


def test_settings_persistence(store):
    """Test that settings persist across store instances."""
    store.add_recent("f1", "c1")
    
    # Create new store instance with same home
    from stanok.services.storage import ProjectStore
    new_store = ProjectStore(store.home_dir)
    recent = new_store.get_recent()
    assert len(recent) == 1
    assert recent[0].config == "c1"


def test_path_validation_traversal(store):
    """Test that path traversal is rejected."""
    with pytest.raises(StorageError, match="traversal"):
        store._validate_path(Path("../../../etc/passwd"))


def test_absolute_path_rejected(store):
    """Test that absolute paths outside home are rejected."""
    with pytest.raises(StorageError, match="must be under home"):
        store._validate_path(Path("/etc/passwd"))


def test_empty_home_no_settings(store):
    """Test get_recent on empty home."""
    recent = store.get_recent()
    assert recent == []


def test_concurrent_save(store, sample_pj, tmp_path):
    """Test concurrent save operations."""
    import threading
    
    results = []
    def save_task():
        try:
            store.save(sample_pj)
            results.append("ok")
        except Exception as e:
            results.append(str(e))
    
    threads = [threading.Thread(target=save_task) for _ in range(5)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    
    # All should succeed
    assert all(r == "ok" for r in results)


def test_resolve_by_folder_name(store, sample_pj, tmp_path):
    """Test resolving by folder name in Home."""
    path = store.save(sample_pj)
    pj, resolved = resolve_project(path.stem, store)
    assert pj.version == "0.0.0"
    assert resolved == path


def test_resolve_by_config_name(store, sample_pj):
    """Test resolving by config name."""
    path = store.save(sample_pj)
    pj, resolved = resolve_project(path.stem, store)
    assert pj.version == "0.0.0"


def test_empty_project_name_fallback(store):
    """Test fallback when project name is empty."""
    pj = store._get_config_path("")
    assert pj.name == "project.stanok"


def test_atomic_write_creates_backup(store, sample_pj, tmp_path):
    """Test that atomic write creates .bak file."""
    path = store.save(sample_pj)
    bak_path = path.with_suffix(path.suffix + ".bak")
    # Backup should exist after second save
    store.save(sample_pj)
    assert bak_path.exists()


def test_save_updates_existing(store, sample_pj):
    """Test that save updates existing config."""
    path1 = store.save(sample_pj)
    path2 = store.save(sample_pj)
    assert path1 == path2


def test_load_validates_pj(store, tmp_path):
    """Test that load validates PJ structure."""
    # Create invalid config
    invalid_path = store.home_dir / "invalid.stanok"
    invalid_path.write_text(json.dumps({"version": "0.0.0"}))  # missing required fields

    with pytest.raises(StorageError):
        store.load("invalid")


def test_default_home_dir(monkeypatch, tmp_path):
    """Test default home is ~/.stanok."""
    monkeypatch.setattr(Path, "home", lambda: tmp_path)
    store = ProjectStore()
    assert store.home_dir == (tmp_path / ".stanok").resolve()
    assert (store.home_dir / "settings.json").exists()


def test_atomic_rename_fallback(monkeypatch, store, sample_pj):
    """Test Windows copy fallback when rename fails."""
    real_rename = Path.rename

    def boom_once(self, target):
        if self.suffix == ".tmp":
            raise OSError("locked")
        return real_rename(self, target)

    monkeypatch.setattr(Path, "rename", boom_once)
    path = store.save(sample_pj)
    assert path.exists()
    assert not path.with_suffix(path.suffix + ".tmp").exists()


def test_extract_name_variants(store, sample_pj):
    """Test project name extraction fallbacks."""
    assert store._extract_project_name(sample_pj) == "output_"
    raw = sample_pj.model_dump(mode="json")
    raw["filename_template"] = ""
    raw["templates"] = {"t": {"file": "Шаблоны/Договор.docx", "fields": {}}}
    pj = validate_pj(raw)
    assert store._extract_project_name(pj) == "Договор"
    pj2 = sample_pj.model_copy(update={"filename_template": "", "templates": {}})
    assert store._extract_project_name(pj2) == "project"


def test_load_reraise_storage(monkeypatch, store, sample_pj):
    """Test StorageError from read propagates unwrapped."""
    store.save(sample_pj)
    monkeypatch.setattr(
        store, "_read_json", lambda p: (_ for _ in ()).throw(StorageError("x", ["y"]))
    )
    with pytest.raises(StorageError):
        store.load("output_")


def test_delete_removes_bak(store, sample_pj):
    """Test delete removes config and backup."""
    path = store.save(sample_pj)
    store.save(sample_pj)
    bak = path.with_suffix(path.suffix + ".bak")
    assert bak.exists()
    store.delete(path.stem)
    assert not path.exists()
    assert not bak.exists()


def test_list_skips_invalid(store, sample_pj):
    """Test list skips invalid configs."""
    store.save(sample_pj)
    (store.home_dir / "broken.stanok").write_text("{not json")
    assert len(store.list()) == 1


def test_migrate_traversal_rejected(store, tmp_path):
    """Test migrate rejects traversal paths."""
    with pytest.raises(StorageError):
        store.migrate_from_project(Path("../outside"))


def test_migrate_missing_legacy(store, tmp_path):
    """Test migrate without legacy config."""
    d = tmp_path / "empty_proj"
    d.mkdir()
    with pytest.raises(StorageError, match="legacy config not found"):
        store.migrate_from_project(d)


def test_resolve_folder_fallback(store, sample_pj):
    """Test resolve via folder-name match in filename_template."""
    path = store.save(sample_pj)
    pj, resolved = resolve_project(path.stem, store)
    assert resolved == path


def test_get_recent_no_settings(store):
    """Test get_recent without settings file."""
    (store.home_dir / "settings.json").unlink()
    assert store.get_recent() == []

def _src_files(tmp_path, name="src"):
    """Build one xlsx + one docx with {{ФИО}}/{{Сумма}} placeholders."""
    from docx import Document as DocxDocument
    from openpyxl import Workbook

    src = tmp_path / name
    src.mkdir()
    wb = Workbook()
    ws = wb.active
    ws.append(["ФИО", "Сумма"])
    ws.append(["Иван", 100])
    xlsx = src / "data.xlsx"
    wb.save(xlsx)
    doc = DocxDocument()
    doc.add_paragraph("Договор {{ФИО}} на сумму {{Сумма}}")
    tpl = src / "tpl.docx"
    doc.save(tpl)
    return xlsx, tpl


def test_init_copies_and_builds_pj(store, tmp_path):
    xlsx, tpl = _src_files(tmp_path)
    pj = store.init_project(tmp_path / "proj", "nov", [xlsx], [tpl])
    assert (tmp_path / "proj" / "Данные" / "data.xlsx").is_file()
    assert (tmp_path / "proj" / "Шаблоны" / "tpl.docx").is_file()
    assert pj.templates["tpl"].file == "Шаблоны/tpl.docx"
    assert pj.templates["tpl"].fields["ФИО"].source == "table"
    assert pj.data_sources[0].file == "Данные/data.xlsx"
    assert pj.counters == {}
    assert (store.home_dir / "nov.stanok").is_file()
    assert any(r.config == "nov" for r in store.get_recent())


def test_init_no_copy_links_in_place(store, tmp_path):
    xlsx, tpl = _src_files(tmp_path)
    folder = tmp_path / "proj"
    (folder / "Данные").mkdir(parents=True)
    (folder / "Шаблоны").mkdir(parents=True)
    in_xlsx = folder / "Данные" / "data.xlsx"
    in_tpl = folder / "Шаблоны" / "tpl.docx"
    xlsx.rename(in_xlsx)
    tpl.rename(in_tpl)
    pj = store.init_project(folder, "loc", [in_xlsx], [in_tpl], copy_files=False)
    assert pj.data_sources[0].file == "Данные/data.xlsx"
    assert pj.templates["tpl"].file == "Шаблоны/tpl.docx"


def test_init_no_copy_outside_rejected(store, tmp_path):
    xlsx, tpl = _src_files(tmp_path)
    with pytest.raises(StorageError, match="outside project"):
        store.init_project(
            tmp_path / "proj", "out", [xlsx], [tpl], copy_files=False
        )


def test_init_name_taken_rejected(store, tmp_path, sample_pj):
    store.save(sample_pj, name="taken")
    xlsx, tpl = _src_files(tmp_path)
    with pytest.raises(StorageError, match="already exists"):
        store.init_project(tmp_path / "proj", "taken", [xlsx], [tpl])
    with pytest.raises(StorageError, match="invalid project name"):
        store.init_project(tmp_path / "proj", "  ", [xlsx], [tpl])


def test_init_copy_collision_uniquified(store, tmp_path):
    xlsx, tpl = _src_files(tmp_path, "a")
    xlsx_b, _ = _src_files(tmp_path, "b")
    pj = store.init_project(tmp_path / "proj", "dup", [xlsx, xlsx_b], [tpl])
    assert pj.data_sources[0].file == "Данные/data.xlsx"
    assert pj.data_sources[1].file == "Данные/data (1).xlsx"


def test_init_broken_template_aborts(store, tmp_path):
    xlsx, _ = _src_files(tmp_path)
    bad = tmp_path / "src" / "bad.docx"
    bad.write_text("not a zip")
    with pytest.raises(StorageError, match="cannot read template bad.docx"):
        store.init_project(tmp_path / "proj", "bad", [xlsx], [bad])


def test_init_placeholders_from_table(store, tmp_path):
    from docx import Document as DocxDocument

    xlsx, _ = _src_files(tmp_path)
    doc = DocxDocument()
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Сумма {{Итого}}"
    tpl = tmp_path / "src" / "tab.docx"
    doc.save(tpl)
    pj = store.init_project(tmp_path / "proj", "tab", [xlsx], [tpl])
    assert set(pj.templates["tab"].fields) == {"Итого"}


def test_init_second_copy_collision(store, tmp_path):
    xlsx, tpl = _src_files(tmp_path, "a")
    _src_files(tmp_path, "b")
    _src_files(tmp_path, "c")
    pj = store.init_project(
        tmp_path / "proj", "col",
        [xlsx, tmp_path / "b" / "data.xlsx", tmp_path / "c" / "data.xlsx"],
        [tpl],
    )
    assert [ds.file for ds in pj.data_sources] == [
        "Данные/data.xlsx",
        "Данные/data (1).xlsx",
        "Данные/data (2).xlsx",
    ]


def test_init_empty_lists_rejected(store, tmp_path):
    xlsx, tpl = _src_files(tmp_path)
    with pytest.raises(StorageError, match="no Excel"):
        store.init_project(tmp_path / "p1", "e1", [], [tpl])
    with pytest.raises(StorageError, match="no Word"):
        store.init_project(tmp_path / "p2", "e2", [xlsx], [])


def test_init_wrong_suffix_rejected(store, tmp_path):
    xlsx, _ = _src_files(tmp_path)
    bad = tmp_path / "src" / "notes.txt"
    bad.write_text("hi")
    with pytest.raises(StorageError, match="wrong file type"):
        store.init_project(tmp_path / "proj", "ws", [bad], [xlsx])


def test_init_missing_source_rejected(store, tmp_path):
    _, tpl = _src_files(tmp_path)
    with pytest.raises(StorageError, match="file not found"):
        store.init_project(
            tmp_path / "proj", "miss", [tmp_path / "src" / "gone.xlsx"], [tpl]
        )


def test_init_duplicate_template_stems(store, tmp_path):
    from docx import Document as DocxDocument

    folder = tmp_path / "proj"
    (folder / "Данные").mkdir(parents=True)
    xlsx, _ = _src_files(tmp_path)
    xlsx.rename(folder / "Данные" / "data.xlsx")
    for sub in ("a", "b"):
        d = folder / "Шаблоны" / sub
        d.mkdir(parents=True)
        doc = DocxDocument()
        doc.add_paragraph("Hi {{Имя}}")
        doc.save(d / "tpl.docx")
    pj = store.init_project(
        folder, "dst",
        [folder / "Данные" / "data.xlsx"],
        [folder / "Шаблоны" / "a" / "tpl.docx",
         folder / "Шаблоны" / "b" / "tpl.docx"],
        copy_files=False,
    )
    assert set(pj.templates) == {"tpl", "tpl (2)"}


def test_init_mkdir_failure(store, tmp_path, monkeypatch):
    import pathlib

    xlsx, tpl = _src_files(tmp_path)
    real_mkdir = pathlib.Path.mkdir

    def boom(self, *a, **k):
        raise OSError("denied")

    monkeypatch.setattr(pathlib.Path, "mkdir", boom)
    try:
        with pytest.raises(StorageError, match="cannot create folder"):
            store.init_project(tmp_path / "proj", "mk", [xlsx], [tpl])
    finally:
        monkeypatch.setattr(pathlib.Path, "mkdir", real_mkdir)


def test_init_copy_failure(store, tmp_path, monkeypatch):
    import shutil

    xlsx, tpl = _src_files(tmp_path)

    def boom(*a, **k):
        raise OSError("disk full")

    monkeypatch.setattr(shutil, "copy2", boom)
    with pytest.raises(StorageError, match="cannot copy"):
        store.init_project(tmp_path / "proj", "cp", [xlsx], [tpl])


def test_init_no_copy_wrong_subdir(store, tmp_path):
    xlsx, tpl = _src_files(tmp_path)
    folder = tmp_path / "proj"
    (folder / "Данные").mkdir(parents=True)
    xlsx.rename(folder / "Данные" / "data.xlsx")
    misplaced = folder / "tpl.docx"
    tpl.rename(misplaced)
    with pytest.raises(StorageError, match="must be in"):
        store.init_project(
            folder, "wd", [folder / "Данные" / "data.xlsx"], [misplaced],
            copy_files=False,
        )
    ghost = folder / "Шаблоны" / "ghost.docx"
    with pytest.raises(StorageError, match="file not found"):
        store.init_project(
            folder, "gh", [folder / "Данные" / "data.xlsx"], [ghost],
            copy_files=False,
        )
