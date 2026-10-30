# Copyright (C) 2026 Kirill Borovoy
# SPDX-License-Identifier: GPL-3.0-or-later
"""Text editing of a template .docx (027)."""

import pytest
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from stanok.services.template_edit import (
    TemplateDocument,
    TemplateEditError,
    is_editable,
    paragraph_text,
    set_paragraph_text,
)
from stanok.services.template_layout import Field, Text, build_layout


def _para(doc, *runs):
    """A paragraph with runs given as (text, {'bold': ..., 'italic': ...})."""
    para = doc.add_paragraph()
    for text, fmt in runs:
        run = para.add_run(text)
        run.bold = fmt.get("bold")
        run.italic = fmt.get("italic")
    return para


def _styles_of(doc, element):
    """Text pieces of the paragraph as {text: style} from a fresh layout."""
    layout = build_layout(doc)
    for block in layout.body:
        if getattr(block, "element", None) is element:
            return {p.text: p.style for p in block.pieces if isinstance(p, Text)}
    raise AssertionError("paragraph not in layout")


def test_template_edit_paragraph_text_marks_breaks_and_tabs():
    doc = Document()
    para = doc.add_paragraph("a")
    para.runs[0].add_break()
    para.add_run("b")
    para.runs[-1].add_tab()
    para.add_run("c")
    para.add_run().add_break(WD_BREAK.PAGE)
    para.add_run("d")

    assert paragraph_text(para._p) == "a\nb\t" + "c\fd"


def test_template_edit_insert_takes_formatting_of_left_neighbour():
    doc = Document()
    para = _para(doc, ("Привет ", {"bold": True}), ("мир", {"italic": True}))

    set_paragraph_text(para._p, "Привет, мир!")

    styles = _styles_of(doc, para._p)
    assert paragraph_text(para._p) == "Привет, мир!"
    bold_text = "".join(t for t, s in styles.items() if s.bold)
    italic_text = "".join(t for t, s in styles.items() if s.italic)
    assert "Привет" in bold_text
    assert "мир" in italic_text
    assert "!" in italic_text  # typed after the italic run: takes its formatting


def test_template_edit_delete_keeps_formatting_of_rest():
    doc = Document()
    para = _para(doc, ("Жирный", {"bold": True}), (" обычный", {}))

    set_paragraph_text(para._p, "Жирный")

    assert paragraph_text(para._p) == "Жирный"
    style = _styles_of(doc, para._p)["Жирный"]
    assert style.bold is True


def test_template_edit_keeps_field_placeholders_in_text():
    doc = Document()
    para = _para(doc, ("Договор {{ФИО}}", {}))

    set_paragraph_text(para._p, "Договор {{ФИО}} от {{Дата}}")

    layout = build_layout(doc)
    names = [p.name for p in layout.body[0].pieces if isinstance(p, Field)]
    assert names == ["ФИО", "Дата"]


def test_template_edit_line_break_is_written_as_break_element():
    doc = Document()
    para = _para(doc, ("одна", {}))

    set_paragraph_text(para._p, "одна\nдве")

    assert paragraph_text(para._p) == "одна\nдве"
    assert para._p.findall(".//" + qn("w:br")), "line break element expected"


def test_template_edit_page_break_is_written_as_page_break():
    doc = Document()
    para = _para(doc, ("a", {}))

    set_paragraph_text(para._p, "a\fb")

    breaks = para._p.findall(".//" + qn("w:br"))
    assert [b.get(qn("w:type")) for b in breaks] == ["page"]


def test_template_edit_unchanged_text_returns_false():
    doc = Document()
    para = _para(doc, ("тот же", {}))

    assert set_paragraph_text(para._p, "тот же") is False


def test_template_edit_keeps_picture_runs_in_place():
    doc = Document()
    para = _para(doc, ("текст", {}))
    picture_run = OxmlElement("w:r")
    picture_run.append(OxmlElement("w:drawing"))
    para._p.append(picture_run)

    set_paragraph_text(para._p, "другой")

    assert para._p.findall(".//" + qn("w:drawing"))
    assert paragraph_text(para._p) == "другой"


def test_template_edit_refuses_paragraph_with_hyperlink():
    doc = Document()
    para = _para(doc, ("ссылка", {}))
    para._p.append(OxmlElement("w:hyperlink"))

    assert is_editable(para._p) is False
    with pytest.raises(TemplateEditError):
        set_paragraph_text(para._p, "новый")


def test_template_document_save_writes_text_and_keeps_backup(tmp_path):
    path = tmp_path / "tpl.docx"
    doc = Document()
    doc.add_paragraph("старый текст")
    doc.save(path)
    template = TemplateDocument.load(path)
    element = template.layout().body[0].element

    assert template.set_paragraph_text(element, "новый {{ФИО}}") is True
    assert template.dirty is True
    template.save()

    assert template.dirty is False
    saved = Document(str(path))
    assert saved.paragraphs[0].text == "новый {{ФИО}}"
    backup = Document(str(tmp_path / "tpl.docx.bak"))
    assert backup.paragraphs[0].text == "старый текст"
    assert not (tmp_path / "tpl.docx.tmp").exists()


def test_template_document_save_failure_raises_and_keeps_file(tmp_path, monkeypatch):
    path = tmp_path / "tpl.docx"
    doc = Document()
    doc.add_paragraph("текст")
    doc.save(path)
    template = TemplateDocument.load(path)
    template.set_paragraph_text(template.layout().body[0].element, "правка")

    def boom(target):
        raise OSError("disk full")

    monkeypatch.setattr(template.doc, "save", boom)
    with pytest.raises(TemplateEditError):
        template.save()

    assert template.dirty is True
    assert Document(str(path)).paragraphs[0].text == "текст"
