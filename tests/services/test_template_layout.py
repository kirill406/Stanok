# Copyright (C) 2026 Kirill Borovoy
# SPDX-License-Identifier: GPL-3.0-or-later
"""Layout model of a docx template (027)."""

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt
from lxml import etree

from stanok.services.template_layout import (
    Field,
    LineBreak,
    PageBreak,
    Paragraph,
    Table,
    Text,
    TemplateLayoutError,
    build_layout,
    load_template_layout,
)


WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _doc(*texts: str):
    doc = Document()
    for text in texts:
        doc.add_paragraph(text)
    return doc


def _layout(doc, tmp_path):
    path = tmp_path / "tpl.docx"
    doc.save(path)
    return load_template_layout(path)


def test_template_layout_plain_paragraph_has_text_piece(tmp_path):
    layout = _layout(_doc("Договор"), tmp_path)

    para = layout.body[0]
    assert isinstance(para, Paragraph)
    assert [p.text for p in para.pieces if isinstance(p, Text)] == ["Договор"]


def test_template_layout_placeholder_becomes_field_piece(tmp_path):
    layout = _layout(_doc("Договор {{ ФИО }} на {{Сумма}}."), tmp_path)

    names = [p.name for p in layout.body[0].pieces if isinstance(p, Field)]
    assert names == ["ФИО", "Сумма"]
    texts = [p.text for p in layout.body[0].pieces if isinstance(p, Text)]
    assert texts == ["Договор ", " на ", "."]


def test_template_layout_placeholder_split_across_runs_is_one_field(tmp_path):
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("Фио: {{Фа")
    para.add_run("милия}} конец")
    layout = _layout(doc, tmp_path)

    fields = [p for p in layout.body[0].pieces if isinstance(p, Field)]
    assert [f.name for f in fields] == ["Фамилия"]


def test_template_layout_alignment_center_and_right_are_kept(tmp_path):
    doc = Document()
    doc.add_paragraph("a").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("b").alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph("c")
    layout = _layout(doc, tmp_path)

    assert [p.align for p in layout.body] == ["center", "right", "left"]


def test_template_layout_justify_is_marked_justify(tmp_path):
    doc = Document()
    doc.add_paragraph("длинный текст").alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    layout = _layout(doc, tmp_path)

    assert layout.body[0].align == "justify"


def test_template_layout_indents_and_spacing_are_points(tmp_path):
    doc = Document()
    para = doc.add_paragraph("текст")
    para.paragraph_format.left_indent = Pt(36)
    para.paragraph_format.right_indent = Pt(12)
    para.paragraph_format.first_line_indent = Pt(18)
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after = Pt(10)
    layout = _layout(doc, tmp_path)

    p = layout.body[0]
    assert (p.left_pt, p.right_pt, p.first_pt) == (36, 12, 18)
    assert (p.before_pt, p.after_pt) == (6, 10)


def test_template_layout_run_formatting_is_kept(tmp_path):
    doc = Document()
    para = doc.add_paragraph()
    bold = para.add_run("жирный")
    bold.bold = True
    italic = para.add_run("курсив")
    italic.italic = True
    underlined = para.add_run("подчёркнут")
    underlined.underline = True
    big = para.add_run("крупно")
    big.font.size = Pt(16)
    layout = _layout(doc, tmp_path)

    styles = {p.text: p.style for p in layout.body[0].pieces if isinstance(p, Text)}
    assert styles["жирный"].bold and not styles["жирный"].italic
    assert styles["курсив"].italic
    assert styles["подчёркнут"].underline
    assert styles["крупно"].size_pt == 16


def test_template_layout_style_size_is_inherited_by_runs(tmp_path):
    doc = Document()
    doc.add_paragraph("Раздел", style="Heading 1")
    layout = _layout(doc, tmp_path)

    style = layout.body[0].pieces[0].style
    assert style.bold is True
    assert style.size_pt == 14  # Heading 1 of the default template is 14 pt


def test_template_layout_line_break_is_line_piece(tmp_path):
    doc = Document()
    para = doc.add_paragraph("первая")
    para.runs[0].add_break()
    para.add_run("вторая")
    layout = _layout(doc, tmp_path)

    kinds = [type(p) for p in layout.body[0].pieces]
    assert kinds == [Text, LineBreak, Text]


def test_template_layout_page_break_is_page_piece(tmp_path):
    doc = Document()
    para = doc.add_paragraph("до")
    para.runs[0].add_break(WD_BREAK.PAGE)
    para.add_run("после")
    layout = _layout(doc, tmp_path)

    assert any(isinstance(p, PageBreak) for p in layout.body[0].pieces)


def test_template_layout_page_break_before_is_flag(tmp_path):
    doc = Document()
    doc.add_paragraph("обычный")
    doc.add_paragraph("с новой страницы").paragraph_format.page_break_before = True
    layout = _layout(doc, tmp_path)

    assert [p.page_break_before for p in layout.body] == [False, True]


def test_template_layout_tab_stays_in_text(tmp_path):
    doc = Document()
    para = doc.add_paragraph("а")
    para.runs[0].add_tab()
    para.add_run("б")
    layout = _layout(doc, tmp_path)

    text = "".join(p.text for p in layout.body[0].pieces if isinstance(p, Text))
    assert text == "а\tб"


def test_template_layout_empty_paragraph_has_no_pieces(tmp_path):
    doc = Document()
    doc.add_paragraph()
    layout = _layout(doc, tmp_path)

    assert layout.body[0].pieces == []


def test_template_layout_table_keeps_widths_spans_and_borders(tmp_path):
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.style = doc.styles["Table Grid"]
    table.cell(0, 0).paragraphs[0].text = "{{ФИО}}"
    merged = table.cell(1, 0).merge(table.cell(1, 1))
    merged.paragraphs[0].text = "итог"
    table.columns[0].width = Pt(100)
    table.columns[1].width = Pt(200)
    layout = _layout(doc, tmp_path)

    block = layout.body[0]
    assert isinstance(block, Table)
    assert block.borders is True
    assert block.col_widths_pt == pytest.approx([100, 200])
    assert [c.span for c in block.rows[1]] == [2]
    field_names = [
        piece.name
        for piece in block.rows[0][0].blocks[0].pieces
        if isinstance(piece, Field)
    ]
    assert field_names == ["ФИО"]


def test_template_layout_table_without_borders_is_flagged(tmp_path):
    doc = Document()
    doc.add_table(rows=1, cols=1)
    layout = _layout(doc, tmp_path)

    assert layout.body[0].borders is False


def test_template_layout_body_keeps_paragraph_and_table_order(tmp_path):
    doc = Document()
    doc.add_paragraph("до таблицы")
    doc.add_table(rows=1, cols=1)
    doc.add_paragraph("после таблицы")
    layout = _layout(doc, tmp_path)

    assert [type(b) for b in layout.body] == [Paragraph, Table, Paragraph]


def test_template_layout_header_and_footer_are_separate_blocks(tmp_path):
    doc = Document()
    section = doc.sections[0]
    section.header.paragraphs[0].text = "Шапка {{Организация}}"
    section.footer.paragraphs[0].text = "Подвал"
    doc.add_paragraph("тело")
    layout = _layout(doc, tmp_path)

    header_fields = [
        p.name for p in layout.header[0].pieces if isinstance(p, Field)
    ]
    assert header_fields == ["Организация"]
    assert layout.footer[0].pieces[0].text == "Подвал"
    assert layout.body[0].pieces[0].text == "тело"


def test_template_layout_page_size_and_margins_come_from_section(tmp_path):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Pt(400)
    section.page_height = Pt(600)
    section.left_margin = Pt(20)
    section.right_margin = Pt(30)
    section.top_margin = Pt(40)
    section.bottom_margin = Pt(50)
    layout = _layout(doc, tmp_path)

    assert (layout.width_pt, layout.height_pt) == (400, 600)
    assert layout.margins_pt == (20, 30, 40, 50)


def test_template_layout_placeholders_are_unique_and_in_document_order(tmp_path):
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "{{Шапка}}"
    doc.add_paragraph("{{Б}} и {{А}}")
    doc.add_paragraph("снова {{Б}}")
    layout = _layout(doc, tmp_path)

    assert layout.placeholders() == ["Шапка", "Б", "А"]


def test_template_layout_unreadable_file_raises_layout_error(tmp_path):
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a docx")

    with pytest.raises(TemplateLayoutError):
        load_template_layout(path)


def test_template_layout_build_layout_reads_open_document(tmp_path):
    doc = Document()
    doc.add_paragraph("строка")
    layout = build_layout(doc)

    assert layout.body[0].pieces[0].text == "строка"


def test_template_layout_default_size_comes_from_doc_defaults(tmp_path):
    doc = Document()
    doc.styles.element.insert(0, etree.fromstring(
        f'<w:docDefaults xmlns:w="{WORD_NS}">'
        '<w:rPrDefault><w:rPr><w:sz w:val="28"/></w:rPr></w:rPrDefault>'
        "</w:docDefaults>"
    ))
    doc.add_paragraph("без размера")
    layout = _layout(doc, tmp_path)

    assert layout.body[0].pieces[0].style.size_pt == 14
