# -*- coding: utf-8 -*-
"""Regression tests for B1 «Генерируемые проекты» (002-stabilization).

Covers: Home snapshots named after the project with (1),(2),… collision
rule (a), prefilled documents in generated projects (б), and the
«Поля шаблона для генерируемых проектов» section contract (в).

Run with: python -m pytest tests/test_002_projects.py -q
"""

import json
import logging
import os

import pytest

from docxforge import generate as gen_module
from docxforge.engine.schema import (
    BatchSourceConfig,
    FieldMapping,
    FieldType,
    Project,
    RowIterationMode,
    TemplateConfig,
)


logger = logging.getLogger(__name__)


def _make_flat_source(tmp_dir: str) -> str:
    """Minimal flat source: 1 template, 1 SEQUENTIAL table, 3 rows."""
    from docx import Document
    from openpyxl import Workbook

    project_dir = os.path.join(tmp_dir, 'source_project')
    data_dir = os.path.join(project_dir, 'Данные')
    tmpl_dir = os.path.join(project_dir, 'Шаблоны')
    os.makedirs(data_dir, exist_ok=True)
    os.makedirs(tmpl_dir, exist_ok=True)

    doc = Document()
    doc.add_paragraph('Клиент: {{ client_name }}')
    doc.add_paragraph('Организация: {{ org }}')
    doc.save(os.path.join(tmpl_dir, 'contract.docx'))

    wb = Workbook()
    ws = wb.active
    ws.append(['client_name'])
    ws.append(['ООО Альфа'])
    ws.append(['ООО Бета'])
    ws.append(['ИП Гамма'])
    wb.save(os.path.join(data_dir, 'clients.xlsx'))

    prj = Project()
    tc = TemplateConfig()
    tc.fields['client_name'] = FieldMapping(
        type=FieldType.TABLE, file='clients.xlsx', column='client_name')
    tc.fields['org'] = FieldMapping(type=FieldType.CONSTANT, value='ORG_VALUE')
    tc.batch_sources = {
        'clients.xlsx': BatchSourceConfig(
            file='clients.xlsx', mode=RowIterationMode.SEQUENTIAL),
    }
    tc.filename_template = '{{ client_name }}_договор.docx'
    prj.templates['contract.docx'] = tc
    prj.to_file(os.path.join(project_dir, 'проект.docxforge'))
    return project_dir


def _make_nested_source(tmp_dir: str) -> str:
    """Minimal nested source: employee + project_name batch columns."""
    from docx import Document
    from openpyxl import Workbook

    project_dir = os.path.join(tmp_dir, 'nested_source')
    data_dir = os.path.join(project_dir, 'Данные')
    tmpl_dir = os.path.join(project_dir, 'Шаблоны')
    os.makedirs(data_dir, exist_ok=True)
    os.makedirs(tmpl_dir, exist_ok=True)

    doc = Document()
    doc.add_paragraph('Сотрудник: {{ employee }}')
    doc.add_paragraph('Проект: {{ project_name }}')
    doc.add_paragraph('Организация: {{ org }}')
    doc.save(os.path.join(tmpl_dir, 'contract.docx'))

    wb = Workbook()
    ws = wb.active
    ws.append(['employee', 'project_name'])
    ws.append(['Ivanov_Ivan', 'Dogovor_001'])
    ws.append(['Petrov_Petr', 'Dogovor_002'])
    wb.save(os.path.join(data_dir, 'work.xlsx'))

    prj = Project()
    tc = TemplateConfig()
    tc.fields['employee'] = FieldMapping(
        type=FieldType.TABLE, file='work.xlsx', column='employee')
    tc.fields['project_name'] = FieldMapping(
        type=FieldType.TABLE, file='work.xlsx', column='project_name')
    tc.fields['org'] = FieldMapping(type=FieldType.CONSTANT, value='ORG_VALUE')
    tc.batch_sources = {
        'work.xlsx': BatchSourceConfig(
            file='work.xlsx', mode=RowIterationMode.SEQUENTIAL),
    }
    tc.filename_template = '{{ project_name }}.docx'
    prj.templates['contract.docx'] = tc
    prj.to_file(os.path.join(project_dir, 'проект.docxforge'))
    return project_dir


def _read_docx_text(path: str) -> str:
    from docx import Document

    doc = Document(path)
    return '\n'.join(p.text for p in doc.paragraphs)


# --- (a) Home snapshots -----------------------------------------------------

def test_002_unique_home_project_file_no_collision(tmp_path):
    """No collision → plain `<name>.docxforge` in Home."""
    home = str(tmp_path / 'home')
    os.makedirs(home, exist_ok=True)
    path = gen_module.unique_home_project_file('Договор_001', home_dir=home)
    assert path == os.path.join(home, 'Договор_001.docxforge')
    assert not os.path.exists(path)


def test_002_unique_home_project_file_collision_renames_new(tmp_path):
    """Collision → `name (1)`, `name (2)`; existing file untouched."""
    home = str(tmp_path / 'home')
    os.makedirs(home, exist_ok=True)
    existing = os.path.join(home, 'Договор_001.docxforge')
    with open(existing, 'w', encoding='utf-8') as f:
        f.write('SENTINEL')
    first = gen_module.unique_home_project_file('Договор_001', home_dir=home)
    assert first == os.path.join(home, 'Договор_001 (1).docxforge')
    with open(first, 'w', encoding='utf-8') as f:
        f.write('NEW')
    second = gen_module.unique_home_project_file('Договор_001', home_dir=home)
    assert second == os.path.join(home, 'Договор_001 (2).docxforge')
    with open(existing, encoding='utf-8') as f:
        assert f.read() == 'SENTINEL'


def test_002_flat_home_snapshots_named_after_project(tmp_path):
    """Creation writes `<folder>.docxforge` in folder; open migrates to Home."""
    source_dir = _make_flat_source(str(tmp_path))
    home = str(tmp_path / 'home')
    os.makedirs(home, exist_ok=True)
    projects_dir, count = gen_module.create_projects_from_template(
        source_dir, 'contract.docx', '{{client_name}}', home_dir=home)
    assert count == 3
    folders = sorted(
        d for d in os.listdir(projects_dir)
        if os.path.isdir(os.path.join(projects_dir, d)))
    assert len(folders) == 3
    assert os.listdir(home) == []
    for folder in folders:
        folder_file = os.path.join(projects_dir, folder, folder + '.docxforge')
        assert os.path.isfile(folder_file), folder_file
        assert not os.path.exists(
            os.path.join(projects_dir, folder, 'проект.docxforge'))
        live = Project.from_file(folder_file)
        assert 'contract.docx' in live.templates
        migrated = gen_module.migrate_project_configs_to_home(
            os.path.join(projects_dir, folder), home)
        assert migrated == [os.path.join(home, folder + '.docxforge')]
        assert not os.path.exists(folder_file)
        snap_project = Project.from_file(migrated[0])
        assert (snap_project.templates['contract.docx'].fields['client_name'].value
                == live.templates['contract.docx'].fields['client_name'].value)


def test_002_flat_home_snapshot_collision_keeps_existing(tmp_path):
    """Pre-existing Home file keeps bytes; migrated copy gets (1)."""
    source_dir = _make_flat_source(str(tmp_path))
    home = str(tmp_path / 'home')
    os.makedirs(home, exist_ok=True)
    preoccupied = os.path.join(home, 'ООО Альфа.docxforge')
    with open(preoccupied, 'wb') as f:
        f.write(b'SENTINEL-BYTES')
    projects_dir, _count = gen_module.create_projects_from_template(
        source_dir, 'contract.docx', '{{client_name}}',
        max_projects=1, home_dir=home)
    folder = os.listdir(projects_dir)[0]
    migrated = gen_module.migrate_project_configs_to_home(
        os.path.join(projects_dir, folder), home)
    with open(preoccupied, 'rb') as f:
        assert f.read() == b'SENTINEL-BYTES'
    renamed = os.path.join(home, 'ООО Альфа (1).docxforge')
    assert migrated == [renamed]
    assert os.path.isfile(renamed)
    snap_project = Project.from_file(renamed)
    assert 'contract.docx' in snap_project.templates


def test_002_nested_home_snapshots_named_after_project(tmp_path):
    """Nested creation writes named configs; open migrates them to Home."""
    source_dir = _make_nested_source(str(tmp_path))
    home = str(tmp_path / 'home')
    os.makedirs(home, exist_ok=True)
    projects_dir, n_employees, n_projects = (
        gen_module.create_nested_employee_projects(
            source_dir, 'contract.docx', '{{employee}}/{{project_name}}',
            home_dir=home))
    assert (n_employees, n_projects) == (2, 2)
    assert os.listdir(home) == []
    for emp in sorted(os.listdir(projects_dir)):
        emp_dir = os.path.join(projects_dir, emp)
        if not os.path.isdir(emp_dir):
            continue
        for proj in sorted(os.listdir(emp_dir)):
            pdir = os.path.join(emp_dir, proj)
            if not os.path.isdir(pdir):
                continue
            assert os.path.isfile(os.path.join(pdir, proj + '.docxforge'))
            migrated = gen_module.migrate_project_configs_to_home(pdir, home)
            assert migrated == [os.path.join(home, proj + '.docxforge')]
    snapshots = sorted(
        f for f in os.listdir(home) if f.endswith('.docxforge'))
    assert snapshots == ['Dogovor_001.docxforge', 'Dogovor_002.docxforge']
    for snap_name in snapshots:
        snap_project = Project.from_file(os.path.join(home, snap_name))
        assert 'contract.docx' in snap_project.templates


# --- (б) Prefilled documents -------------------------------------------------

def test_002_flat_generated_projects_prefilled(tmp_path):
    """Each flat project renders ≥1 doc with row value + constant, no {{ }}."""
    source_dir = _make_flat_source(str(tmp_path))
    home = str(tmp_path / 'home')
    os.makedirs(home, exist_ok=True)
    projects_dir, count = gen_module.create_projects_from_template(
        source_dir, 'contract.docx', '{{client_name}}', home_dir=home)
    assert count == 3
    seen_clients = set()
    for folder in sorted(os.listdir(projects_dir)):
        result_dir = os.path.join(projects_dir, folder, 'Результат')
        docs = [f for f in os.listdir(result_dir) if f.endswith('.docx')]
        assert len(docs) >= 1, folder
        text = _read_docx_text(os.path.join(result_dir, docs[0]))
        assert '{{' not in text and '}}' not in text, text
        assert 'ORG_VALUE' in text
        cfg = Project.from_file(
            gen_module.resolve_project_file(os.path.join(projects_dir, folder))
        ).templates['contract.docx']
        client = cfg.fields['client_name'].value
        assert client in text
        seen_clients.add(client)
    assert seen_clients == {'ООО Альфа', 'ООО Бета', 'ИП Гамма'}


def test_002_nested_generated_projects_prefilled(tmp_path):
    """Each nested project renders ≥1 doc with constant, no {{ }}."""
    source_dir = _make_nested_source(str(tmp_path))
    home = str(tmp_path / 'home')
    os.makedirs(home, exist_ok=True)
    projects_dir, n_employees, n_projects = (
        gen_module.create_nested_employee_projects(
            source_dir, 'contract.docx', '{{employee}}/{{project_name}}',
            home_dir=home))
    assert (n_employees, n_projects) == (2, 2)
    checked = 0
    for emp in sorted(os.listdir(projects_dir)):
        emp_dir = os.path.join(projects_dir, emp)
        if not os.path.isdir(emp_dir):
            continue
        for proj in sorted(os.listdir(emp_dir)):
            pdir = os.path.join(emp_dir, proj)
            if not os.path.isdir(pdir):
                continue
            result_dir = os.path.join(pdir, 'Результат')
            docs = [f for f in os.listdir(result_dir) if f.endswith('.docx')]
            assert len(docs) >= 1, pdir
            text = _read_docx_text(os.path.join(result_dir, docs[0]))
            assert '{{' not in text and '}}' not in text, text
            assert 'ORG_VALUE' in text
            checked += 1
    assert checked == 2


# --- (в) Generated-project template fields ------------------------------------

def test_002_generated_fields_schema_round_trip(tmp_path):
    """generated_project_fields persist; old files default to [] (= all)."""
    prj = Project()
    tc = TemplateConfig()
    tc.fields['a'] = FieldMapping(type=FieldType.CONSTANT, value='1')
    tc.fields['b'] = FieldMapping(type=FieldType.CONSTANT, value='2')
    tc.generated_project_fields = ['a']
    prj.templates['t.docx'] = tc
    path = os.path.join(str(tmp_path), 'проект.docxforge')
    prj.to_file(path)
    loaded = Project.from_file(path)
    assert loaded.templates['t.docx'].generated_project_fields == ['a']
    with open(path, encoding='utf-8') as f:
        raw = json.load(f)
    assert raw['templates']['t.docx']['generated_project_fields'] == ['a']

    prj2 = Project()
    tc2 = TemplateConfig()
    tc2.fields['a'] = FieldMapping(type=FieldType.CONSTANT, value='1')
    prj2.templates['t.docx'] = tc2
    path2 = os.path.join(str(tmp_path), 'old.docxforge')
    prj2.to_file(path2)
    with open(path2, encoding='utf-8') as f:
        raw2 = json.load(f)
    assert 'generated_project_fields' not in raw2['templates']['t.docx']
    assert Project.from_file(path2).templates['t.docx'].generated_project_fields == []


def test_002_build_project_config_respects_include_fields():
    """Non-empty include_fields → only subset in generated config."""
    tc = TemplateConfig()
    tc.fields['keep'] = FieldMapping(
        type=FieldType.TABLE, file='b.xlsx', column='keep')
    tc.fields['drop'] = FieldMapping(
        type=FieldType.TABLE, file='b.xlsx', column='drop')
    tc.fields['const'] = FieldMapping(type=FieldType.CONSTANT, value='C')
    row = {'keep': 'K', 'drop': 'D'}
    full = gen_module._build_project_config(tc, row, 'b.xlsx')
    assert set(full.fields) == {'keep', 'drop', 'const'}
    subset = gen_module._build_project_config(
        tc, row, 'b.xlsx', include_fields=['keep'])
    assert set(subset.fields) == {'keep'}
    assert subset.fields['keep'].type == FieldType.CONSTANT
    assert subset.fields['keep'].value == 'K'
    empty_means_all = gen_module._build_project_config(
        tc, row, 'b.xlsx', include_fields=[])
    assert set(empty_means_all.fields) == {'keep', 'drop', 'const'}


def test_002_flat_creation_applies_stored_subset(tmp_path):
    """Stored generated_project_fields subset flows into created projects."""
    source_dir = _make_flat_source(str(tmp_path))
    home = str(tmp_path / 'home')
    os.makedirs(home, exist_ok=True)
    prj = Project.from_file(os.path.join(source_dir, 'проект.docxforge'))
    prj.templates['contract.docx'].generated_project_fields = ['client_name']
    prj.to_file(os.path.join(source_dir, 'проект.docxforge'))
    projects_dir, count = gen_module.create_projects_from_template(
        source_dir, 'contract.docx', '{{client_name}}', home_dir=home)
    assert count == 3
    for folder in sorted(os.listdir(projects_dir)):
        cfg = Project.from_file(
            gen_module.resolve_project_file(os.path.join(projects_dir, folder))
        ).templates['contract.docx']
        assert set(cfg.fields) == {'client_name'}
        assert cfg.fields['client_name'].type == FieldType.CONSTANT


@pytest.mark.gui
class TestGeneratedFieldsSection:
    """FillForm section «Поля шаблона для генерируемых проектов»."""

    def test_002_section_present_all_checked_by_default(
            self, qtbot, sample_project):
        """Section exists, one checked box per field, STRINGS title."""
        from PyQt5.QtWidgets import QGroupBox

        from docxforge.gui.fill_form import FillForm
        from docxforge.gui.strings import STRINGS

        dlg = FillForm(sample_project, 'all_fields.docx')
        qtbot.addWidget(dlg)
        dlg.show()

        groups = [g for g in dlg.findChildren(QGroupBox)
                  if g.title() == STRINGS['fill_generated_fields_section']]
        assert len(groups) == 1
        assert set(dlg.generated_field_checks) == set(dlg.field_widgets)
        assert all(box.isChecked() for box in dlg.generated_field_checks.values())
        assert dlg._collect_config().generated_project_fields == []

    def test_002_unchecked_field_excluded_from_collected_config(
            self, qtbot, sample_project):
        """Unchecking one field stores the rest as the subset."""
        from docxforge.gui.fill_form import FillForm

        dlg = FillForm(sample_project, 'all_fields.docx')
        qtbot.addWidget(dlg)
        dlg.show()

        names = sorted(dlg.generated_field_checks)
        assert len(names) >= 2
        dlg.generated_field_checks[names[0]].setChecked(False)
        collected = dlg._collect_config()
        assert set(collected.generated_project_fields) == set(names[1:])
        assert names[0] not in collected.generated_project_fields

        for box in dlg.generated_field_checks.values():
            box.setChecked(True)
        assert dlg._collect_config().generated_project_fields == []
