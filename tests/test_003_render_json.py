# -*- coding: utf-8 -*-
"""Tests for Renderer.render_from_json (003-json, Phase 1, P1).

Covers: Filling JSON validation, Excel-free render from JSON,
line-break preservation, and table-cell rendering.
"""

import os
import shutil
import zipfile

import pytest
from docx import Document
from lxml import etree

from docxforge.engine.data_reader import DataReader
from docxforge.engine.renderer import Renderer

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def _make_work(tmp_path, template_name='template.docx'):
    """Create project work dir with Шаблоны/ and return (work, template)."""
    work = str(tmp_path / 'work')
    templates_dir = os.path.join(work, 'Шаблоны')
    os.makedirs(templates_dir, exist_ok=True)
    return work, os.path.join(templates_dir, template_name)


def _render(work, filling, **kwargs):
    renderer = Renderer(work, DataReader())
    renderer.load_project()
    return renderer.render_from_json(filling, **kwargs)


def _docx_paragraph_texts(path):
    doc = Document(path)
    return [p.text for p in doc.paragraphs]


class TestRenderFromJsonValidation:
    def test_render_from_json_not_dict_raises(self, tmp_path):
        work, _ = _make_work(tmp_path)
        for bad in (None, [], 'template.docx', 42):
            with pytest.raises(ValueError):
                _render(work, bad, output_dir=str(tmp_path / 'out'))

    def test_render_from_json_missing_template_raises(self, tmp_path):
        work, _ = _make_work(tmp_path)
        for bad in ({}, {'fields': {}}, {'fields': {'a': 'b'}},
                    {'template': '', 'fields': {}},
                    {'template': '   ', 'fields': {}},
                    {'template': None, 'fields': {}},
                    {'template': 42, 'fields': {}}):
            with pytest.raises(ValueError):
                _render(work, bad, output_dir=str(tmp_path / 'out'))

    def test_render_from_json_missing_fields_raises(self, tmp_path):
        work, _ = _make_work(tmp_path)
        with pytest.raises(ValueError):
            _render(work, {'template': 'template.docx'},
                    output_dir=str(tmp_path / 'out'))

    def test_render_from_json_fields_not_dict_raises(self, tmp_path):
        work, _ = _make_work(tmp_path)
        for bad_fields in (None, [], 'value', 42):
            with pytest.raises(ValueError):
                _render(work, {'template': 'template.docx',
                               'fields': bad_fields},
                        output_dir=str(tmp_path / 'out'))


class TestRenderFromJsonRender:
    def test_render_from_json_basic_no_excel(self, tmp_path):
        """Filling values render as constants; no Excel files involved."""
        work, template_path = _make_work(tmp_path)
        doc = Document()
        doc.add_paragraph('Client: {{ client }}, Doc #{{ doc_number }}')
        doc.save(template_path)

        filling = {'template': 'template.docx',
                   'fields': {'client': 'OOO Romashka', 'doc_number': 5}}
        outputs = _render(work, filling, output_dir=str(tmp_path / 'out'))

        assert len(outputs) == 1
        assert os.path.exists(outputs[0])
        assert os.listdir(os.path.join(work, 'Шаблоны')) == ['template.docx']
        assert _docx_paragraph_texts(outputs[0]) == [
            'Client: OOO Romashka, Doc #5']

    def test_render_from_json_dist_dir_default(self, tmp_path):
        """Without output_dir, dist dirname is used under project_dir."""
        work, template_path = _make_work(tmp_path)
        doc = Document()
        doc.add_paragraph('Hello {{ name }}')
        doc.save(template_path)

        filling = {'template': 'template.docx',
                   'dist': 'result_dir/out.docx',
                   'fields': {'name': 'Ivan'}}
        outputs = _render(work, filling)

        assert len(outputs) == 1
        assert os.path.dirname(outputs[0]) == os.path.join(work, 'result_dir')
        assert _docx_paragraph_texts(outputs[0]) == ['Hello Ivan']

    def test_render_from_json_linebreak_preserved(self, tmp_path):
        """w:br before a placeholder stays before the substituted value."""
        work, template_path = _make_work(tmp_path)
        doc = Document()
        p = doc.add_paragraph()
        p.add_run('Line1')
        p.add_run().add_break()
        p.add_run('Client: {{ client }}')
        doc.save(template_path)

        outputs = _render(work, {'template': 'template.docx',
                                 'fields': {'client': 'OOO Romashka'}},
                          output_dir=str(tmp_path / 'out'))

        assert len(outputs) == 1
        with zipfile.ZipFile(outputs[0], 'r') as zf:
            xml = etree.parse(zf.open('word/document.xml'))
        assert len(xml.findall('.//{%s}br' % W)) == 1
        texts = _docx_paragraph_texts(outputs[0])
        assert texts == ['Line1\nClient: OOO Romashka']

    def test_render_from_json_table_cells(self, tmp_path):
        """Placeholders inside table cells render from Filling JSON."""
        work, template_path = _make_work(tmp_path)
        doc = Document()
        doc.add_paragraph('List:')
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = '{{ item1 }}'
        table.cell(0, 1).text = 'static'
        table.cell(1, 0).text = '{{ item2 }}'
        table.cell(1, 1).text = '{{ item1 }}'
        doc.save(template_path)

        outputs = _render(work, {'template': 'template.docx',
                                 'fields': {'item1': 'A', 'item2': 'B'}},
                          output_dir=str(tmp_path / 'out'))

        assert len(outputs) == 1
        result = Document(outputs[0])
        cells = [cell.text for row in result.tables[0].rows
                 for cell in row.cells]
        assert cells == ['A', 'static', 'B', 'A']
