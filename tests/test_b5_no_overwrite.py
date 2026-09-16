# -*- coding: utf-8 -*-
"""B5 regression tests (002-stabilization): generation never overwrites files.

If the target file exists, the created file is renamed
(``файл (1)``, ``файл (2)``, …) and the existing file is left untouched.
Recent-projects format is out of scope and unchanged (read-only).
"""

import os
import tempfile

from docxforge.engine.render_loop import resolve_unique_output_path
from docxforge.generate import generate_project


def _make_single_doc_source(tmp_dir):
    """Minimal project rendering one doc per run (TABLE + SEQUENTIAL)."""
    from docx import Document
    from openpyxl import Workbook

    from docxforge.engine.schema import (
        BatchSourceConfig,
        FieldMapping,
        FieldType,
        Project,
        RowIterationMode,
        TemplateConfig,
    )

    project_dir = os.path.join(tmp_dir, 'b5_source')
    data_dir = os.path.join(project_dir, 'Данные')
    tmpl_dir = os.path.join(project_dir, 'Шаблоны')
    os.makedirs(data_dir, exist_ok=True)
    os.makedirs(tmpl_dir, exist_ok=True)

    doc = Document()
    doc.add_paragraph('Client: {{ client_name }}')
    doc.save(os.path.join(tmpl_dir, 'contract.docx'))

    wb = Workbook()
    ws = wb.active
    ws.append(['client_name'])
    ws.append(['Альфа'])
    ws.append(['Бета'])  # two rows: repeated generation (resume) still renders
    wb.save(os.path.join(data_dir, 'clients.xlsx'))

    prj = Project()
    tc = TemplateConfig()
    tc.fields['client_name'] = FieldMapping(
        type=FieldType.TABLE, file='clients.xlsx', column='client_name')
    tc.batch_sources = {
        'clients.xlsx': BatchSourceConfig(
            file='clients.xlsx', mode=RowIterationMode.SEQUENTIAL),
    }
    prj.templates['contract.docx'] = tc
    prj.to_file(os.path.join(project_dir, 'проект.docxforge'))
    return project_dir


class TestResolveUniqueOutputPath:
    def test_render_loop_missing_file_passthrough(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, 'doc.docx')
            assert resolve_unique_output_path(path) == path

    def test_render_loop_existing_file_renamed_incrementally(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, 'doc.docx')
            open(path, 'w').close()
            first = resolve_unique_output_path(path)
            assert first == os.path.join(tmp, 'doc (1).docx')
            open(first, 'w').close()
            second = resolve_unique_output_path(path)
            assert second == os.path.join(tmp, 'doc (2).docx')


class TestGenerateNoOverwrite:
    def test_generate_existing_output_renamed_both_kept(self):
        with tempfile.TemporaryDirectory() as tmp:
            project_dir = _make_single_doc_source(tmp)

            first_run = generate_project(project_dir, num_docs=1)
            assert len(first_run) == 1
            first_path = first_run[0]
            assert os.path.isfile(first_path)
            with open(first_path, 'rb') as f:
                first_bytes = f.read()

            second_run = generate_project(project_dir, num_docs=1)
            assert len(second_run) == 1
            second_path = second_run[0]
            assert second_path != first_path
            assert os.path.basename(second_path).endswith(' (1).docx')

            # Existing file untouched, both outputs on disk.
            with open(first_path, 'rb') as f:
                assert f.read() == first_bytes
            assert os.path.isfile(second_path)
