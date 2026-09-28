# -*- coding: utf-8 -*-
"""B6 regression: creation-from-generation advances the source counter.

Covers the three creation entry points:
- ``generate.create_projects_from_template`` (flat),
- ``generate.create_nested_employee_projects`` (composite),
- ``schema.create_projects`` (engine level).

Each test records counter value before/after (было/стало).
"""

import logging
import os
import tempfile

from docxforge.engine.schema import (
    BatchSourceConfig,
    FieldMapping,
    FieldType,
    Project,
    RowIterationMode,
    TemplateConfig,
    advance_counter_after_creation,
    create_projects,
)
from docxforge.generate import (
    create_nested_employee_projects,
    create_projects_from_template,
)


logger = logging.getLogger(__name__)


def _write_docx(path, *paragraphs):
    from docx import Document
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(path)


def _write_xlsx(path, header, rows):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for row in rows:
        ws.append(row)
    wb.save(path)


def _make_flat_source(tmp_dir, last_counter=2):
    """Source project: COUNTER + SEQUENTIAL batch with 3 rows."""
    project_dir = os.path.join(tmp_dir, 'source_project')
    os.makedirs(os.path.join(project_dir, 'Данные'), exist_ok=True)
    os.makedirs(os.path.join(project_dir, 'Шаблоны'), exist_ok=True)
    _write_docx(
        os.path.join(project_dir, 'Шаблоны', 'contract.docx'),
        'Doc {{ doc_number }} {{ client_name }}',
    )
    _write_xlsx(
        os.path.join(project_dir, 'Данные', 'clients.xlsx'),
        ['client_name'],
        [['Alpha'], ['Beta'], ['Gamma']],
    )
    prj = Project()
    tc = TemplateConfig()
    tc.fields['doc_number'] = FieldMapping(
        type=FieldType.COUNTER, start=1, format='0001')
    tc.fields['client_name'] = FieldMapping(
        type=FieldType.TABLE, file='clients.xlsx', column='client_name')
    tc.batch_sources = {
        'clients.xlsx': BatchSourceConfig(
            file='clients.xlsx', mode=RowIterationMode.SEQUENTIAL)
    }
    tc.resume.last_counter_value = last_counter
    prj.templates['contract.docx'] = tc
    prj.to_file(os.path.join(project_dir, 'проект.docxforge'))
    return project_dir


def _make_nested_source(tmp_dir, last_counter=1):
    """Source project with employee/project_name batch columns (6 rows)."""
    project_dir = os.path.join(tmp_dir, 'source_project')
    os.makedirs(os.path.join(project_dir, 'Данные'), exist_ok=True)
    os.makedirs(os.path.join(project_dir, 'Шаблоны'), exist_ok=True)
    _write_docx(
        os.path.join(project_dir, 'Шаблоны', 'report.docx'),
        'Employee: {{ employee }}',
        'Project: {{ project_name }}',
        'Doc: {{ doc_number }}',
    )
    _write_xlsx(
        os.path.join(project_dir, 'Данные', 'batch.xlsx'),
        ['employee', 'project_name'],
        [
            ['Ivanov', 'Deal_001'],
            ['Ivanov', 'Deal_002'],
            ['Ivanov', 'Deal_003'],
            ['Petrov', 'Deal_004'],
            ['Petrov', 'Deal_005'],
            ['Petrov', 'Deal_006'],
        ],
    )
    prj = Project()
    tc = TemplateConfig()
    tc.fields['employee'] = FieldMapping(
        type=FieldType.TABLE, file='batch.xlsx', column='employee')
    tc.fields['project_name'] = FieldMapping(
        type=FieldType.TABLE, file='batch.xlsx', column='project_name')
    tc.fields['doc_number'] = FieldMapping(
        type=FieldType.COUNTER, start=1, format='0001')
    tc.batch_sources = {
        'batch.xlsx': BatchSourceConfig(
            file='batch.xlsx', mode=RowIterationMode.SEQUENTIAL)
    }
    tc.resume.last_counter_value = last_counter
    prj.templates['report.docx'] = tc
    prj.to_file(os.path.join(project_dir, 'проект.docxforge'))
    return project_dir


def _read_counter(project_dir, template_name):
    prj = Project.from_file(os.path.join(project_dir, 'проект.docxforge'))
    return prj.templates[template_name].resume.last_counter_value


class TestB6HelperMath:
    def test_helper_create_projects_advances_counter_by_created(self):
        from docxforge.engine.schema import ResumeState
        resume = ResumeState(last_counter_value=2, continue_from_last=True)
        assert advance_counter_after_creation(resume, 2) == 4
        assert resume.last_counter_value == 4

    def test_helper_create_projects_no_continue_resets_base(self):
        from docxforge.engine.schema import ResumeState
        resume = ResumeState(last_counter_value=5, continue_from_last=False)
        assert advance_counter_after_creation(resume, 2) == 2

    def test_helper_create_projects_zero_created_is_noop(self):
        from docxforge.engine.schema import ResumeState
        resume = ResumeState(last_counter_value=5, continue_from_last=True)
        assert advance_counter_after_creation(resume, 0) == 5


class TestB6FlatCreation:
    def test_generate_create_projects_advances_source_counter(self):
        with tempfile.TemporaryDirectory() as tmp:
            src = _make_flat_source(tmp, last_counter=2)
            before = _read_counter(src, 'contract.docx')
            _, created = create_projects_from_template(
                src, 'contract.docx', '{{client_name}}', max_projects=2)
            after = _read_counter(src, 'contract.docx')
            assert created == 2
            assert before == 2
            assert after == before + created == 4


class TestB6NestedCreation:
    def test_generate_create_nested_advances_source_counter(self):
        with tempfile.TemporaryDirectory() as tmp:
            src = _make_nested_source(tmp, last_counter=1)
            before = _read_counter(src, 'report.docx')
            _, _, created = create_nested_employee_projects(
                src, 'report.docx', '{{employee}}/{{project_name}}')
            after = _read_counter(src, 'report.docx')
            assert created == 6
            assert before == 1
            assert after == before + created == 7


class TestB6EngineCreation:
    def test_schema_create_projects_advances_source_counter(self):
        with tempfile.TemporaryDirectory() as tmp:
            src = _make_flat_source(tmp, last_counter=0)
            before = _read_counter(src, 'contract.docx')
            created = create_projects(
                source_project_dir=src,
                output_base_dir=os.path.join(tmp, 'output_projects'),
                template_name='contract.docx',
                batch_source_name='clients.xlsx',
            )
            after = _read_counter(src, 'contract.docx')
            assert len(created) == 3
            assert before == 0
            assert after == before + len(created) == 3
