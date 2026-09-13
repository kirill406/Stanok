# -*- coding: utf-8 -*-
"""Tests for batch generation — each document must get DIFFERENT data."""

import os, tempfile, zipfile

from lxml import etree
import openpyxl
from docx import Document

from docxforge.engine.schema import (
    Project, TemplateConfig, FieldMapping, FieldType,
)
from docxforge.engine import Renderer
from docxforge.engine.data_reader import DataReader

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def _read_output_text(path):
    with zipfile.ZipFile(path, 'r') as zf:
        doc = etree.parse(zf.open('word/document.xml'))
        lines = []
        for p in doc.findall('.//{%s}p' % W):
            text = ''
            for r in p.findall('{%s}r' % W):
                for t in r.findall('{%s}t' % W):
                    if t.text:
                        text += t.text
            if text:
                lines.append(text)
        return '\n'.join(lines)


class TestBatch:
    """Test: batch generation produces DIFFERENT documents."""

    def test_batch_two_clients_produce_two_different_docs(self):
        """2 client rows → 2 documents, each with different client name."""
        with tempfile.TemporaryDirectory() as tmp:
            data_dir = os.path.join(tmp, 'Данные')
            tmpl_dir = os.path.join(tmp, 'Шаблоны')
            os.makedirs(data_dir)
            os.makedirs(tmpl_dir)

            # clients.xlsx with 2 rows
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(['название', 'ИНН'])
            ws.append(['ООО Альфа', '111'])
            ws.append(['ООО Бета',  '222'])
            wb.save(os.path.join(data_dir, 'клиенты.xlsx'))

            # Template
            doc = Document()
            doc.add_paragraph('Клиент: {{ клиент }}')
            doc.add_paragraph('ИНН: {{ ИНН }}')
            doc.save(os.path.join(tmpl_dir, 't.docx'))

            # Config
            prj = Project()
            tc = TemplateConfig()
            tc.fields['клиент'] = FieldMapping(type=FieldType.TABLE, file='клиенты.xlsx', column='название')
            tc.fields['ИНН'] = FieldMapping(type=FieldType.TABLE, file='клиенты.xlsx', column='ИНН', linked_to='клиент')
            prj.templates['t.docx'] = tc
            prj.to_file(os.path.join(tmp, 'проект.docxforge'))

            reader = DataReader()
            renderer = Renderer(tmp, reader)
            renderer.load_project()
            outputs = renderer.render('t.docx', {}, batch_table='клиенты.xlsx')

            assert len(outputs) == 2, 'Expected 2 output files, got %d' % len(outputs)

            text1 = _read_output_text(outputs[0])
            text2 = _read_output_text(outputs[1])

            print('Doc 1:', repr(text1))
            print('Doc 2:', repr(text2))

            assert 'ООО Альфа' in text1, 'Doc 1 must contain Альфа: %r' % text1
            assert '111' in text1, 'Doc 1 must contain 111'

            assert 'ООО Бета' in text2, 'Doc 2 must contain Бета: %r' % text2
            assert '222' in text2, 'Doc 2 must contain 222'

            assert text1 != text2, 'Documents must be DIFFERENT, got identical content'

    def test_batch_three_rows_three_docs(self):
        """3 rows → 3 documents, all different."""
        with tempfile.TemporaryDirectory() as tmp:
            data_dir = os.path.join(tmp, 'Данные')
            tmpl_dir = os.path.join(tmp, 'Шаблоны')
            os.makedirs(data_dir)
            os.makedirs(tmpl_dir)

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(['имя'])
            ws.append(['Анна'])
            ws.append(['Борис'])
            ws.append(['Вера'])
            wb.save(os.path.join(data_dir, 'люди.xlsx'))

            doc = Document()
            doc.add_paragraph('Имя: {{ имя }}')
            doc.save(os.path.join(tmpl_dir, 't.docx'))

            prj = Project()
            tc = TemplateConfig()
            tc.fields['имя'] = FieldMapping(type=FieldType.TABLE, file='люди.xlsx', column='имя')
            prj.templates['t.docx'] = tc
            prj.to_file(os.path.join(tmp, 'проект.docxforge'))

            reader = DataReader()
            renderer = Renderer(tmp, reader)
            renderer.load_project()
            outputs = renderer.render('t.docx', {}, batch_table='люди.xlsx')

            assert len(outputs) == 3
            texts = [_read_output_text(o) for o in outputs]
            names = set()
            for t in texts:
                # Extract name after "Имя: "
                names.add(t.replace('Имя: ', ''))
            assert names == {'Анна', 'Борис', 'Вера'}, 'Got: %s' % names

    def test_batch_counter_increments(self):
        """Each document gets incremental counter."""
        with tempfile.TemporaryDirectory() as tmp:
            data_dir = os.path.join(tmp, 'Данные')
            tmpl_dir = os.path.join(tmp, 'Шаблоны')
            os.makedirs(data_dir)
            os.makedirs(tmpl_dir)

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(['имя'])
            ws.append(['X'])
            ws.append(['Y'])
            wb.save(os.path.join(data_dir, 'data.xlsx'))

            doc = Document()
            doc.add_paragraph('Приказ № {{ doc_number }}')
            doc.save(os.path.join(tmpl_dir, 't.docx'))

            prj = Project()
            tc = TemplateConfig()
            tc.fields['doc_number'] = FieldMapping(type=FieldType.COUNTER, start=100, format='1')
            prj.templates['t.docx'] = tc
            prj.to_file(os.path.join(tmp, 'проект.docxforge'))

            reader = DataReader()
            renderer = Renderer(tmp, reader)
            renderer.load_project()
            outputs = renderer.render('t.docx', {}, batch_table='data.xlsx')

            assert len(outputs) == 2
            t1 = _read_output_text(outputs[0])
            t2 = _read_output_text(outputs[1])
            assert 'Приказ № 100' in t1, 'Got: %r' % t1
            assert 'Приказ № 101' in t2, 'Got: %r' % t2
