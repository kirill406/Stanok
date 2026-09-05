# -*- coding: utf-8 -*-
"""Coverage-boosting tests: edge cases in renderer, data_reader, template_parser."""

import os, sys, tempfile, zipfile
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from docx import Document
from lxml import etree

from docxforge.engine.schema import (
    Project, TemplateConfig, FieldMapping, FieldType,
    CycleMapping, AggregationMapping, AggregationFunction,
    BatchSourceConfig, RowIterationMode,
)
from docxforge.engine import (
    Renderer, format_counter, format_today, compute_aggregation,
    merge_and_replace_paragraph, row_contains_placeholder,
    row_has_placeholders, expand_table_cycle, clone_element,
)
from docxforge.engine.data_reader import DataReader
from docxforge.engine.template_parser import scan_template

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W_NS = "{%s}" % W


def _make_dirs(tmp):
    os.makedirs(os.path.join(tmp, "Данные"), exist_ok=True)
    os.makedirs(os.path.join(tmp, "Шаблоны"), exist_ok=True)


def _read_output_text(path):
    with zipfile.ZipFile(path, "r") as zf:
        doc = etree.parse(zf.open("word/document.xml"))
        lines = []
        for p in doc.findall(".//{%s}p" % W):
            text = ""
            for r in p.findall("{%s}r" % W):
                for t in r.findall("{%s}t" % W):
                    if t.text:
                        text += t.text
            if text:
                lines.append(text)
        return "\n".join(lines)


# ============================================================
# Renderer edge cases: headers/footers, template in subdir, etc.
# ============================================================

class TestRendererHeadersFooters:
    """Test that headers/footers get placeholders replaced."""

    def test_header_with_placeholder(self):
        with tempfile.TemporaryDirectory() as tmp:
            _make_dirs(tmp)
            doc = Document()
            section = doc.sections[0]
            header = section.header
            hp = header.paragraphs[0]
            hp.text = "Договор № {{ doc_number }}"
            doc.add_paragraph("Body text")
            doc.save(os.path.join(tmp, "Шаблоны", "t.docx"))

            prj = Project()
            tc = TemplateConfig()
            tc.fields["doc_number"] = FieldMapping(type=FieldType.COUNTER, start=1, format="1")
            prj.templates["t.docx"] = tc
            prj.to_file(os.path.join(tmp, "проект.docxforge"))

            reader = DataReader()
            renderer = Renderer(tmp, reader)
            renderer.load_project()
            outputs = renderer.render("t.docx", {})
            assert len(outputs) == 1
            # Check header XML
            with zipfile.ZipFile(outputs[0], "r") as zf:
                for name in zf.namelist():
                    if "header" in name:
                        part = etree.parse(zf.open(name))
                        text = ""
                        for p in part.findall(".//{%s}p" % W):
                            for r in p.findall("{%s}r" % W):
                                for t in r.findall("{%s}t" % W):
                                    if t.text:
                                        text += t.text
                        # doc_number placeholder should be replaced
                        assert "{{" not in text, "Header still has placeholder: %r" % text
                        assert "1" in text


class TestRendererSubdirTemplate:
    """Test get_template_path when template is in a subdirectory."""

    def test_template_in_subdir(self):
        with tempfile.TemporaryDirectory() as tmp:
            _make_dirs(tmp)
            subdir = os.path.join(tmp, "Шаблоны", "subdir")
            os.makedirs(subdir, exist_ok=True)
            doc = Document()
            doc.add_paragraph("Hello {{ name }}")
            doc.save(os.path.join(subdir, "t.docx"))

            prj = Project()
            tc = TemplateConfig()
            tc.fields["name"] = FieldMapping(type=FieldType.CONSTANT, value="World")
            prj.templates["subdir/t.docx"] = tc
            prj.to_file(os.path.join(tmp, "проект.docxforge"))

            reader = DataReader()
            renderer = Renderer(tmp, reader)
            renderer.load_project()
            outputs = renderer.render("subdir/t.docx", {})
            assert len(outputs) == 1

    def test_missing_data_file_returns_empty(self):
        """_read_table_data with non-existent file returns []."""
        with tempfile.TemporaryDirectory() as tmp:
            _make_dirs(tmp)
            doc = Document()
            doc.add_paragraph("Hello {{ name }}")
            doc.save(os.path.join(tmp, "Шаблоны", "t.docx"))

            prj = Project()
            tc = TemplateConfig()
            tc.fields["name"] = FieldMapping(type=FieldType.TABLE, file="missing.xlsx", column="name")
            prj.templates["t.docx"] = tc
            prj.to_file(os.path.join(tmp, "проект.docxforge"))

            reader = DataReader()
            renderer = Renderer(tmp, reader)
            renderer.load_project()
            # Should not crash — returns empty list
            outputs = renderer.render("t.docx", {})
            assert len(outputs) == 1
            text = _read_output_text(outputs[0])
            # name stays unreplaced since no data
            assert "{{ name }}" in text or "name" not in text


class TestRendererCycleExpansion:
    """Test expand_table_cycle with actual table in template."""

    def test_cycle_expands_rows(self):
        with tempfile.TemporaryDirectory() as tmp:
            _make_dirs(tmp)
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["товар", "кол-во"])
            ws.append(["Товар A", "10"])
            ws.append(["Товар B", "20"])
            wb.save(os.path.join(tmp, "Данные", "spec.xlsx"))

            doc = Document()
            doc.add_paragraph("Договор")
            table = doc.add_table(rows=2, cols=2)
            table.style = "Table Grid"
            table.rows[0].cells[0].text = "Наименование"
            table.rows[0].cells[1].text = "Кол-во"
            table.rows[1].cells[0].text = "{{ товар }}"
            table.rows[1].cells[1].text = "{{ кол-во }}"
            doc.save(os.path.join(tmp, "Шаблоны", "t.docx"))

            prj = Project()
            tc = TemplateConfig()
            tc.cycles.append(CycleMapping(
                table="spec.xlsx",
                columns={"товар": "товар", "кол-во": "кол-во"}))
            prj.templates["t.docx"] = tc
            prj.to_file(os.path.join(tmp, "проект.docxforge"))

            reader = DataReader()
            renderer = Renderer(tmp, reader)
            renderer.load_project()
            outputs = renderer.render("t.docx", {})
            assert len(outputs) == 1

            with zipfile.ZipFile(outputs[0], "r") as zf:
                doc_xml = etree.parse(zf.open("word/document.xml"))
                # Should have original header row + 2 expanded rows = 3 rows total
                rows = doc_xml.findall(".//{%s}tr" % W)
                data_rows = [r for r in rows if any(
                    t.text and "Товар" in t.text
                    for tr_r in r.findall(".//{%s}r" % W)
                    for t in tr_r.findall("{%s}t" % W))]
                assert len(data_rows) == 2, "Expected 2 data rows, got %d" % len(data_rows)


class TestRendererBatchEdgeCases:
    """Batch mode edge cases: empty data, single doc fallback."""

    def test_batch_single_with_empty_data(self):
        """SINGLE mode with empty data file still produces 1 doc."""
        with tempfile.TemporaryDirectory() as tmp:
            _make_dirs(tmp)
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["name"])
            # No data rows
            wb.save(os.path.join(tmp, "Данные", "empty.xlsx"))

            doc = Document()
            doc.add_paragraph("Hello {{ name }}")
            doc.save(os.path.join(tmp, "Шаблоны", "t.docx"))

            prj = Project()
            tc = TemplateConfig()
            tc.fields["name"] = FieldMapping(type=FieldType.TABLE, file="empty.xlsx", column="name")
            prj.templates["t.docx"] = tc
            prj.to_file(os.path.join(tmp, "проект.docxforge"))

            reader = DataReader()
            renderer = Renderer(tmp, reader)
            renderer.load_project()

            bsc = BatchSourceConfig(file="empty.xlsx", mode=RowIterationMode.CONSTANT)
            outputs = renderer.render("t.docx", {},
                batch_table="empty.xlsx",
                batch_configs={"empty.xlsx": bsc})
            assert len(outputs) == 1

    def test_batch_configs_without_matching_table(self):
        """batch_configs without matching batch_table falls through to legacy."""
        with tempfile.TemporaryDirectory() as tmp:
            _make_dirs(tmp)
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["name"])
            ws.append(["A"])
            ws.append(["B"])
            wb.save(os.path.join(tmp, "Данные", "data.xlsx"))

            doc = Document()
            doc.add_paragraph("Name: {{ name }}")
            doc.save(os.path.join(tmp, "Шаблоны", "t.docx"))

            prj = Project()
            tc = TemplateConfig()
            tc.fields["name"] = FieldMapping(type=FieldType.TABLE, file="data.xlsx", column="name")
            prj.templates["t.docx"] = tc
            prj.to_file(os.path.join(tmp, "проект.docxforge"))

            reader = DataReader()
            renderer = Renderer(tmp, reader)
            renderer.load_project()

            # batch_configs has OTHER file, not data.xlsx
            bsc = BatchSourceConfig(file="other.xlsx", mode=RowIterationMode.CONSTANT)
            outputs = renderer.render("t.docx", {},
                batch_table="data.xlsx",
                batch_configs={"other.xlsx": bsc})
            # Falls through to legacy: all rows
            assert len(outputs) == 2


class TestAggregationEdgeCases:
    """Edge cases for compute_aggregation."""

    def test_sum_with_non_numeric_values(self):
        agg = AggregationMapping(function=AggregationFunction.SUM, table="x", column="цена")
        data = [{"цена": "abc"}, {"цена": "100"}, {"цена": "xyz"}]
        result = compute_aggregation(agg, data)
        assert result == "100"

    def test_max_all_non_numeric(self):
        agg = AggregationMapping(function=AggregationFunction.MAX, table="x", column="цена")
        data = [{"цена": "abc"}, {"цена": "xyz"}]
        result = compute_aggregation(agg, data)
        assert result == "0"

    def test_sum_multiply_without_multiplier(self):
        agg = AggregationMapping(
            function=AggregationFunction.SUM_MULTIPLY,
            table="x", column="цена", multiplier=None)
        data = [{"цена": "100"}, {"цена": "200"}]
        result = compute_aggregation(agg, data)
        assert result == "300"

    def test_sum_integer_result(self):
        agg = AggregationMapping(function=AggregationFunction.SUM, table="x", column="цена")
        data = [{"цена": "100"}, {"цена": "200"}]
        result = compute_aggregation(agg, data)
        # Integer result should not have decimals
        assert result == "300"
        assert "." not in result


class TestFormatTodayEdgeCases:
    """Edge cases for format_today."""

    def test_dd_MM_yyyy_HH_mm_format(self):
        from datetime import datetime
        dt = datetime(2026, 9, 4, 9, 5, 0)
        result = format_today("dd.MM.yyyy HH:mm", dt)
        assert result == "04.09.2026 09:05"
    def test_seconds_format(self):
        from datetime import datetime
        dt = datetime(2026, 9, 4, 14, 30, 45)
        result = format_today("ss", dt)
        assert result == "45"


class TestDataReaderEdgeCases:
    """Additional DataReader edge cases."""

    def test_read_excel_with_none_header_cell(self):
        with tempfile.TemporaryDirectory() as tmp:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append([None, "name", "age"])  # first cell is None
            ws.append(["x", "Anna", "25"])
            path = os.path.join(tmp, "data.xlsx")
            wb.save(path)
            reader = DataReader()
            rows = reader.read_excel(path)
            assert len(rows) == 1
            # Check that the None header becomes col_0
            assert "col_0" in rows[0]

    def test_get_columns_empty_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append([])
            path = os.path.join(tmp, "empty.xlsx")
            wb.save(path)
            reader = DataReader()
            cols = reader.get_columns(path)
            assert cols == []

    def test_get_distinct_values_empty_column(self):
        with tempfile.TemporaryDirectory() as tmp:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["name"])
            ws.append([""])
            ws.append([""])
            path = os.path.join(tmp, "data.xlsx")
            wb.save(path)
            reader = DataReader()
            vals = reader.get_distinct_values(path, "name")
            # Empty strings are skipped
            assert vals == []


class TestTemplateParserEdgeCases:
    """Template parser edge cases."""

    def test_scan_template_with_table(self):
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            doc.add_paragraph("Hello {{ name }}")
            table = doc.add_table(rows=2, cols=1)
            table.rows[0].cells[0].text = "Header"
            table.rows[1].cells[0].text = "{{ item }}"
            path = os.path.join(tmp, "t.docx")
            doc.save(path)
            result = scan_template(path)
            assert "name" in result["simple"]
            assert "item" in result["simple"]

    def test_scan_reserved_without_format(self):
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            doc.add_paragraph("Date: {{ now }}")
            doc.add_paragraph("Page: {{ page }}")
            path = os.path.join(tmp, "t.docx")
            doc.save(path)
            result = scan_template(path)
            # 'now' and 'page' are reserved — should go to doc_number
            assert "now" in result["doc_number"]
            assert "page" in result["doc_number"]
