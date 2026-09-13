# -*- coding: utf-8 -*-
"""Unit tests for create_projects mode.

Run with: python -m pytest tests/test_create_projects.py -v
"""

import os
import tempfile
import shutil
from pathlib import Path

import pytest


from docxforge.engine.schema import (
    Project, TemplateConfig, FieldMapping, FieldType,
    RowIterationMode, create_project, create_projects,
)
from docxforge.engine.data_reader import DataReader


def _make_source_project(tmp_dir: str) -> str:
    """Create a source project with template and batch data for testing."""
    from docx import Document
    from openpyxl import Workbook

    project_dir = os.path.join(tmp_dir, 'source_project')
    data_dir = os.path.join(project_dir, 'Данные')
    tmpl_dir = os.path.join(project_dir, 'Шаблоны')
    os.makedirs(data_dir, exist_ok=True)
    os.makedirs(tmpl_dir, exist_ok=True)

    # Create template
    doc = Document()
    doc.add_paragraph('Договор № {{ doc_number }}')
    doc.add_paragraph('Клиент: {{ client_name }}')
    doc.add_paragraph('Сумма: {{ amount }}')
    doc.add_paragraph('Регион: {{ region }}')
    doc.add_paragraph('Город: {{ city }}')
    doc.save(os.path.join(tmpl_dir, 'contract.docx'))

    # Create data file with 3 rows
    wb = Workbook()
    ws = wb.active
    ws.append(['client_name', 'amount', 'region', 'city'])
    ws.append(['ООО Альфа', '10000', 'Москва', 'Москва'])
    ws.append(['ООО Бета', '20000', 'СПб', 'Санкт-Петербург'])
    ws.append(['ИП Гамма', '15000', 'Урал', 'Екатеринбург'])
    wb.save(os.path.join(data_dir, 'clients.xlsx'))

    # Create project config
    prj = Project()
    tc = TemplateConfig()
    tc.fields['doc_number'] = FieldMapping(
        type=FieldType.COUNTER, start=1, format='0001')
    tc.fields['client_name'] = FieldMapping(
        type=FieldType.TABLE, file='clients.xlsx', column='client_name')
    tc.fields['amount'] = FieldMapping(
        type=FieldType.TABLE, file='clients.xlsx', column='amount')
    tc.fields['region'] = FieldMapping(
        type=FieldType.TABLE, file='clients.xlsx', column='region')
    tc.fields['city'] = FieldMapping(
        type=FieldType.TABLE, file='clients.xlsx', column='city')
    tc.fields['constant_field'] = FieldMapping(
        type=FieldType.CONSTANT, value='CONST_VALUE')
    tc.batch_sources['clients.xlsx'] = FieldMapping(
        type=FieldType.TABLE, file='clients.xlsx', column='client_name')
    # Actually use BatchSourceConfig for batch_sources
    from docxforge.engine.schema import BatchSourceConfig
    tc.batch_sources = {
        'clients.xlsx': BatchSourceConfig(
            file='clients.xlsx',
            mode=RowIterationMode.SEQUENTIAL,
        )
    }
    tc.directory_template = '{{ region }}/{{ city }}'
    tc.filename_template = '{{ client_name }}_договор.docx'
    prj.templates['contract.docx'] = tc
    prj.to_file(os.path.join(project_dir, 'проект.docxforge'))

    return project_dir


class TestCreateProjectsBasic:
    """Test basic create_projects functionality."""

    def test_create_projects_basic(self):
        """1 template, 1 table, 3 rows → 3 projects."""
        with tempfile.TemporaryDirectory() as tmp:
            source_dir = _make_source_project(tmp)
            output_base = os.path.join(tmp, 'output_projects')

            created = create_projects(
                source_project_dir=source_dir,
                output_base_dir=output_base,
                template_name='contract.docx',
                batch_source_name='clients.xlsx',
            )

            assert len(created) == 3
            for proj_dir in created:
                assert os.path.exists(os.path.join(proj_dir, 'проект.docxforge'))
                assert os.path.exists(os.path.join(proj_dir, 'Данные'))
                assert os.path.exists(os.path.join(proj_dir, 'Шаблоны'))
                assert os.path.exists(os.path.join(proj_dir, 'Результат'))
                assert os.path.exists(os.path.join(proj_dir, 'Шаблоны', 'contract.docx'))
                assert os.path.exists(os.path.join(proj_dir, 'Данные', 'clients.xlsx'))

    def test_create_projects_returns_project_paths(self):
        """Created projects list contains full paths."""
        with tempfile.TemporaryDirectory() as tmp:
            source_dir = _make_source_project(tmp)
            output_base = os.path.join(tmp, 'output_projects')

            created = create_projects(
                source_project_dir=source_dir,
                output_base_dir=output_base,
                template_name='contract.docx',
                batch_source_name='clients.xlsx',
            )

            for path in created:
                assert os.path.isabs(path)
                assert path.startswith(output_base)

            # Check that relative paths match expected (normalize separators)
            rel_paths = [os.path.relpath(p, output_base).replace('\\', '/') for p in created]
            expected = ['Москва/Москва', 'СПб/Санкт-Петербург', 'Урал/Екатеринбург']
            assert set(rel_paths) == set(expected)

    def test_create_projects_with_max_limit(self):
        """max_projects limit is respected."""
        with tempfile.TemporaryDirectory() as tmp:
            source_dir = _make_source_project(tmp)
            output_base = os.path.join(tmp, 'output_projects')

            created = create_projects(
                source_project_dir=source_dir,
                output_base_dir=output_base,
                template_name='contract.docx',
                batch_source_name='clients.xlsx',
                max_projects=2,
            )

            assert len(created) == 2


class TestConstantsPreserved:
    """Test that constant fields are copied unchanged."""

    def test_constants_preserved(self):
        """Constant fields copied unchanged to each project."""
        with tempfile.TemporaryDirectory() as tmp:
            source_dir = _make_source_project(tmp)
            output_base = os.path.join(tmp, 'output_projects')

            created = create_projects(
                source_project_dir=source_dir,
                output_base_dir=output_base,
                template_name='contract.docx',
                batch_source_name='clients.xlsx',
            )

            for proj_dir in created:
                prj = Project.from_file(os.path.join(proj_dir, 'проект.docxforge'))
                tc = prj.templates['contract.docx']
                fm = tc.fields['constant_field']
                assert fm.type == FieldType.CONSTANT
                assert fm.value == 'CONST_VALUE'


class TestTableFieldsPopulated:
    """Test that table fields get correct row values."""

    def test_table_fields_populated(self):
        """Each project gets correct row values for table fields."""
        with tempfile.TemporaryDirectory() as tmp:
            source_dir = _make_source_project(tmp)
            output_base = os.path.join(tmp, 'output_projects')

            created = create_projects(
                source_project_dir=source_dir,
                output_base_dir=output_base,
                template_name='contract.docx',
                batch_source_name='clients.xlsx',
            )

            expected_clients = ['ООО Альфа', 'ООО Бета', 'ИП Гамма']
            expected_amounts = ['10000', '20000', '15000']
            expected_regions = ['Москва', 'СПб', 'Урал']
            expected_cities = ['Москва', 'Санкт-Петербург', 'Екатеринбург']

            for i, proj_dir in enumerate(created):
                prj = Project.from_file(os.path.join(proj_dir, 'проект.docxforge'))
                tc = prj.templates['contract.docx']

                assert tc.fields['client_name'].type == FieldType.CONSTANT
                assert tc.fields['client_name'].value == expected_clients[i]

                assert tc.fields['amount'].type == FieldType.CONSTANT
                assert tc.fields['amount'].value == expected_amounts[i]

                assert tc.fields['region'].type == FieldType.CONSTANT
                assert tc.fields['region'].value == expected_regions[i]

                assert tc.fields['city'].type == FieldType.CONSTANT
                assert tc.fields['city'].value == expected_cities[i]


class TestCounterReset:
    """Test that counter fields are reset to start value."""

    def test_counter_reset(self):
        """Each project counter starts at start value."""
        with tempfile.TemporaryDirectory() as tmp:
            source_dir = _make_source_project(tmp)
            output_base = os.path.join(tmp, 'output_projects')

            created = create_projects(
                source_project_dir=source_dir,
                output_base_dir=output_base,
                template_name='contract.docx',
                batch_source_name='clients.xlsx',
            )

            for proj_dir in created:
                prj = Project.from_file(os.path.join(proj_dir, 'проект.docxforge'))
                tc = prj.templates['contract.docx']
                fm = tc.fields['doc_number']
                assert fm.type == FieldType.COUNTER
                assert fm.start == 1
                assert fm.format == '0001'


class TestFolderNameTemplate:
    """Test that folder name template is resolved with row data."""

    def test_folder_name_template(self):
        """Project folder names resolved from directory_template with row data."""
        with tempfile.TemporaryDirectory() as tmp:
            source_dir = _make_source_project(tmp)
            output_base = os.path.join(tmp, 'output_projects')

            created = create_projects(
                source_project_dir=source_dir,
                output_base_dir=output_base,
                template_name='contract.docx',
                batch_source_name='clients.xlsx',
            )

            expected_folders = [
                'Москва/Москва',
                'СПб/Санкт-Петербург',
                'Урал/Екатеринбург',
            ]

            rel_paths = [os.path.relpath(p, output_base).replace('\\', '/') for p in created]
            assert set(rel_paths) == set(expected_folders)


class TestMaxProjectsLimit:
    """Test that max_projects limit is respected."""

    def test_max_projects_limit(self):
        """User-specified max_projects limit is respected."""
        with tempfile.TemporaryDirectory() as tmp:
            source_dir = _make_source_project(tmp)
            output_base = os.path.join(tmp, 'output_projects')

            created = create_projects(
                source_project_dir=source_dir,
                output_base_dir=output_base,
                template_name='contract.docx',
                batch_source_name='clients.xlsx',
                max_projects=1,
            )

            assert len(created) == 1

    def test_max_projects_zero(self):
        """max_projects=0 returns empty list."""
        with tempfile.TemporaryDirectory() as tmp:
            source_dir = _make_source_project(tmp)
            output_base = os.path.join(tmp, 'output_projects')

            created = create_projects(
                source_project_dir=source_dir,
                output_base_dir=output_base,
                template_name='contract.docx',
                batch_source_name='clients.xlsx',
                max_projects=0,
            )

            assert len(created) == 0

    def test_max_projects_none_creates_all(self):
        """max_projects=None creates projects for all rows."""
        with tempfile.TemporaryDirectory() as tmp:
            source_dir = _make_source_project(tmp)
            output_base = os.path.join(tmp, 'output_projects')

            created = create_projects(
                source_project_dir=source_dir,
                output_base_dir=output_base,
                template_name='contract.docx',
                batch_source_name='clients.xlsx',
                max_projects=None,
            )

            assert len(created) == 3


class TestNoBatchSources:
    """Test graceful handling when no batch sources configured."""

    def test_no_batch_sources_raises_error(self):
        """create_projects raises ValueError when batch source not found."""
        with tempfile.TemporaryDirectory() as tmp:
            from docx import Document
            from openpyxl import Workbook

            project_dir = os.path.join(tmp, 'source_project')
            data_dir = os.path.join(project_dir, 'Данные')
            tmpl_dir = os.path.join(project_dir, 'Шаблоны')
            os.makedirs(data_dir, exist_ok=True)
            os.makedirs(tmpl_dir, exist_ok=True)

            doc = Document()
            doc.add_paragraph('Тест: {{ field }}')
            doc.save(os.path.join(tmpl_dir, 'test.docx'))

            prj = Project()
            tc = TemplateConfig()
            tc.fields['field'] = FieldMapping(
                type=FieldType.CONSTANT, value='value')
            # No batch_sources configured
            prj.templates['test.docx'] = tc
            prj.to_file(os.path.join(project_dir, 'проект.docxforge'))

            output_base = os.path.join(tmp, 'output_projects')

            with pytest.raises(ValueError, match='Batch source not found'):
                create_projects(
                    source_project_dir=project_dir,
                    output_base_dir=output_base,
                    template_name='test.docx',
                    batch_source_name='nonexistent.xlsx',
                )

    def test_non_sequential_mode_raises_error(self):
        """create_projects raises ValueError for non-SEQUENTIAL mode."""
        with tempfile.TemporaryDirectory() as tmp:
            source_dir = _make_source_project(tmp)

            # Modify batch source to CONSTANT mode
            prj = Project.from_file(os.path.join(source_dir, 'проект.docxforge'))
            from docxforge.engine.schema import BatchSourceConfig
            prj.templates['contract.docx'].batch_sources['clients.xlsx'] = BatchSourceConfig(
                file='clients.xlsx',
                mode=RowIterationMode.CONSTANT,
            )
            prj.to_file(os.path.join(source_dir, 'проект.docxforge'))

            output_base = os.path.join(tmp, 'output_projects')

            with pytest.raises(ValueError, match='SEQUENTIAL mode'):
                create_projects(
                    source_project_dir=source_dir,
                    output_base_dir=output_base,
                    template_name='contract.docx',
                    batch_source_name='clients.xlsx',
                )

    def test_empty_batch_source_returns_empty(self):
        """Empty batch source returns empty list."""
        with tempfile.TemporaryDirectory() as tmp:
            from docx import Document
            from openpyxl import Workbook

            project_dir = os.path.join(tmp, 'source_project')
            data_dir = os.path.join(project_dir, 'Данные')
            tmpl_dir = os.path.join(project_dir, 'Шаблоны')
            os.makedirs(data_dir, exist_ok=True)
            os.makedirs(tmpl_dir, exist_ok=True)

            doc = Document()
            doc.add_paragraph('Тест: {{ field }}')
            doc.save(os.path.join(tmpl_dir, 'test.docx'))

            # Empty data file (only header)
            wb = Workbook()
            ws = wb.active
            ws.append(['client_name', 'amount'])
            wb.save(os.path.join(data_dir, 'empty.xlsx'))

            prj = Project()
            tc = TemplateConfig()
            tc.fields['client_name'] = FieldMapping(
                type=FieldType.TABLE, file='empty.xlsx', column='client_name')
            from docxforge.engine.schema import BatchSourceConfig
            tc.batch_sources = {
                'empty.xlsx': BatchSourceConfig(
                    file='empty.xlsx',
                    mode=RowIterationMode.SEQUENTIAL,
                )
            }
            prj.templates['test.docx'] = tc
            prj.to_file(os.path.join(project_dir, 'проект.docxforge'))

            output_base = os.path.join(tmp, 'output_projects')

            created = create_projects(
                source_project_dir=project_dir,
                output_base_dir=output_base,
                template_name='test.docx',
                batch_source_name='empty.xlsx',
            )

            assert created == []


class TestCreateProjectsEdgeCases:
    """Additional edge case tests."""

    def test_counter_fields_not_shared(self):
        """Each project has independent counter state."""
        with tempfile.TemporaryDirectory() as tmp:
            source_dir = _make_source_project(tmp)
            output_base = os.path.join(tmp, 'output_projects')

            created = create_projects(
                source_project_dir=source_dir,
                output_base_dir=output_base,
                template_name='contract.docx',
                batch_source_name='clients.xlsx',
            )

            # All projects should have counter start=1
            for proj_dir in created:
                prj = Project.from_file(os.path.join(proj_dir, 'проект.docxforge'))
                tc = prj.templates['contract.docx']
                fm = tc.fields['doc_number']
                assert fm.start == 1

    def test_template_file_copied(self):
        """Template .docx file is copied to each project."""
        with tempfile.TemporaryDirectory() as tmp:
            source_dir = _make_source_project(tmp)
            output_base = os.path.join(tmp, 'output_projects')

            created = create_projects(
                source_project_dir=source_dir,
                output_base_dir=output_base,
                template_name='contract.docx',
                batch_source_name='clients.xlsx',
            )

            for proj_dir in created:
                tmpl_path = os.path.join(proj_dir, 'Шаблоны', 'contract.docx')
                assert os.path.exists(tmpl_path)

    def test_data_file_copied(self):
        """Batch source data file is copied to each project."""
        with tempfile.TemporaryDirectory() as tmp:
            source_dir = _make_source_project(tmp)
            output_base = os.path.join(tmp, 'output_projects')

            created = create_projects(
                source_project_dir=source_dir,
                output_base_dir=output_base,
                template_name='contract.docx',
                batch_source_name='clients.xlsx',
            )

            for proj_dir in created:
                data_path = os.path.join(proj_dir, 'Данные', 'clients.xlsx')
                assert os.path.exists(data_path)

    def test_other_batch_sources_preserved_as_constant(self):
        """Other batch sources are converted to CONSTANT mode."""
        with tempfile.TemporaryDirectory() as tmp:
            from docx import Document
            from openpyxl import Workbook

            project_dir = os.path.join(tmp, 'source_project')
            data_dir = os.path.join(project_dir, 'Данные')
            tmpl_dir = os.path.join(project_dir, 'Шаблоны')
            os.makedirs(data_dir, exist_ok=True)
            os.makedirs(tmpl_dir, exist_ok=True)

            doc = Document()
            doc.add_paragraph('{{ field1 }} - {{ field2 }}')
            doc.save(os.path.join(tmpl_dir, 'test.docx'))

            wb1 = Workbook()
            ws1 = wb1.active
            ws1.append(['field1'])
            ws1.append(['A1'])
            ws1.append(['A2'])
            wb1.save(os.path.join(data_dir, 'source1.xlsx'))

            wb2 = Workbook()
            ws2 = wb2.active
            ws2.append(['field2'])
            ws2.append(['B1'])
            ws2.append(['B2'])
            wb2.save(os.path.join(data_dir, 'source2.xlsx'))

            prj = Project()
            tc = TemplateConfig()
            tc.fields['field1'] = FieldMapping(
                type=FieldType.TABLE, file='source1.xlsx', column='field1')
            tc.fields['field2'] = FieldMapping(
                type=FieldType.TABLE, file='source2.xlsx', column='field2')
            from docxforge.engine.schema import BatchSourceConfig
            tc.batch_sources = {
                'source1.xlsx': BatchSourceConfig(
                    file='source1.xlsx',
                    mode=RowIterationMode.SEQUENTIAL,
                ),
                'source2.xlsx': BatchSourceConfig(
                    file='source2.xlsx',
                    mode=RowIterationMode.SEQUENTIAL,
                ),
            }
            prj.templates['test.docx'] = tc
            prj.to_file(os.path.join(project_dir, 'проект.docxforge'))

            output_base = os.path.join(tmp, 'output_projects')

            created = create_projects(
                source_project_dir=project_dir,
                output_base_dir=output_base,
                template_name='test.docx',
                batch_source_name='source1.xlsx',
            )

            for proj_dir in created:
                prj2 = Project.from_file(os.path.join(proj_dir, 'проект.docxforge'))
                tc2 = prj2.templates['test.docx']
                # source1 (iterated) should be CONSTANT
                assert tc2.batch_sources['source1.xlsx'].mode == RowIterationMode.CONSTANT
                # source2 (other) should be CONSTANT too (converted)
                assert tc2.batch_sources['source2.xlsx'].mode == RowIterationMode.CONSTANT


if __name__ == '__main__':
    import pytest
    pytest.main([__file__, '-v'])