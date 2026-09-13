# -*- coding: utf-8 -*-
"""Tests for DataReader and Schema (serialization round-trip)."""

import os
import tempfile
import zipfile


import openpyxl
from docx import Document
from lxml import etree

from docxforge.engine.schema import (
    Project, TemplateConfig, FieldMapping, FieldType,
    CycleMapping, AggregationMapping, AggregationFunction,
    create_project,
)
from docxforge.engine import Renderer, format_counter, format_today, compute_aggregation
from docxforge.engine.data_reader import DataReader
from docxforge.engine.template_parser import scan_template

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def _make_dirs(tmp):
    os.makedirs(os.path.join(tmp, 'Данные'), exist_ok=True)
    os.makedirs(os.path.join(tmp, 'Шаблоны'), exist_ok=True)


def _read_output_text(path):
    with zipfile.ZipFile(path, 'r') as zf:
        doc = etree.parse(zf.open('word/document.xml'))
        lines = []
        for p in doc.findall('.//{%s}p' % W):
            text = ''
            for r in p.findall('{%s}r' % W):
                for t in r.findall('{%s}t' % W):
                    if t.text:
                        text += t.text
            if text:
                lines.append(text)
        return '\n'.join(lines)


# ============================================================
# DataReader tests
# ============================================================

class TestDataReader:
    """Tests for DataReader: reading Excel, getting columns, distinct values."""

    def test_read_excel_basic(self):
        with tempfile.TemporaryDirectory() as tmp:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(['имя', 'возраст'])
            ws.append(['Анна', '25'])
            ws.append(['Борис', '30'])
            path = os.path.join(tmp, 'data.xlsx')
            wb.save(path)

            reader = DataReader()
            rows = reader.read_excel(path)
            assert len(rows) == 2
            assert rows[0]['имя'] == 'Анна'
            assert rows[0]['возраст'] == '25'
            assert rows[1]['имя'] == 'Борис'

    def test_read_excel_skips_empty_rows(self):
        with tempfile.TemporaryDirectory() as tmp:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(['имя'])
            ws.append(['Анна'])
            ws.append([])  # empty row
            ws.append(['Борис'])
            path = os.path.join(tmp, 'data.xlsx')
            wb.save(path)

            reader = DataReader()
            rows = reader.read_excel(path)
            assert len(rows) == 2
            assert rows[0]['имя'] == 'Анна'
            assert rows[1]['имя'] == 'Борис'

    def test_read_excel_empty_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            wb = openpyxl.Workbook()
            ws = wb.active
            # No data at all
            path = os.path.join(tmp, 'empty.xlsx')
            wb.save(path)

            reader = DataReader()
            rows = reader.read_excel(path)
            assert rows == []

    def test_get_columns(self):
        with tempfile.TemporaryDirectory() as tmp:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(['название', 'ИНН', 'адрес'])
            ws.append(['ООО Альфа', '111', 'ул. 1'])
            path = os.path.join(tmp, 'data.xlsx')
            wb.save(path)

            reader = DataReader()
            cols = reader.get_columns(path)
            assert cols == ['название', 'ИНН', 'адрес']

    def test_get_distinct_values(self):
        with tempfile.TemporaryDirectory() as tmp:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(['город'])
            ws.append(['Москва'])
            ws.append(['Питер'])
            ws.append(['Москва'])
            path = os.path.join(tmp, 'data.xlsx')
            wb.save(path)

            reader = DataReader()
            vals = reader.get_distinct_values(path, 'город')
            assert vals == ['Москва', 'Питер']

    def test_read_excel_numeric_values_become_strings(self):
        with tempfile.TemporaryDirectory() as tmp:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(['число'])
            ws.append([42])
            ws.append([3.14])
            path = os.path.join(tmp, 'data.xlsx')
            wb.save(path)

            reader = DataReader()
            rows = reader.read_excel(path)
            assert rows[0]['число'] == '42'
            assert rows[1]['число'] == '3.14'


# ============================================================
# Schema tests
# ============================================================

class TestSchema:
    """Tests for schema: round-trip save/load, field types."""

    def test_project_round_trip(self):
        with tempfile.TemporaryDirectory() as tmp:
            prj = Project()
            tc = TemplateConfig()
            tc.fields['имя'] = FieldMapping(type=FieldType.CONSTANT, value='Тест')
            tc.fields['клиент'] = FieldMapping(type=FieldType.TABLE, file='клиенты.xlsx', column='название')
            tc.fields['doc_number'] = FieldMapping(type=FieldType.COUNTER, start=5, format='001')
            tc.fields['дата'] = FieldMapping(type=FieldType.TODAY, format='dd.MM.yyyy')
            tc.fields['печать'] = FieldMapping(type=FieldType.IMAGE, value='seal.png')
            tc.cycles.append(CycleMapping(table='спецификация.xlsx', columns={'товар': 'наименование'}))
            tc.aggregations['итого'] = AggregationMapping(
                function=AggregationFunction.SUM, table='спецификация.xlsx', column='цена')
            prj.templates['contract.docx'] = tc

            path = os.path.join(tmp, 'проект.docxforge')
            prj.to_file(path)
            prj2 = Project.from_file(path)

            assert len(prj2.templates) == 1
            tc2 = prj2.templates['contract.docx']
            assert tc2.fields['имя'].type == FieldType.CONSTANT
            assert tc2.fields['имя'].value == 'Тест'
            assert tc2.fields['клиент'].type == FieldType.TABLE
            assert tc2.fields['клиент'].file == 'клиенты.xlsx'
            assert tc2.fields['doc_number'].start == 5
            assert tc2.fields['doc_number'].format == '001'
            assert tc2.fields['дата'].type == FieldType.TODAY
            assert tc2.fields['печать'].type == FieldType.IMAGE
            assert len(tc2.cycles) == 1
            assert len(tc2.aggregations) == 1

    def test_create_project(self):
        with tempfile.TemporaryDirectory() as tmp:
            proj_file = create_project(os.path.join(tmp, 'myproject'))
            assert os.path.exists(proj_file)
            prj = Project.from_file(proj_file)
            assert prj.version == 2
            assert len(prj.templates) == 0

    def test_multiply_aggregation_round_trip(self):
        with tempfile.TemporaryDirectory() as tmp:
            prj = Project()
            tc = TemplateConfig()
            tc.aggregations['с_ндс'] = AggregationMapping(
                function=AggregationFunction.SUM_MULTIPLY,
                table='spec.xlsx', column='цена', multiplier=1.2)
            prj.templates['t.docx'] = tc
            path = os.path.join(tmp, 'проект.docxforge')
            prj.to_file(path)
            prj2 = Project.from_file(path)
            agg = prj2.templates['t.docx'].aggregations['с_ндс']
            assert agg.function == AggregationFunction.SUM_MULTIPLY
            assert agg.multiplier == 1.2


# ============================================================
# Formatter tests
# ============================================================

