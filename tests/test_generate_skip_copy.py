# -*- coding: utf-8 -*-
"""Regression tests for 002 B4 «skip copy» (excluded tables).

Excluded tables are not copied into the Данные/ folder of generated
projects; the rest are copied as before. Copy flags live in
``docxforge.generate`` (new functions only); the per-table checkbox lives
in ``BatchSourceRow`` (batch_section.py).

Run with: python -m pytest tests/test_generate_skip_copy.py -v
"""

import os

import pytest

from docxforge import generate as gen_module
from docxforge.engine.schema import BatchSourceConfig, RowIterationMode
from docxforge.generate import copy_data_tree, get_skip_copy_tables
from docxforge.gui.strings import STRINGS


def _make_data_dir(base, files):
    src = os.path.join(base, 'Данные')
    os.makedirs(src, exist_ok=True)
    for name, content in files.items():
        with open(os.path.join(src, name), 'w', encoding='utf-8') as f:
            f.write(content)
    return src


def test_generate_skip_copy_excluded_not_copied(tmp_path):
    """Excluded tables are missing in dst, the rest are copied."""
    src = _make_data_dir(str(tmp_path), {
        'clients.xlsx': 'clients-data',
        'orders.xlsx': 'orders-data',
        'archive.xlsx': 'archive-data',
    })
    dst = os.path.join(str(tmp_path), 'out', 'Данные')
    copy_data_tree(src, dst, exclude_names={'orders.xlsx'})
    assert sorted(os.listdir(dst)) == ['archive.xlsx', 'clients.xlsx']


def test_generate_skip_copy_rest_copied_as_before(tmp_path):
    """Without excludes the copy matches the source byte-for-byte."""
    payload = {'a.xlsx': 'aaa', 'b.xlsx': 'bbb'}
    src = _make_data_dir(str(tmp_path), payload)
    dst = os.path.join(str(tmp_path), 'out', 'Данные')
    copy_data_tree(src, dst)
    assert sorted(os.listdir(dst)) == ['a.xlsx', 'b.xlsx']
    for name, content in payload.items():
        with open(os.path.join(dst, name), encoding='utf-8') as f:
            assert f.read() == content


def test_generate_skip_copy_flags_from_batch_sources():
    """Only sources flagged skip_copy are reported (getattr-based)."""
    keep = BatchSourceConfig(file='clients.xlsx',
                             mode=RowIterationMode.SEQUENTIAL)
    skip = BatchSourceConfig(file='orders.xlsx',
                             mode=RowIterationMode.SEQUENTIAL)
    skip.skip_copy = True
    skipped = get_skip_copy_tables({'clients.xlsx': keep,
                                    'orders.xlsx': skip})
    assert skipped == {'orders.xlsx'}
    assert get_skip_copy_tables({}) == set()
    assert get_skip_copy_tables(None) == set()


@pytest.mark.gui
def test_batch_skip_copy_checkbox_per_table(qtbot, sample_project):
    """Every batch table row has an unchecked «skip copy» checkbox."""
    from docxforge.gui.fill_form import FillForm
    dlg = FillForm(sample_project, 'all_fields.docx')
    qtbot.addWidget(dlg)
    assert dlg.batch_source_widgets, 'expected at least one batch table row'
    for df, widgets in dlg.batch_source_widgets.items():
        checkbox = widgets.get('chk_skip_copy')
        assert checkbox is not None, f'no skip-copy checkbox for {df}'
        assert checkbox.text() == STRINGS['batch_skip_copy']
        assert not checkbox.isChecked()


def _make_skip_flat_source(tmp_dir: str) -> str:
    """Flat source with 2 data files; orders.xlsx flagged skip_copy."""
    from docx import Document
    from openpyxl import Workbook

    from docxforge.engine.schema import (
        FieldMapping,
        FieldType,
        Project,
        TemplateConfig,
    )

    project_dir = os.path.join(tmp_dir, 'skip_source')
    data_dir = os.path.join(project_dir, 'Данные')
    tmpl_dir = os.path.join(project_dir, 'Шаблоны')
    os.makedirs(data_dir, exist_ok=True)
    os.makedirs(tmpl_dir, exist_ok=True)

    doc = Document()
    doc.add_paragraph('Клиент: {{ client_name }}')
    doc.save(os.path.join(tmpl_dir, 'contract.docx'))

    wb = Workbook()
    ws = wb.active
    ws.append(['client_name'])
    ws.append(['ООО Альфа'])
    wb.save(os.path.join(data_dir, 'clients.xlsx'))

    wb2 = Workbook()
    ws2 = wb2.active
    ws2.append(['order_id'])
    ws2.append([1])
    wb2.save(os.path.join(data_dir, 'orders.xlsx'))

    prj = Project()
    tc = TemplateConfig()
    tc.fields['client_name'] = FieldMapping(
        type=FieldType.TABLE, file='clients.xlsx', column='client_name')
    tc.batch_sources = {
        'clients.xlsx': BatchSourceConfig(
            file='clients.xlsx', mode=RowIterationMode.SEQUENTIAL),
        'orders.xlsx': BatchSourceConfig(
            file='orders.xlsx', mode=RowIterationMode.CONSTANT,
            skip_copy=True),
    }
    prj.templates['contract.docx'] = tc
    prj.to_file(os.path.join(project_dir, 'проект.docxforge'))
    return project_dir


def test_generate_skip_copy_wired_into_flat_creation(tmp_path):
    """End-to-end: flagged table missing in generated Данные/, rest copied."""
    source_dir = _make_skip_flat_source(str(tmp_path))
    home = str(tmp_path / 'home')
    os.makedirs(home, exist_ok=True)
    projects_dir, count = gen_module.create_projects_from_template(
        source_dir, 'contract.docx', '{{client_name}}',
        max_projects=1, home_dir=home)
    assert count == 1
    folders = os.listdir(projects_dir)
    assert len(folders) == 1
    data_files = sorted(os.listdir(
        os.path.join(projects_dir, folders[0], 'Данные')))
    assert data_files == ['clients.xlsx']


def test_generate_skip_copy_wired_into_nested_creation(tmp_path):
    """End-to-end (nested): flagged table missing in generated Данные/."""
    from docx import Document
    from openpyxl import Workbook

    from docxforge.engine.schema import (
        FieldMapping,
        FieldType,
        Project,
        TemplateConfig,
    )

    project_dir = os.path.join(str(tmp_path), 'nested_skip_source')
    data_dir = os.path.join(project_dir, 'Данные')
    tmpl_dir = os.path.join(project_dir, 'Шаблоны')
    os.makedirs(data_dir, exist_ok=True)
    os.makedirs(tmpl_dir, exist_ok=True)

    doc = Document()
    doc.add_paragraph('Сотрудник: {{ employee }}')
    doc.save(os.path.join(tmpl_dir, 'contract.docx'))

    wb = Workbook()
    ws = wb.active
    ws.append(['employee', 'project_name'])
    ws.append(['Ivanov_Ivan', 'Dogovor_001'])
    wb.save(os.path.join(data_dir, 'work.xlsx'))

    wb2 = Workbook()
    ws2 = wb2.active
    ws2.append(['note'])
    ws2.append(['x'])
    wb2.save(os.path.join(data_dir, 'extra.xlsx'))

    prj = Project()
    tc = TemplateConfig()
    tc.fields['employee'] = FieldMapping(
        type=FieldType.TABLE, file='work.xlsx', column='employee')
    tc.fields['project_name'] = FieldMapping(
        type=FieldType.TABLE, file='work.xlsx', column='project_name')
    tc.batch_sources = {
        'work.xlsx': BatchSourceConfig(
            file='work.xlsx', mode=RowIterationMode.SEQUENTIAL),
        'extra.xlsx': BatchSourceConfig(
            file='extra.xlsx', mode=RowIterationMode.CONSTANT,
            skip_copy=True),
    }
    prj.templates['contract.docx'] = tc
    prj.to_file(os.path.join(project_dir, 'проект.docxforge'))

    home = str(tmp_path / 'home')
    os.makedirs(home, exist_ok=True)
    projects_dir, _emp_count, count = gen_module.create_projects_from_template(
        project_dir, 'contract.docx', '{{employee}}/{{project_name}}',
        max_projects=1, home_dir=home)
    assert count == 1
    employee_dirs = os.listdir(projects_dir)
    assert len(employee_dirs) == 1
    inner = os.path.join(projects_dir, employee_dirs[0])
    project_folders = [d for d in os.listdir(inner)
                       if os.path.isdir(os.path.join(inner, d))]
    assert len(project_folders) == 1
    data_files = sorted(os.listdir(
        os.path.join(inner, project_folders[0], 'Данные')))
    assert data_files == ['work.xlsx']
