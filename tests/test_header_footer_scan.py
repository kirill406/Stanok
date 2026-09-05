# -*- coding: utf-8 -*-
"""Tests for header/footer placeholder scanning in template_parser."""

import os
import tempfile
import zipfile
from docx import Document
from lxml import etree

from docxforge.engine.template_parser import scan_template


def _add_header_with_text(docx_path, header_text):
    """Add a header with text to an existing .docx file."""
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    RELS_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
    CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"

    with zipfile.ZipFile(docx_path, "r") as zin:
        zdata = {name: zin.read(name) for name in zin.namelist()}

    # Create header XML
    header_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:hdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:p><w:r><w:t>' + header_text + '</w:t></w:r></w:p>'
        '</w:hdr>'
    )
    zdata["word/header1.xml"] = header_xml.encode("utf-8")

    # Add relationship
    rels_path = "word/_rels/document.xml.rels"
    if rels_path in zdata:
        rels_root = etree.fromstring(zdata[rels_path])
    else:
        rels_root = etree.Element("{%s}Relationships" % RELS_NS)
    rel = etree.SubElement(rels_root, "{%s}Relationship" % RELS_NS)
    rel.set("Id", "rIdHeader1")
    rel.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/header")
    rel.set("Target", "header1.xml")
    zdata[rels_path] = etree.tostring(rels_root, xml_declaration=True, encoding="UTF-8", standalone=True)

    # Update [Content_Types].xml
    ct_path = "[Content_Types].xml"
    if ct_path in zdata:
        ct_root = etree.fromstring(zdata[ct_path])
        override = etree.SubElement(ct_root, "{%s}Override" % CT_NS)
        override.set("PartName", "/word/header1.xml")
        override.set("ContentType", "application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml")
        zdata[ct_path] = etree.tostring(ct_root, xml_declaration=True, encoding="UTF-8", standalone=True)

    with zipfile.ZipFile(docx_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in zdata.items():
            zout.writestr(name, data)


class TestHeaderFooterScan:
    """Test that template_parser scans headers and footers."""

    def test_header_placeholder_found(self):
        """Placeholder in header is detected by scan_template."""
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            doc.add_paragraph("Body text")
            path = os.path.join(tmp, "t.docx")
            doc.save(path)
            _add_header_with_text(path, "Doc {{ company }}")

            result = scan_template(path)
            assert "company" in result["simple"]
            hf_phs = result["header_footer_placeholders"]
            assert len(hf_phs) >= 1
            assert any("company" in ph for ph, _, _ in hf_phs)

    def test_footer_placeholder_found(self):
        """Placeholder in footer is detected."""
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            doc.add_paragraph("Body")
            path = os.path.join(tmp, "t.docx")
            doc.save(path)
            # Add footer
            W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
            RELS_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
            footer_xml = (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<w:ftr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                '<w:p><w:r><w:t>Page {{ page_num }}</w:t></w:r></w:p>'
                '</w:ftr>'
            )
            with zipfile.ZipFile(path, "r") as zin:
                zdata = {n: zin.read(n) for n in zin.namelist()}
            zdata["word/footer1.xml"] = footer_xml.encode("utf-8")
            if "word/_rels/document.xml.rels" in zdata:
                rels_root = etree.fromstring(zdata["word/_rels/document.xml.rels"])
            else:
                rels_root = etree.Element("{%s}Relationships" % RELS_NS)
            rel = etree.SubElement(rels_root, "{%s}Relationship" % RELS_NS)
            rel.set("Id", "rIdFooter1")
            rel.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footer")
            rel.set("Target", "footer1.xml")
            zdata["word/_rels/document.xml.rels"] = etree.tostring(rels_root, xml_declaration=True, encoding="UTF-8", standalone=True)
            with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
                for name, data in zdata.items():
                    zout.writestr(name, data)

            result = scan_template(path)
            assert "page_num" in result["simple"]
            hf_phs = result["header_footer_placeholders"]
            assert any("page_num" in ph for ph, _, _ in hf_phs)

    def test_no_header_footer_returns_empty_list(self):
        """Template without headers/footers has empty header_footer_placeholders."""
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            doc.add_paragraph("Just body {{ field }}")
            path = os.path.join(tmp, "t.docx")
            doc.save(path)

            result = scan_template(path)
            assert result["header_footer_placeholders"] == []
            assert "field" in result["simple"]

    def test_header_footer_placeholders_have_part_name(self):
        """Each header/footer placeholder includes the XML part name."""
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            doc.add_paragraph("Body")
            path = os.path.join(tmp, "t.docx")
            doc.save(path)
            _add_header_with_text(path, "{{ org }}")

            result = scan_template(path)
            hf_phs = result["header_footer_placeholders"]
            assert len(hf_phs) >= 1
            ph_name, raw_text, part_name = hf_phs[0]
            assert part_name == "word/header1.xml"
