# -*- coding: utf-8 -*-
"""Tests for image insertion in the renderer."""

import os, tempfile, zipfile, struct

from docx import Document
from lxml import etree

from docxforge.engine.schema import Project, TemplateConfig, FieldMapping, FieldType
from docxforge.engine import Renderer
from docxforge.engine.data_reader import DataReader
from docxforge.engine.image_utils import (
    get_image_content_type, read_image_file, build_drawing_xml,
    add_image_to_zdata, add_image_relationship, insert_image_in_paragraph,
    DEFAULT_WIDTH_EMU, DEFAULT_HEIGHT_EMU,
)

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W_NS = "{%s}" % W

DATA = "Данные"
TMPL = "Шаблоны"
PJ = "проект.docxforge"

def _make_dirs(tmp):
    os.makedirs(os.path.join(tmp, DATA), exist_ok=True)
    os.makedirs(os.path.join(tmp, TMPL), exist_ok=True)

def _make_minimal_png():
    import zlib
    def chunk(ct, d):
        c = ct + d
        return struct.pack(">I", len(d)) + c + struct.pack(">I", zlib.crc32(c) & 0xFFFFFFFF)
    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    raw = b"\x00\xff\x00\x00"
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


class TestImageUtils:
    def test_get_image_content_type_png(self):
        assert get_image_content_type("logo.png") == "image/png"
    def test_get_image_content_type_jpg(self):
        assert get_image_content_type("photo.jpg") == "image/jpeg"
    def test_get_image_content_type_unknown(self):
        assert get_image_content_type("file.xyz") == "image/png"
    def test_read_image_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            png = _make_minimal_png()
            path = os.path.join(tmp, "test.png")
            with open(path, "wb") as f2:
                f2.write(png)
            data, ct = read_image_file(path)
            assert data == png
            assert ct == "image/png"
    def test_build_drawing_xml(self):
        xml = build_drawing_xml("rId1", 1000, 2000, "logo")
        assert b"w:drawing" in xml
        assert b"wp:inline" in xml
        assert b"rId1" in xml
        assert b"logo" in xml
    def test_build_drawing_xml_has_inline(self):
        xml = build_drawing_xml("rId1")
        assert b"wp:inline" in xml
        assert b"wp:extent" in xml
    def test_add_image_to_zdata(self):
        png = _make_minimal_png()
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "test.png")
            with open(path, "wb") as f2:
                f2.write(png)
            zdata = {}
            mp, ct = add_image_to_zdata(zdata, path, 1)
            assert mp == "word/media/image1.png"
            assert ct == "image/png"
            assert zdata[mp] == png
    def test_add_image_to_zdata_increments(self):
        png = _make_minimal_png()
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "test.png")
            with open(path, "wb") as f2:
                f2.write(png)
            zdata = {}
            mp1, _ = add_image_to_zdata(zdata, path, 5)
            assert mp1 == "word/media/image5.png"
            mp2, _ = add_image_to_zdata(zdata, path, 10)
            assert mp2 == "word/media/image10.png"
    def test_add_image_relationship(self):
        zdata = {}
        add_image_relationship(zdata, "rId1", "word/media/image1.png")
        assert "word/_rels/document.xml.rels" in zdata
        root = etree.fromstring(zdata["word/_rels/document.xml.rels"])
        ns = "http://schemas.openxmlformats.org/package/2006/relationships"
        rels = root.findall("{%s}Relationship" % ns)
        assert len(rels) == 1
        assert rels[0].get("Id") == "rId1"
        assert "image" in rels[0].get("Type")
        assert rels[0].get("Target") == "media/image1.png"
    def test_add_image_relationship_appends(self):
        zdata = {}
        add_image_relationship(zdata, "rId1", "word/media/image1.png")
        add_image_relationship(zdata, "rId2", "word/media/image2.png")
        root = etree.fromstring(zdata["word/_rels/document.xml.rels"])
        ns = "http://schemas.openxmlformats.org/package/2006/relationships"
        rels = root.findall("{%s}Relationship" % ns)
        assert len(rels) == 2
    def test_insert_image_in_paragraph(self):
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("{{ image:logo }}")
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "t.docx")
            doc.save(path)
            with zipfile.ZipFile(path, "r") as zf:
                doc_xml = etree.parse(zf.open("word/document.xml"))
            p_el = doc_xml.find(".//{%s}p" % W)
            insert_image_in_paragraph(p_el, "rId1", name="logo")
            drawings = p_el.findall(".//{%s}drawing" % W)
            assert len(drawings) == 1
            assert len(p_el.findall(W_NS + "r")) == 1


class TestImageRenderEndToEnd:
    def test_image_placeholder_replaced_by_drawing(self):
        with tempfile.TemporaryDirectory() as tmp:
            _make_dirs(tmp)
            png = _make_minimal_png()
            img_path = os.path.join(tmp, DATA, "logo.png")
            with open(img_path, "wb") as f2:
                f2.write(png)
            doc = Document()
            doc.add_paragraph("Header")
            doc.add_paragraph("{{ image:logo }}")
            doc.add_paragraph("Footer")
            doc.save(os.path.join(tmp, TMPL, "t.docx"))
            prj = Project()
            tc = TemplateConfig()
            tc.fields["logo"] = FieldMapping(type=FieldType.IMAGE, value=os.path.join(DATA, "logo.png"))
            prj.templates["t.docx"] = tc
            prj.to_file(os.path.join(tmp, PJ))
            rdr = DataReader()
            renderer = Renderer(tmp, rdr)
            renderer.load_project()
            out = renderer.render("t.docx", {})
            assert len(out) == 1
            with zipfile.ZipFile(out[0], "r") as zf:
                doc_xml = etree.parse(zf.open("word/document.xml"))
                assert len(doc_xml.findall(".//{%s}drawing" % W)) >= 1
                assert len(doc_xml.findall(".//{%s}inline" % "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing")) >= 1
                assert any("media" in n for n in zf.namelist())
                # Check image relationship exists
                rels_xml = etree.parse(zf.open("word/_rels/document.xml.rels"))
                rns = "http://schemas.openxmlformats.org/package/2006/relationships"
                img_type = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
                for rel in rels_xml.findall("{%s}Relationship" % rns):
                    if rel.get("Type") == img_type:
                        break
                else:
                    pytest.fail("No image relationship found")
            import pytest

    def test_image_missing_file_does_not_crash(self):
        with tempfile.TemporaryDirectory() as tmp:
            _make_dirs(tmp)
            doc = Document()
            doc.add_paragraph("{{ image:missing }}")
            doc.save(os.path.join(tmp, TMPL, "t.docx"))
            prj = Project()
            tc = TemplateConfig()
            tc.fields["missing"] = FieldMapping(type=FieldType.IMAGE, value=os.path.join(DATA, "nope.png"))
            prj.templates["t.docx"] = tc
            prj.to_file(os.path.join(tmp, PJ))
            rdr = DataReader()
            renderer = Renderer(tmp, rdr)
            renderer.load_project()
            out = renderer.render("t.docx", {})
            assert len(out) == 1

    def test_no_image_fields_renders_normally(self):
        with tempfile.TemporaryDirectory() as tmp:
            _make_dirs(tmp)
            doc = Document()
            doc.add_paragraph("Hello {{ name }}")
            doc.save(os.path.join(tmp, TMPL, "t.docx"))
            prj = Project()
            tc = TemplateConfig()
            tc.fields["name"] = FieldMapping(type=FieldType.CONSTANT, value="World")
            prj.templates["t.docx"] = tc
            prj.to_file(os.path.join(tmp, PJ))
            rdr = DataReader()
            renderer = Renderer(tmp, rdr)
            renderer.load_project()
            out = renderer.render("t.docx", {})
            assert len(out) == 1
            with zipfile.ZipFile(out[0], "r") as zf:
                doc_xml = etree.parse(zf.open("word/document.xml"))
                txt = "".join(t.text or "" for t in doc_xml.iter("{%s}t" % W))
                assert "World" in txt

    def test_image_and_text_placeholders_together(self):
        with tempfile.TemporaryDirectory() as tmp:
            _make_dirs(tmp)
            png = _make_minimal_png()
            img_path = os.path.join(tmp, DATA, "stamp.png")
            with open(img_path, "wb") as f2:
                f2.write(png)
            doc = Document()
            doc.add_paragraph("Contract: {{ number }}")
            doc.add_paragraph("{{ image:stamp }}")
            doc.add_paragraph("Signed: {{ name }}")
            doc.save(os.path.join(tmp, TMPL, "t.docx"))
            prj = Project()
            tc = TemplateConfig()
            tc.fields["number"] = FieldMapping(type=FieldType.CONSTANT, value="123")
            tc.fields["name"] = FieldMapping(type=FieldType.CONSTANT, value="Ivanov")
            tc.fields["stamp"] = FieldMapping(type=FieldType.IMAGE, value=os.path.join(DATA, "stamp.png"))
            prj.templates["t.docx"] = tc
            prj.to_file(os.path.join(tmp, PJ))
            rdr = DataReader()
            renderer = Renderer(tmp, rdr)
            renderer.load_project()
            out = renderer.render("t.docx", {})
            assert len(out) == 1
            with zipfile.ZipFile(out[0], "r") as zf:
                doc_xml = etree.parse(zf.open("word/document.xml"))
                txt = "".join(t.text or "" for t in doc_xml.iter("{%s}t" % W))
                assert "123" in txt
                assert "Ivanov" in txt
                assert len(doc_xml.findall(".//{%s}drawing" % W)) >= 1
                assert any("media" in n for n in zf.namelist())

    def test_image_in_output_zip_has_correct_binary(self):
        with tempfile.TemporaryDirectory() as tmp:
            _make_dirs(tmp)
            png = _make_minimal_png()
            img_path = os.path.join(tmp, DATA, "sig.png")
            with open(img_path, "wb") as f2:
                f2.write(png)
            doc = Document()
            doc.add_paragraph("{{ image:sig }}")
            doc.save(os.path.join(tmp, TMPL, "t.docx"))
            prj = Project()
            tc = TemplateConfig()
            tc.fields["sig"] = FieldMapping(type=FieldType.IMAGE, value=os.path.join(DATA, "sig.png"))
            prj.templates["t.docx"] = tc
            prj.to_file(os.path.join(tmp, PJ))
            rdr = DataReader()
            renderer = Renderer(tmp, rdr)
            renderer.load_project()
            out = renderer.render("t.docx", {})
            assert len(out) == 1
            with zipfile.ZipFile(out[0], "r") as zf:
                assert "word/media/image1.png" in zf.namelist()
                assert zf.read("word/media/image1.png") == png
