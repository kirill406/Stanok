# -*- coding: utf-8 -*-
"""Fixture-driven tests for the JSON layer (003).

Every `tests/json/<name>/` folder holds source data (xlsx, docx, filling.json)
plus `expected.docx`. The test renders the template with Filling JSON values
and compares text with the etalon (see tests/json/README.md).

Run with: python -m pytest tests/test_json_fixtures.py -q
"""

import json
import os
import shutil

import pytest

FIXTURES_ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                             'json')


def _case_dirs():
    return sorted(
        d for d in os.listdir(FIXTURES_ROOT)
        if os.path.isdir(os.path.join(FIXTURES_ROOT, d))
        and not d.startswith(('_', '.')))


def _fj_case_dirs():
    return [d for d in _case_dirs()
            if os.path.isfile(os.path.join(FIXTURES_ROOT, d, 'filling.json'))]


def _pj_case_dirs():
    return [d for d in _case_dirs()
            if os.path.isfile(os.path.join(FIXTURES_ROOT, d, 'project.json'))]


def _docx_texts(path):
    from docx import Document

    doc = Document(path)
    texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            texts.extend(cell.text for cell in row.cells)
    return texts


def assert_docx_text_equal(result_path: str, expected_path: str):
    """Text-level comparison (formatting not compared — see README)."""
    assert _docx_texts(result_path) == _docx_texts(expected_path)


@pytest.mark.parametrize('case', _fj_case_dirs())
def test_json_fixture_case(tmp_path, case):
    """Render template.docx with filling.json → text equals expected.docx."""
    from docxforge.engine.data_reader import DataReader
    from docxforge.engine.renderer import Renderer

    case_dir = os.path.join(FIXTURES_ROOT, case)
    with open(os.path.join(case_dir, 'filling.json'), encoding='utf-8') as f:
        filling = json.load(f)

    work = str(tmp_path / 'work')
    templates_dir = os.path.join(work, 'Шаблоны')
    os.makedirs(templates_dir, exist_ok=True)
    shutil.copy(os.path.join(case_dir, filling['template']), templates_dir)

    renderer = Renderer(work, DataReader())
    renderer.load_project()
    outputs = renderer.render_from_json(
        filling, output_dir=str(tmp_path / 'out'))
    assert len(outputs) == 1, outputs
    assert_docx_text_equal(
        outputs[0], os.path.join(case_dir, 'expected.docx'))


@pytest.mark.parametrize('case', _pj_case_dirs())
def test_json_pj_excel_to_filling(case):
    """PJ + Excel → Filling JSON equals expected_filling.json."""
    from docxforge.engine.data_formatting import (
        read_table_rows,
        resolve_fields,
        resolve_source_row,
    )
    from docxforge.engine.data_reader import DataReader
    from docxforge.engine.schema import FieldMapping, FieldType

    case_dir = os.path.join(FIXTURES_ROOT, case)
    with open(os.path.join(case_dir, 'project.json'), encoding='utf-8') as f:
        pj = json.load(f)
    with open(os.path.join(case_dir, 'expected_filling.json'),
              encoding='utf-8') as f:
        expected = json.load(f)

    mappings = {}
    for name, spec in pj['fields'].items():
        ftype = FieldType(spec['type'])
        mappings[name] = FieldMapping(
            type=ftype, file=spec.get('file'), column=spec.get('column'),
            value=spec.get('value'))

    reader = DataReader()
    rows = read_table_rows(
        reader, os.path.join(case_dir, pj['source']['file']))
    row = resolve_source_row(rows, pj['source'].get('mode', 'sequential'), 0)
    assert resolve_fields(mappings, row=row, doc_index=0) == expected['fields']
