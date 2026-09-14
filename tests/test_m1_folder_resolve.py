# -*- coding: utf-8 -*-
"""Regression tests for 001-review M1, M8, M9 (unified folder-name resolve).

M1: one engine core (schema.substitute_placeholders) shared by
    generate flat/nested, render_loop and schema directory_template.
M8: max_projects None/<=0 = all rows in every entry point.
M9: empty employee values never merge into a shared '' group.
"""
import os
import tempfile

from docx import Document
from openpyxl import Workbook

from docxforge.engine.render_loop import resolve_folder_name_template
from docxforge.engine.schema import (
    BatchSourceConfig, FieldMapping, FieldType, Project, RowIterationMode,
    TemplateConfig, create_projects, limit_rows, substitute_placeholders,
)
from docxforge.generate import (
    _resolve_folder_name_template,
    create_nested_employee_projects,
    create_projects_from_template,
)


def _write_flat_source(tmp_dir, rows):
    project_dir = os.path.join(tmp_dir, 'src')
    data_dir = os.path.join(project_dir, 'Данные')
    tmpl_dir = os.path.join(project_dir, 'Шаблоны')
    os.makedirs(data_dir, exist_ok=True)
    os.makedirs(tmpl_dir, exist_ok=True)
    doc = Document()
    doc.add_paragraph('Client: {{ client_name }}')
    doc.save(os.path.join(tmpl_dir, 'contract.docx'))
    wb = Workbook()
    ws = wb.active
    ws.append(['client_name'])
    for value in rows:
        ws.append([value])
    wb.save(os.path.join(data_dir, 'clients.xlsx'))
    prj = Project()
    tc = TemplateConfig()
    tc.fields['client_name'] = FieldMapping(
        type=FieldType.TABLE, file='clients.xlsx', column='client_name')
    tc.batch_sources = {'clients.xlsx': BatchSourceConfig(
        file='clients.xlsx', mode=RowIterationMode.SEQUENTIAL)}
    prj.templates['contract.docx'] = tc
    prj.to_file(os.path.join(project_dir, 'проект.docxforge'))
    return project_dir


def _write_nested_source(tmp_dir, rows, template='{{employee}}/{{project}}'):
    project_dir = os.path.join(tmp_dir, 'src')
    data_dir = os.path.join(project_dir, 'Данные')
    tmpl_dir = os.path.join(project_dir, 'Шаблоны')
    os.makedirs(data_dir, exist_ok=True)
    os.makedirs(tmpl_dir, exist_ok=True)
    doc = Document()
    doc.add_paragraph('Hi {{employee}} {{project}}')
    doc.save(os.path.join(tmpl_dir, 't.docx'))
    wb = Workbook()
    ws = wb.active
    ws.append(['employee', 'project'])
    for emp, proj in rows:
        ws.append([emp, proj])
    wb.save(os.path.join(data_dir, 'batch.xlsx'))
    prj = Project()
    tc = TemplateConfig()
    tc.fields['employee'] = FieldMapping(
        type=FieldType.TABLE, file='batch.xlsx', column='employee')
    tc.fields['project'] = FieldMapping(
        type=FieldType.TABLE, file='batch.xlsx', column='project')
    tc.batch_sources = {'batch.xlsx': BatchSourceConfig(
        file='batch.xlsx', mode=RowIterationMode.SEQUENTIAL)}
    prj.templates['t.docx'] = tc
    prj.to_file(os.path.join(project_dir, 'проект.docxforge'))
    return project_dir, template


class TestM1SingleCore:
    def test_m1_core_whitespace_tolerant(self):
        for template in ('{{name}}', '{{ name }}', '{{  name  }}',
                         '{{name }}', '{{ name}}'):
            result, unresolved = substitute_placeholders(
                template, {'name': 'X'})
            assert result == 'X', template
            assert unresolved == []

    def test_m1_core_missing_keep_vs_empty(self):
        kept, unresolved = substitute_placeholders('{{a}}/{{b}}', {'a': '1'})
        assert kept == '1/{{b}}'
        assert unresolved == ['b']
        emptied, unresolved = substitute_placeholders(
            '{{a}}/{{b}}', {'a': '1'}, on_missing='empty')
        assert emptied == '1/'
        assert unresolved == ['b']

    def test_m1_generate_and_render_loop_agree(self):
        tc = TemplateConfig()
        tc.fields['client_name'] = FieldMapping(
            type=FieldType.TABLE, file='clients.xlsx',
            column='client_name')
        row = {'client_name': 'Acme'}
        gen = _resolve_folder_name_template(
            '{{  client_name  }}_doc', row, tc)
        loop = resolve_folder_name_template(
            '{{  client_name  }}_doc', row_data=row)
        assert gen == loop == 'Acme_doc'

    def test_m1_render_loop_missing_becomes_empty(self):
        assert resolve_folder_name_template(
            '{{known}}_{{missing}}', row_data={'known': 'K'}) == 'K_'

    def test_m1_render_loop_extra_features_kept(self):
        from datetime import datetime
        out = resolve_folder_name_template(
            '{{c}}_{{n}}', constants={'c': 'C'},
            row_data={'n': 'N'})
        assert out == 'C_N'
        out = resolve_folder_name_template(
            'doc_{{counter}}', counter_value=7, counter_format='0000')
        assert out == 'doc_0007'
        out = resolve_folder_name_template(
            '{{today:yyyy}}', now=datetime(2024, 5, 6))
        assert out == '2024'


class TestM8UnifiedLimit:
    def test_m1_limit_rows_semantics(self):
        rows = [1, 2, 3]
        assert limit_rows(rows, None) == [1, 2, 3]
        assert limit_rows(rows, 0) == [1, 2, 3]
        assert limit_rows(rows, -2) == [1, 2, 3]
        assert limit_rows(rows, 2) == [1, 2]
        assert limit_rows(rows, 99) == [1, 2, 3]

    def test_m8_flat_zero_means_all(self):
        with tempfile.TemporaryDirectory() as tmp:
            src = _write_flat_source(tmp, ['A', 'B', 'C'])
            _projects_dir, created = create_projects_from_template(
                src, 'contract.docx', '{{client_name}}', max_projects=0)
            assert created == 3

    def test_m8_flat_negative_means_all(self):
        with tempfile.TemporaryDirectory() as tmp:
            src = _write_flat_source(tmp, ['A', 'B'])
            _projects_dir, created = create_projects_from_template(
                src, 'contract.docx', '{{client_name}}', max_projects=-1)
            assert created == 2

    def test_m8_nested_limit_still_caps(self):
        with tempfile.TemporaryDirectory() as tmp:
            src, template = _write_nested_source(
                tmp, [('E1', 'P1'), ('E2', 'P2'), ('E3', 'P3')])
            _d, employees, projects = create_nested_employee_projects(
                src, 't.docx', template, max_projects=2)
            assert projects == 2

    def test_m8_schema_zero_means_all(self):
        with tempfile.TemporaryDirectory() as tmp:
            src = _write_flat_source(tmp, ['A', 'B'])
            out = os.path.join(tmp, 'out')
            created = create_projects(
                src, out, 'contract.docx', 'clients.xlsx', max_projects=0)
            assert len(created) == 2


class TestM9EmptyEmployee:
    def test_m9_empty_employees_not_merged(self, caplog):
        import logging
        with tempfile.TemporaryDirectory() as tmp:
            src, template = _write_nested_source(
                tmp, [('', 'P1'), ('', 'P2'), ('E3', 'P3')])
            with caplog.at_level(logging.WARNING):
                _d, employees, projects = create_nested_employee_projects(
                    src, 't.docx', template)
            assert projects == 3
            assert employees == 3
            assert any('Empty employee' in (r.getMessage() or '')
                       for r in caplog.records)

    def test_m9_no_empty_literal_folders(self):
        with tempfile.TemporaryDirectory() as tmp:
            src, template = _write_nested_source(tmp, [('', 'P1'), ('', 'P2')])
            projects_dir, _e, _p = create_nested_employee_projects(
                src, 't.docx', template)
            names = os.listdir(projects_dir)
            assert '' not in names
            assert len(names) == 2
            assert names[0] != names[1]
