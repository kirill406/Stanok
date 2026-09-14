# -*- coding: utf-8 -*-
"""M2 contract: the two row->project entry points share folder-name semantics.

``schema.create_projects`` (engine-level, keeps TABLE mappings) and
``generate.create_projects_from_template`` (full pipeline, freezes
TABLE->CONSTANT) serve different callers, so both stay. What M2 unifies
is their contract: same rows + same template -> same folder sequence
(resolve -> sanitize -> unique, shared core from B5/M1/M8).
"""
import os
import tempfile

from docx import Document
from openpyxl import Workbook

from docxforge.engine.schema import (
    BatchSourceConfig, FieldMapping, FieldType, Project, RowIterationMode,
    TemplateConfig, create_projects,
)
from docxforge.generate import create_projects_from_template


def _make_source(tmp_dir, rows, directory_template):
    project_dir = os.path.join(tmp_dir, 'src')
    data_dir = os.path.join(project_dir, 'Данные')
    tmpl_dir = os.path.join(project_dir, 'Шаблоны')
    os.makedirs(data_dir, exist_ok=True)
    os.makedirs(tmpl_dir, exist_ok=True)
    doc = Document()
    doc.add_paragraph('Name: {{ name }}')
    doc.save(os.path.join(tmpl_dir, 't.docx'))
    wb = Workbook()
    ws = wb.active
    ws.append(['name'])
    for value in rows:
        ws.append([value])
    wb.save(os.path.join(data_dir, 'data.xlsx'))
    prj = Project()
    tc = TemplateConfig()
    tc.fields['name'] = FieldMapping(
        type=FieldType.TABLE, file='data.xlsx', column='name')
    tc.batch_sources = {'data.xlsx': BatchSourceConfig(
        file='data.xlsx', mode=RowIterationMode.SEQUENTIAL)}
    tc.directory_template = directory_template
    prj.templates['t.docx'] = tc
    prj.to_file(os.path.join(project_dir, 'проект.docxforge'))
    return project_dir


class TestM2SharedContract:
    def test_m2_same_rows_same_folders(self):
        rows = ['Alpha', 'Same', 'Same', '  spaced  ']
        with tempfile.TemporaryDirectory() as tmp:
            src = _make_source(tmp, rows, '{{ name }}')
            out = os.path.join(tmp, 'out')
            schema_created = create_projects(
                src, out, 't.docx', 'data.xlsx')
            schema_names = sorted(
                os.path.basename(p) for p in schema_created)

            _projects_dir, count = create_projects_from_template(
                src, 't.docx', '{{ name }}')
            assert count == len(rows)
            gen_names = sorted(
                d for d in os.listdir(os.path.join(src, 'Projects'))
                if os.path.isdir(os.path.join(src, 'Projects', d)))
            assert schema_names == gen_names
            assert 'Same_1' in gen_names
            assert 'spaced' in gen_names

    def test_m2_entrypoints_document_each_other(self):
        assert 'create_projects_from_template' in (
            create_projects.__doc__ or '')
        from docxforge import generate as gen_module
        assert 'schema.create_projects' in (
            gen_module.create_projects_from_template.__doc__ or '')
