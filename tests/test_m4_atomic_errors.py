# -*- coding: utf-8 -*-
"""Regression tests for specs/001-review MAJOR items M4 and M5.

M4: all ``.docxforge`` project writes go through the atomic
    (tmp-write + fsync + os.replace) helper — a crash between tmp-write
    and replace must not leave a half-written project file.
M5: engine failures carry machine-readable codes with a single
    code -> text mapping (GUI/CLI reuse), coexisting with the legacy
    ``GenerationError('No documents generated...')`` message.
"""

import json
import logging
import os
import tempfile

import pytest

from docxforge.engine.errors import (
    EMPTY_SEQUENTIAL,
    TABLE_EXHAUSTED,
    NO_DOCUMENTS_GENERATED,
    ERROR_MESSAGES,
    EngineError,
    RenderError,
    message_for_code,
)
from docxforge.engine.schema import Project, TemplateConfig
from docxforge.generate import GenerationError, generate_project


def _write_initial_project(path):
    prj = Project()
    prj.templates['t.docx'] = TemplateConfig()
    prj.to_file(path)
    with open(path, 'r', encoding='utf-8') as f:
        return f.read()


class TestM4AtomicWrites:
    def test_m4_to_file_crash_between_tmp_and_replace_keeps_original(self, monkeypatch):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, 'проект.docxforge')
            original = _write_initial_project(path)

            changed = Project(version=99)
            changed.templates['other.docx'] = TemplateConfig()

            real_replace = os.replace

            def flaky_replace(src, dst):
                # Simulate a killed process: tmp is fully written, but the
                # final atomic rename never happens.
                assert os.path.exists(src)
                raise OSError('simulated crash before replace')

            monkeypatch.setattr(os, 'replace', flaky_replace)
            with pytest.raises(OSError, match='simulated crash'):
                changed.to_file(path)

            monkeypatch.setattr(os, 'replace', real_replace)
            with open(path, 'r', encoding='utf-8') as f:
                assert f.read() == original

    def test_m4_to_file_writes_tmp_before_replace(self, monkeypatch):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, 'проект.docxforge')
            _write_initial_project(path)

            changed = Project(version=7)
            changed.templates['n.docx'] = TemplateConfig()
            seen = {}

            real_replace = os.replace

            def record_replace(src, dst):
                assert src == path + '.tmp'
                assert dst == path
                with open(src, 'r', encoding='utf-8') as f:
                    seen['tmp_payload'] = json.load(f)
                return real_replace(src, dst)

            monkeypatch.setattr(os, 'replace', record_replace)
            changed.to_file(path)
            assert seen['tmp_payload']['version'] == 7
            assert 'n.docx' in seen['tmp_payload']['templates']
            assert not os.path.exists(path + '.tmp')
            with open(path, 'r', encoding='utf-8') as f:
                assert json.load(f)['version'] == 7

    def test_m4_save_project_stays_atomic_and_keeps_backup(self):
        from docxforge.engine.data_reader import DataReader
        from docxforge.engine.renderer import Renderer

        with tempfile.TemporaryDirectory() as tmp:
            os.makedirs(os.path.join(tmp, 'Шаблоны'), exist_ok=True)
            os.makedirs(os.path.join(tmp, 'Данные'), exist_ok=True)
            path = os.path.join(tmp, 'проект.docxforge')
            original = _write_initial_project(path)

            renderer = Renderer(tmp, DataReader())
            renderer.load_project()
            renderer.project.version = 42
            renderer.save_project()

            with open(path, 'r', encoding='utf-8') as f:
                assert json.load(f)['version'] == 42
            bak = path + '.bak'
            assert os.path.exists(bak)
            with open(bak, 'r', encoding='utf-8') as f:
                assert f.read() == original


class TestM5EngineErrorCodes:
    def test_m5_error_codes_have_text_mapping(self):
        for code in (EMPTY_SEQUENTIAL, TABLE_EXHAUSTED, NO_DOCUMENTS_GENERATED):
            assert code in ERROR_MESSAGES
            assert isinstance(ERROR_MESSAGES[code], str)
            assert ERROR_MESSAGES[code].strip()
        assert 'не имеет данных' in message_for_code(
            EMPTY_SEQUENTIAL, source='data.xlsx')
        assert 'исчерпана' in message_for_code(
            TABLE_EXHAUSTED, source='data.xlsx', rows=0, start=0)
        assert 'No documents generated' in message_for_code(
            NO_DOCUMENTS_GENERATED)

    def test_m5_engine_error_carries_code(self):
        err = RenderError('boom', code=TABLE_EXHAUSTED)
        assert isinstance(err, EngineError)
        assert err.code == TABLE_EXHAUSTED
        assert isinstance(GenerationError('x', code=EMPTY_SEQUENTIAL).code, str)

    def test_m5_generate_project_empty_reports_code(self):
        from docx import Document
        from openpyxl import Workbook

        from docxforge.engine.data_reader import DataReader
        from docxforge.engine.renderer import Renderer  # noqa: F401
        from docxforge.engine.schema import (
            BatchSourceConfig, FieldMapping, FieldType, RowIterationMode,
        )

        with tempfile.TemporaryDirectory() as tmp:
            data_dir = os.path.join(tmp, 'Данные')
            tmpl_dir = os.path.join(tmp, 'Шаблоны')
            os.makedirs(data_dir, exist_ok=True)
            os.makedirs(tmpl_dir, exist_ok=True)
            doc = Document()
            doc.add_paragraph('Name: {{ name }}')
            doc.save(os.path.join(tmpl_dir, 't.docx'))
            wb = Workbook()
            ws = wb.active
            ws.append(['name'])
            wb.save(os.path.join(data_dir, 'data.xlsx'))
            prj = Project()
            tc = TemplateConfig()
            tc.fields['name'] = FieldMapping(
                type=FieldType.TABLE, file='data.xlsx', column='name')
            tc.batch_sources = {
                'data.xlsx': BatchSourceConfig(
                    file='data.xlsx', mode=RowIterationMode.SEQUENTIAL),
            }
            prj.templates['t.docx'] = tc
            prj.to_file(os.path.join(tmp, 'проект.docxforge'))

            with pytest.raises(GenerationError, match='No documents generated') as exc:
                generate_project(tmp)
            assert exc.value.code == EMPTY_SEQUENTIAL

    def test_m5_render_empty_logs_mapped_warning(self, caplog):
        from docx import Document
        from openpyxl import Workbook

        from docxforge.engine.data_reader import DataReader
        from docxforge.engine.renderer import Renderer
        from docxforge.engine.schema import (
            BatchSourceConfig, FieldMapping, FieldType, RowIterationMode,
        )

        with tempfile.TemporaryDirectory() as tmp:
            data_dir = os.path.join(tmp, 'Данные')
            tmpl_dir = os.path.join(tmp, 'Шаблоны')
            os.makedirs(data_dir, exist_ok=True)
            os.makedirs(tmpl_dir, exist_ok=True)
            doc = Document()
            doc.add_paragraph('Name: {{ name }}')
            doc.save(os.path.join(tmpl_dir, 't.docx'))
            wb = Workbook()
            ws = wb.active
            ws.append(['name'])
            wb.save(os.path.join(data_dir, 'data.xlsx'))
            prj = Project()
            tc = TemplateConfig()
            tc.fields['name'] = FieldMapping(
                type=FieldType.TABLE, file='data.xlsx', column='name')
            tc.batch_sources = {
                'data.xlsx': BatchSourceConfig(
                    file='data.xlsx', mode=RowIterationMode.SEQUENTIAL),
            }
            prj.templates['t.docx'] = tc
            prj.to_file(os.path.join(tmp, 'проект.docxforge'))
            renderer = Renderer(tmp, DataReader())
            renderer.load_project()
            with caplog.at_level(logging.WARNING,
                                 logger='docxforge.engine.render_execute'):
                assert renderer.render('t.docx', {}) == []
            expected = message_for_code(EMPTY_SEQUENTIAL, source='data.xlsx')
            assert any(expected in (r.getMessage() or '')
                       for r in caplog.records)
