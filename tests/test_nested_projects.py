# -*- coding: utf-8 -*-
"""Unit tests for nested employee/project generation (Phase 6).

Scope: nested structure Employee -> Projects per SPEC.md.
Source/generate.py API from Phases 2-4 may be missing; such tests skip
with a clear reason (pending) while existing flat-mode helpers are
covered directly.

Run with: python -m pytest tests/test_nested_projects.py -v
"""

import json
import logging
import os
import sys
import tempfile

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import docxforge.generate as gen_module
from docxforge.engine.schema import (
    BatchSourceConfig,
    FieldMapping,
    FieldType,
    Project,
    RowIterationMode,
    TemplateConfig,
)


logger = logging.getLogger(__name__)

HAS_PARSE_COMPOSITE = hasattr(gen_module, 'parse_composite_template')
HAS_NESTED = hasattr(gen_module, 'create_nested_employee_projects')

COMPOSITE_TEMPLATE = '{{employee}}/{{project_name}}'
FLAT_TEMPLATE = '{{client_name}}_ договор'


def _make_nested_source_project(tmp_dir: str) -> str:
    """Create source project with employee + project_name batch columns."""
    from docx import Document
    from openpyxl import Workbook

    project_dir = os.path.join(tmp_dir, 'source_project')
    data_dir = os.path.join(project_dir, 'Данные')
    tmpl_dir = os.path.join(project_dir, 'Шаблоны')
    os.makedirs(data_dir, exist_ok=True)
    os.makedirs(tmpl_dir, exist_ok=True)

    doc = Document()
    doc.add_paragraph('Employee: {{ employee }}')
    doc.add_paragraph('Project: {{ project_name }}')
    doc.add_paragraph('Client: {{ client_name }}')
    doc.add_paragraph('Doc: {{ doc_number }}')
    doc.save(os.path.join(tmpl_dir, 'report.docx'))

    wb = Workbook()
    ws = wb.active
    ws.append(['employee', 'project_name', 'client_name', 'amount'])
    ws.append(['Иванов Иван', 'Договор_001', 'ООО Альфа', '10000'])
    ws.append(['Иванов Иван', 'Договор_002', 'ООО Бета', '20000'])
    ws.append(['Иванов Иван', 'Договор_003', 'ИП Гамма', '15000'])
    ws.append(['Петров Петр', 'Контракт_001', 'ООО Дельта', '30000'])
    ws.append(['Петров Петр', 'Контракт_002', 'ООО Эпсилон', '40000'])
    ws.append(['Петров Петр', 'Контракт_003', 'ИП Дзета', '50000'])
    wb.save(os.path.join(data_dir, 'batch.xlsx'))
    # Extra file to verify full data-folder copy
    wb2 = Workbook()
    ws2 = wb2.active
    ws2.append(['ref_key', 'ref_value'])
    ws2.append(['k1', 'v1'])
    wb2.save(os.path.join(data_dir, 'reference.xlsx'))

    prj = Project()
    tc = TemplateConfig()
    tc.fields['employee'] = FieldMapping(
        type=FieldType.TABLE, file='batch.xlsx', column='employee')
    tc.fields['project_name'] = FieldMapping(
        type=FieldType.TABLE, file='batch.xlsx', column='project_name')
    tc.fields['client_name'] = FieldMapping(
        type=FieldType.TABLE, file='batch.xlsx', column='client_name')
    tc.fields['amount'] = FieldMapping(
        type=FieldType.TABLE, file='batch.xlsx', column='amount')
    tc.fields['doc_number'] = FieldMapping(
        type=FieldType.COUNTER, start=1, step=1, format='0001')
    tc.fields['org'] = FieldMapping(
        type=FieldType.CONSTANT, value='ORG_VALUE')
    tc.batch_sources = {
        'batch.xlsx': BatchSourceConfig(
            file='batch.xlsx', mode=RowIterationMode.SEQUENTIAL),
    }
    tc.filename_template = '{{ project_name }}.docx'
    tc.directory_template = COMPOSITE_TEMPLATE
    prj.templates['report.docx'] = tc
    prj.to_file(os.path.join(project_dir, 'проект.docxforge'))

    return project_dir


def _make_flat_source_project(tmp_dir: str) -> str:
    """Create minimal flat source project for backward-compat checks."""
    from docx import Document
    from openpyxl import Workbook

    project_dir = os.path.join(tmp_dir, 'flat_source')
    data_dir = os.path.join(project_dir, 'Данные')
    tmpl_dir = os.path.join(project_dir, 'Шаблоны')
    os.makedirs(data_dir, exist_ok=True)
    os.makedirs(tmpl_dir, exist_ok=True)

    doc = Document()
    doc.add_paragraph('Client: {{ client_name }}')
    doc.save(os.path.join(tmpl_dir, 'contract.docx'))

    wb = Workbook()
    ws = wb.active
    ws.append(['client_name'])
    ws.append(['ООО Альфа'])
    ws.append(['ООО Бета'])
    wb.save(os.path.join(data_dir, 'clients.xlsx'))

    prj = Project()
    tc = TemplateConfig()
    tc.fields['client_name'] = FieldMapping(
        type=FieldType.TABLE, file='clients.xlsx', column='client_name')
    tc.fields['org'] = FieldMapping(
        type=FieldType.CONSTANT, value='ORG_VALUE')
    tc.batch_sources = {
        'clients.xlsx': BatchSourceConfig(
            file='clients.xlsx', mode=RowIterationMode.SEQUENTIAL),
    }
    prj.templates['contract.docx'] = tc
    prj.to_file(os.path.join(project_dir, 'проект.docxforge'))

    return project_dir


def test_nested_basic_generation_creates_employee_and_project_folders():
    """2 employees x 3 projects -> nested Employee/Project folders."""
    if not HAS_NESTED:
        pytest.skip('pending: create_nested_employee_projects not implemented (Phase 3)')
    with tempfile.TemporaryDirectory() as tmp:
        source_dir = _make_nested_source_project(tmp)
        output_base = os.path.join(tmp, 'Projects')
        result = gen_module.create_nested_employee_projects(
            source_dir, 'report.docx', COMPOSITE_TEMPLATE, output_base)
        logger.info('nested generation result: %s', result)
        employees = [d for d in os.listdir(output_base)
                     if os.path.isdir(os.path.join(output_base, d))]
        assert len(employees) == 2
        total_projects = 0
        for emp in employees:
            emp_dir = os.path.join(output_base, emp)
            projects = [d for d in os.listdir(emp_dir)
                        if os.path.isdir(os.path.join(emp_dir, d))]
            assert len(projects) == 3
            total_projects += len(projects)
            for proj in projects:
                proj_dir = os.path.join(emp_dir, proj)
                assert os.path.exists(os.path.join(proj_dir, 'проект.docxforge'))
        assert total_projects == 6


def test_nested_composite_template_parsing_detects_parts():
    """Composite '{{employee}}/{{project_name}}' parses into two parts."""
    if not HAS_PARSE_COMPOSITE:
        pytest.skip('pending: parse_composite_template not implemented (Phase 2)')
    employee_part, project_part = gen_module.parse_composite_template(COMPOSITE_TEMPLATE)
    assert 'employee' in employee_part
    assert 'project_name' in project_part


def test_nested_composite_template_parsing_fallback_resolves_both_placeholders():
    """Existing resolver handles composite template with row data (no new API)."""
    from docxforge.generate import _resolve_folder_name_template
    tc = TemplateConfig()
    tc.fields['employee'] = FieldMapping(
        type=FieldType.TABLE, file='batch.xlsx', column='employee')
    tc.fields['project_name'] = FieldMapping(
        type=FieldType.TABLE, file='batch.xlsx', column='project_name')
    row = {'employee': 'Иванов Иван', 'project_name': 'Договор_001'}
    resolved = _resolve_folder_name_template(COMPOSITE_TEMPLATE, row, tc, [])
    assert resolved == 'Иванов Иван/Договор_001'


def test_nested_config_transformation_applies_constant_counter_batch():
    """CONSTANT preserved, COUNTER reset, batch->CONSTANT (existing helper)."""
    from docxforge.generate import _build_project_config
    tc = TemplateConfig()
    tc.fields['project_name'] = FieldMapping(
        type=FieldType.TABLE, file='batch.xlsx', column='project_name')
    tc.fields['doc_number'] = FieldMapping(
        type=FieldType.COUNTER, start=5, step=2, format='0001')
    tc.fields['org'] = FieldMapping(
        type=FieldType.CONSTANT, value='ORG_VALUE')
    tc.batch_sources = {
        'batch.xlsx': BatchSourceConfig(
            file='batch.xlsx', mode=RowIterationMode.SEQUENTIAL),
    }
    row = {'project_name': 'Договор_001'}
    new_config = _build_project_config(tc, row, 'batch.xlsx')
    const_fm = new_config.fields['org']
    assert const_fm.type == FieldType.CONSTANT
    assert const_fm.value == 'ORG_VALUE'
    counter_fm = new_config.fields['doc_number']
    assert counter_fm.type == FieldType.COUNTER
    assert counter_fm.start == 5
    assert counter_fm.step == 2
    assert counter_fm.format == '0001'
    assert new_config.batch_sources['batch.xlsx'].mode == RowIterationMode.CONSTANT


@pytest.mark.xfail(
    strict=False,
    reason='TABLE->CONSTANT never fires: _build_project_config checks '
           '"fm.file in row_data" but row_data is column-keyed '
           '(Phase 4 fix pending, source must not change in Phase 6)',
)
def test_nested_config_transformation_applies_table_to_constant():
    """TABLE field becomes CONSTANT with row value (SPEC, currently gaps)."""
    from docxforge.generate import _build_project_config
    tc = TemplateConfig()
    tc.fields['project_name'] = FieldMapping(
        type=FieldType.TABLE, file='batch.xlsx', column='project_name')
    tc.batch_sources = {
        'batch.xlsx': BatchSourceConfig(
            file='batch.xlsx', mode=RowIterationMode.SEQUENTIAL),
    }
    row = {'project_name': 'Договор_001'}
    new_config = _build_project_config(tc, row, 'batch.xlsx')
    table_fm = new_config.fields['project_name']
    assert table_fm.type == FieldType.CONSTANT
    assert table_fm.value == 'Договор_001'


def test_nested_flat_mode_backward_compatible_still_works():
    """Flat template (no '/') still generates projects via existing entry point."""
    with tempfile.TemporaryDirectory() as tmp:
        source_dir = _make_flat_source_project(tmp)
        projects_dir, created = gen_module.create_projects_from_template(
            source_dir, 'contract.docx', FLAT_TEMPLATE)
        assert created == 2
        assert os.path.isdir(projects_dir)
        generated = [d for d in os.listdir(projects_dir)
                     if os.path.isdir(os.path.join(projects_dir, d))]
        assert len(generated) == 2
        for proj in generated:
            assert os.path.exists(
                os.path.join(projects_dir, proj, 'проект.docxforge'))
