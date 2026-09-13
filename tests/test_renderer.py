# -*- coding: utf-8 -*-
"""Tests for renderer bugs — text order, formatting preservation."""

import os, tempfile, zipfile, re

from lxml import etree
import openpyxl
from docx import Document
from docx.shared import Pt

from docxforge.engine.schema import (
    create_project, Project, TemplateConfig, FieldMapping, FieldType,
    CycleMapping, AggregationMapping, AggregationFunction,
)
from docxforge.engine import Renderer
from docxforge.engine.data_reader import DataReader

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def _make_dirs(tmp):
    os.makedirs(os.path.join(tmp, 'Данные'), exist_ok=True)
    os.makedirs(os.path.join(tmp, 'Шаблоны'), exist_ok=True)


class TestTextOrder:
    """Test: {{ placeholder }} must not reorder surrounding text."""

    def test_label_before_placeholder(self):
        with tempfile.TemporaryDirectory() as tmp:
            _make_dirs(tmp)
            doc = Document()
            p = doc.add_paragraph()
            p.add_run('Покупатель: ')
            p.add_run('{{ фио }}')
            doc.save(os.path.join(tmp, 'Шаблоны', 'test.docx'))

            prj = Project()
            tc = TemplateConfig()
            tc.fields['фио'] = FieldMapping(type=FieldType.CONSTANT, value='Петров П.П.')
            prj.templates['test.docx'] = tc
            prj.to_file(os.path.join(tmp, 'проект.docxforge'))

            reader = DataReader()
            renderer = Renderer(tmp, reader)
            renderer.load_project()
            outputs = renderer.render('test.docx', {})

            with zipfile.ZipFile(outputs[0], 'r') as zf:
                doc_xml = etree.parse(zf.open('word/document.xml'))
                text = ''
                for p_el in doc_xml.findall('.//{%s}p' % W):
                    for r in p_el.findall('{%s}r' % W):
                        for t_el in r.findall('{%s}t' % W):
                            if t_el.text:
                                text += t_el.text

            assert text == 'Покупатель: Петров П.П.', \
                'FAIL: got %r, expected "Покупатель: Петров П.П."' % text

    def test_text_before_and_after_placeholder(self):
        with tempfile.TemporaryDirectory() as tmp:
            _make_dirs(tmp)
            doc = Document()
            p = doc.add_paragraph()
            p.add_run('Начало ')
            p.add_run('{{ поле }}')
            p.add_run(' конец')
            doc.save(os.path.join(tmp, 'Шаблоны', 'test.docx'))

            prj = Project()
            tc = TemplateConfig()
            tc.fields['поле'] = FieldMapping(type=FieldType.CONSTANT, value='СЕРЕДИНА')
            prj.templates['test.docx'] = tc
            prj.to_file(os.path.join(tmp, 'проект.docxforge'))

            reader = DataReader()
            renderer = Renderer(tmp, reader)
            renderer.load_project()
            outputs = renderer.render('test.docx', {})

            with zipfile.ZipFile(outputs[0], 'r') as zf:
                doc_xml = etree.parse(zf.open('word/document.xml'))
                text = ''
                for p_el in doc_xml.findall('.//{%s}p' % W):
                    for r in p_el.findall('{%s}r' % W):
                        for t_el in r.findall('{%s}t' % W):
                            if t_el.text:
                                text += t_el.text

            assert text == 'Начало СЕРЕДИНА конец', \
                'FAIL: got %r, expected "Начало СЕРЕДИНА конец"' % text

    def test_multiple_placeholders_in_order(self):
        with tempfile.TemporaryDirectory() as tmp:
            _make_dirs(tmp)
            doc = Document()
            p = doc.add_paragraph()
            p.add_run('{{ a }}')
            p.add_run(' ')
            p.add_run('{{ b }}')
            p.add_run(' ')
            p.add_run('{{ c }}')
            doc.save(os.path.join(tmp, 'Шаблоны', 'test.docx'))

            prj = Project()
            tc = TemplateConfig()
            tc.fields['a'] = FieldMapping(type=FieldType.CONSTANT, value='A')
            tc.fields['b'] = FieldMapping(type=FieldType.CONSTANT, value='B')
            tc.fields['c'] = FieldMapping(type=FieldType.CONSTANT, value='C')
            prj.templates['test.docx'] = tc
            prj.to_file(os.path.join(tmp, 'проект.docxforge'))

            reader = DataReader()
            renderer = Renderer(tmp, reader)
            renderer.load_project()
            outputs = renderer.render('test.docx', {})

            with zipfile.ZipFile(outputs[0], 'r') as zf:
                doc_xml = etree.parse(zf.open('word/document.xml'))
                text = ''
                for p_el in doc_xml.findall('.//{%s}p' % W):
                    for r in p_el.findall('{%s}r' % W):
                        for t_el in r.findall('{%s}t' % W):
                            if t_el.text:
                                text += t_el.text

            assert text == 'A B C', \
                'FAIL: got %r, expected "A B C"' % text

    def test_real_contract_text_order(self):
        with tempfile.TemporaryDirectory() as tmp:
            data_dir = os.path.join(tmp, 'Данные')
            tmpl_dir = os.path.join(tmp, 'Шаблоны')
            os.makedirs(data_dir)
            os.makedirs(tmpl_dir)

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = 'клиенты'
            ws.append(['название', 'ИНН', 'фио'])
            ws.append(['ООО Альфа', '7712345678', 'Петров П.П.'])
            wb.save(os.path.join(data_dir, 'клиенты.xlsx'))

            doc = Document()
            style = doc.styles['Normal']
            style.font.name = 'Times New Roman'
            style.font.size = Pt(12)

            p = doc.add_paragraph()
            p.add_run('Покупатель: ')
            p.add_run('{{ фио }}')

            p2 = doc.add_paragraph()
            p2.add_run('ИНН: ')
            p2.add_run('{{ ИНН }}')

            p3 = doc.add_paragraph()
            p3.add_run('Продавец: ')
            p3.add_run('{{ продавец }}')

            doc.save(os.path.join(tmpl_dir, 'contract.docx'))

            prj = Project()
            tc = TemplateConfig()
            tc.fields['фио'] = FieldMapping(
                type=FieldType.TABLE, file='клиенты.xlsx', column='фио')
            tc.fields['ИНН'] = FieldMapping(
                type=FieldType.TABLE, file='клиенты.xlsx',
                column='ИНН', linked_to='фио')
            tc.fields['продавец'] = FieldMapping(
                type=FieldType.CONSTANT, value='Иванов И.И.')
            prj.templates['contract.docx'] = tc
            prj.to_file(os.path.join(tmp, 'проект.docxforge'))

            reader = DataReader()
            renderer = Renderer(tmp, reader)
            renderer.load_project()
            outputs = renderer.render('contract.docx', {})

            with zipfile.ZipFile(outputs[0], 'r') as zf:
                doc_xml = etree.parse(zf.open('word/document.xml'))
                lines = []
                for p_el in doc_xml.findall('.//{%s}p' % W):
                    line = ''
                    for r in p_el.findall('{%s}r' % W):
                        for t_el in r.findall('{%s}t' % W):
                            if t_el.text:
                                line += t_el.text
                    if line:
                        lines.append(line)

            assert lines[0] == 'Покупатель: Петров П.П.', \
                'FAIL line 0: got %r' % lines[0]
            assert lines[1] == 'ИНН: 7712345678', \
                'FAIL line 1: got %r' % lines[1]
            assert lines[2] == 'Продавец: Иванов И.И.', \
                'FAIL line 2: got %r' % lines[2]
