# -*- coding: utf-8 -*-
"""Regression tests for specs/001-review/REPORT.md BLOCKER fixes (render path).

Covers: B2 (merge keeps w:br/w:drawing runs in place), B4 (empty table_data
keeps template row), B3 (row placeholder search over merged row text),
B7 ([Content_Types].xml Override for images), B6 (all images per paragraph
inserted with unique rId/media, no overwrite of template media).
"""

import os
import struct
import tempfile
import zipfile
import zlib

from docx import Document
from lxml import etree

from docxforge.engine.schema import Project, TemplateConfig, FieldMapping, FieldType, CycleMapping
from docxforge.engine import Renderer
from docxforge.engine.data_reader import DataReader
from docxforge.engine.merge import merge_and_replace_paragraph, expand_table_cycle
from docxforge.engine.xml_utils import row_contains_placeholder, row_has_placeholders
from docxforge.engine.image_utils import add_image_to_zdata

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W_NS = "{%s}" % W
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"

DATA = "Данные"
TMPL = "Шаблоны"
PJ = "проект.docxforge"


def _make_dirs(tmp):
    os.makedirs(os.path.join(tmp, DATA), exist_ok=True)
    os.makedirs(os.path.join(tmp, TMPL), exist_ok=True)


def _make_minimal_png():
    def chunk(ct, d):
        c = ct + d
        return struct.pack(">I", len(d)) + c + struct.pack(">I", zlib.crc32(c) & 0xFFFFFFFF)
    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    raw = b"\x00\xff\x00\x00"
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr)
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def _run_with_text(text):
    r = etree.Element(W_NS + "r")
    t = etree.SubElement(r, W_NS + "t")
    t.text = text
    return r


def _para_text(p_el):
    return "".join(t.text or "" for t in p_el.iter(W_NS + "t"))


class TestReviewB2BrRunsKeptInPlace:
    def test_merge_br_run_stays_before_following_text(self):
        """Pure w:br run is preserved at its position, not moved to the end."""
        p = etree.Element(W_NS + "p")
        p.append(_run_with_text("{{ name }}"))
        br_run = etree.Element(W_NS + "r")
        etree.SubElement(br_run, W_NS + "br")
        p.append(br_run)
        p.append(_run_with_text(" after"))

        merge_and_replace_paragraph(p, {"name": "X"})

        children = list(p)
        br_pos = children.index(br_run)
        texts = [_para_text(c) for c in children if c.tag == W_NS + "r"]
        assert len(p.findall(".//{%s}br" % W)) == 1
        assert "".join(texts) == "X after"
        assert br_pos == 1, "br run moved, order: %r" % (texts,)

    def test_merge_mixed_run_br_survives(self):
        """Run with text + w:br keeps its br, text is replaced in place."""
        p = etree.Element(W_NS + "p")
        r = etree.Element(W_NS + "r")
        t1 = etree.SubElement(r, W_NS + "t")
        t1.text = "Hi "
        etree.SubElement(r, W_NS + "br")
        t2 = etree.SubElement(r, W_NS + "t")
        t2.text = "{{ name }}!"
        p.append(r)

        merge_and_replace_paragraph(p, {"name": "X"})

        assert len(p.findall(".//{%s}br" % W)) == 1
        assert _para_text(p) == "Hi X!"


class TestReviewB4EmptyTableDataKeepsRow:
    def test_expand_table_cycle_empty_data_keeps_template_row(self):
        with tempfile.TemporaryDirectory() as tmp:
            _make_dirs(tmp)
            doc = Document()
            tbl = doc.add_table(rows=2, cols=1)
            tbl.cell(0, 0).text = "Header"
            tbl.cell(1, 0).text = "{{ item }}"
            path = os.path.join(tmp, TMPL, "t.docx")
            doc.save(path)
            with zipfile.ZipFile(path) as zf:
                root = etree.fromstring(zf.read("word/document.xml"))
            tbl_el = root.find(".//%s" % (W_NS + "tbl"))
            cycle = CycleMapping(table="f.xlsx", columns={"item": "col"})
            expand_table_cycle(tbl_el, cycle, [], {})
            rows = tbl_el.findall(W_NS + "tr")
            assert len(rows) == 2
            assert "{{ item }}" in "".join(
                t.text or "" for t in rows[1].iter(W_NS + "t"))


class TestReviewB3SplitPlaceholderInRow:
    def _split_row(self):
        row = etree.Element(W_NS + "tr")
        tc = etree.SubElement(row, W_NS + "tc")
        p = etree.SubElement(tc, W_NS + "p")
        p.append(_run_with_text("{{"))
        p.append(_run_with_text("item"))
        p.append(_run_with_text("}}"))
        return row

    def test_review_row_contains_placeholder_split_across_runs(self):
        assert row_contains_placeholder(self._split_row(), "item") is True
        assert row_contains_placeholder(self._split_row(), "other") is False

    def test_review_row_has_placeholders_split_across_runs(self):
        assert row_has_placeholders(self._split_row()) is True


class TestReviewB7ContentTypeOverride:
    def _zdata_no_png_default(self):
        ct = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            "</Types>"
        ).encode("utf-8")
        return {"[Content_Types].xml": ct}

    def test_review_add_image_adds_override(self):
        with tempfile.TemporaryDirectory() as tmp:
            img = os.path.join(tmp, "a.png")
            with open(img, "wb") as f:
                f.write(_make_minimal_png())
            zdata = self._zdata_no_png_default()
            mp, ct = add_image_to_zdata(zdata, img, 1)
            assert mp == "word/media/image1.png"
            root = etree.fromstring(zdata["[Content_Types].xml"])
            overrides = [el for el in root
                         if el.tag == "{%s}Override" % CT_NS
                         and el.get("PartName") == "/word/media/image1.png"]
            assert len(overrides) == 1
            assert overrides[0].get("ContentType") == "image/png"

    def test_review_add_image_skips_override_when_default_exists(self):
        with tempfile.TemporaryDirectory() as tmp:
            img = os.path.join(tmp, "a.png")
            with open(img, "wb") as f:
                f.write(_make_minimal_png())
            zdata = self._zdata_no_png_default()
            root = etree.fromstring(zdata["[Content_Types].xml"])
            d = etree.SubElement(root, "{%s}Default" % CT_NS)
            d.set("Extension", "png")
            d.set("ContentType", "image/png")
            zdata["[Content_Types].xml"] = etree.tostring(
                root, xml_declaration=True, encoding="UTF-8", standalone=True)
            add_image_to_zdata(zdata, img, 1)
            root2 = etree.fromstring(zdata["[Content_Types].xml"])
            assert not [el for el in root2
                        if el.tag == "{%s}Override" % CT_NS]


class TestReviewB6MultipleImagesUnique:
    def _render(self, tmp, doc, fields):
        _make_dirs(tmp)
        doc.save(os.path.join(tmp, TMPL, "t.docx"))
        prj = Project()
        tc = TemplateConfig()
        tc.fields.update(fields)
        prj.templates["t.docx"] = tc
        prj.to_file(os.path.join(tmp, PJ))
        renderer = Renderer(tmp, DataReader())
        renderer.load_project()
        return renderer.render("t.docx", {})

    def test_review_two_images_in_one_paragraph(self):
        with tempfile.TemporaryDirectory() as tmp:
            _make_dirs(tmp)
            png = _make_minimal_png()
            p1 = os.path.join(tmp, DATA, "a.png")
            p2 = os.path.join(tmp, DATA, "b.png")
            for path, payload in ((p1, png), (p2, png + b"\x00")):
                with open(path, "wb") as f:
                    f.write(payload)
            doc = Document()
            doc.add_paragraph("{{ image:img1 }} mid {{ image:img2 }}")
            out = self._render(tmp, doc, {
                "img1": FieldMapping(type=FieldType.IMAGE, value=os.path.join(DATA, "a.png")),
                "img2": FieldMapping(type=FieldType.IMAGE, value=os.path.join(DATA, "b.png")),
            })
            with zipfile.ZipFile(out[0]) as zf:
                names = zf.namelist()
                assert "word/media/image1.png" in names
                assert "word/media/image2.png" in names
                assert zf.read("word/media/image1.png") == png
                assert zf.read("word/media/image2.png") == png + b"\x00"
                doc_xml = etree.parse(zf.open("word/document.xml"))
                assert len(doc_xml.findall(".//{%s}drawing" % W)) == 2
                rels = etree.parse(zf.open("word/_rels/document.xml.rels"))
                rns = "http://schemas.openxmlformats.org/package/2006/relationships"
                ids = [r.get("Id") for r in rels.findall("{%s}Relationship" % rns)]
                assert len(set(ids)) == len(ids) >= 2
                ct_root = etree.fromstring(zf.read("[Content_Types].xml"))
                defaults = {(el.get("Extension") or "").lower() for el in ct_root
                            if el.tag == "{%s}Default" % CT_NS}
                overrides = {el.get("PartName") for el in ct_root
                             if el.tag == "{%s}Override" % CT_NS}
                for mp in ("word/media/image1.png", "word/media/image2.png"):
                    ext = mp.rsplit(".", 1)[1].lower()
                    assert ext in defaults or ("/" + mp) in overrides

    def test_review_template_media_not_overwritten(self):
        with tempfile.TemporaryDirectory() as tmp:
            tpl_png = _make_minimal_png()
            tpl_img = os.path.join(tmp, "tpl.png")
            with open(tpl_img, "wb") as f:
                f.write(tpl_png)
            new_png = tpl_png + b"\x01\x02"
            os.makedirs(os.path.join(tmp, DATA), exist_ok=True)
            new_img = os.path.join(tmp, DATA, "logo.png")
            with open(new_img, "wb") as f:
                f.write(new_png)
            doc = Document()
            doc.add_picture(tpl_img)  # template ships word/media/image1.png
            doc.add_paragraph("{{ image:logo }}")
            _make_dirs(tmp)
            doc.save(os.path.join(tmp, TMPL, "t.docx"))
            with zipfile.ZipFile(os.path.join(tmp, TMPL, "t.docx")) as zf:
                assert "word/media/image1.png" in zf.namelist()
            out = self._render(tmp, doc, {
                "logo": FieldMapping(type=FieldType.IMAGE, value=os.path.join(DATA, "logo.png")),
            })
            with zipfile.ZipFile(out[0]) as zf:
                assert zf.read("word/media/image1.png") == tpl_png
                assert zf.read("word/media/image2.png") == new_png
